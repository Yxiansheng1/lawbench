"""注记 致B-ORCH-注记-LibreOffice三件事-20260929-2218：LibreOffice 配置目录在 <应用数据>\\临时\\lo\\<8 位随机>，
用完删除、启动清理、路径超长不调用、0xC0000409 报"转换程序异常退出"；长案件路径下 doc / xls / 无缓存值 xlsx 都成功；
有外链关系的 xlsx 不交给 LibreOffice 重算。"""
from __future__ import annotations

import json
import os
import pathlib
import shutil
import sys
import zipfile

import openpyxl
import pytest

from lawbench.config import REPO_ROOT
from lawbench.ingest import REASONS, links
from lawbench.ingest import libreoffice as lo

from fakes import CountingListener

FIXTURES = REPO_ROOT / "tests" / "fixtures"
needs_lo = pytest.mark.skipif(lo.find_soffice() is None, reason="本机没有 LibreOffice")


def scan(client, root: pathlib.Path) -> dict:
    cid = client.post("/api/case/open", json={"path": str(root)}).json()["value"]["case_id"]
    r = client.post("/api/materials/scan", json={"case_id": cid}).json()
    assert r["ok"], r
    idx = json.loads((root / "工作区" / "材料" / "index.json").read_text(encoding="utf-8"))
    return {m["name"]: m for m in idx["materials"]}


@pytest.fixture(scope="module")
def samples(tmp_path_factory):
    if lo.find_soffice() is None:
        pytest.skip("本机没有 LibreOffice")
    base = tmp_path_factory.mktemp("s")
    with lo.Converter(base / "t").session() as s:
        (base / "通知.doc").write_bytes(s.convert(FIXTURES / "tender-01" / "补充通知.docx", "doc").read_bytes())
        (base / "流水.xls").write_bytes(s.convert(FIXTURES / "civil-01" / "银行流水.xlsx", "xls").read_bytes())
    shutil.copy(FIXTURES / "civil-01" / "银行流水.xlsx", base / "银行流水.xlsx")  # 汇总!B4 无缓存值
    return base


# ---------- 配置目录的位置、删除、启动清理 ----------

@needs_lo
def test_profile_under_appdata_and_removed(make_client, appdata, cases_dir, samples, monkeypatch):
    used: list[pathlib.Path] = []
    real = lo.write_profile
    monkeypatch.setattr(lo, "write_profile", lambda p: (used.append(p), real(p)))
    root = cases_dir / "位置"
    root.mkdir()
    shutil.copy(samples / "通知.doc", root / "通知.doc")
    mats = scan(make_client(), root)
    assert mats["通知"]["status"] == "parsed"
    base = appdata / "临时" / "lo"
    assert used and all(p.parent.parent == base.resolve() for p in used)   # <应用数据>\临时\lo\<8 位>\p
    assert all(len(p.parent.name) == 8 for p in used)
    assert list(base.iterdir()) == []                                      # 用完删除
    assert list((root / "工作区" / "临时").iterdir()) == []                  # 副本和结果也删了


def test_startup_cleanup(make_client, appdata):
    stale = appdata / "临时" / "lo" / "abcd1234" / "p" / "user"
    stale.mkdir(parents=True)
    (stale / "registrymodifications.xcu").write_text("最近打开的文件记录", encoding="utf-8")
    make_client()
    assert list((appdata / "临时" / "lo").iterdir()) == []


def test_remove_tree_failure_is_logged(tmp_path, monkeypatch):
    from lawbench import logs
    logs.setup(tmp_path / "ad")
    d = tmp_path / "删不掉"
    d.mkdir()
    monkeypatch.setattr(lo.shutil, "rmtree", lambda *a, **k: None)
    monkeypatch.setattr(lo.time, "sleep", lambda s: None)
    assert lo.remove_tree(d) is False
    logs.close()
    text = "".join(p.read_text(encoding="utf-8") for p in (tmp_path / "ad" / "logs").glob("*"))
    assert '"op": "cleanup", "status": "fail"' in text


# ---------- 路径超长不调用；崩溃退出码 ----------

def test_profile_path_too_long_not_called(tmp_path, samples, monkeypatch):
    started = []
    monkeypatch.setattr(lo.subprocess, "Popen", lambda *a, **k: started.append(a) or (_ for _ in ()).throw(OSError()))
    long_base = tmp_path / ("很长的应用数据目录" * 8)
    conv = lo.Converter(tmp_path / "case-temp", lo_base=long_base)
    assert len(str(long_base / "12345678" / "p")) > lo.MAX_PROFILE_PATH
    with pytest.raises(lo.ParseError) as ei:
        with conv.session() as s:
            s.convert(samples / "通知.doc", "docx")
    assert ei.value.reason == "appdata_too_long" and ei.value.message == REASONS["appdata_too_long"]
    assert started == [] and not long_base.exists()


@pytest.mark.skipif(os.name != "nt", reason="Windows 退出码")
def test_crash_exit_code_reason(make_client, cases_dir, tmp_path, monkeypatch):
    cmd = tmp_path / "crash.cmd"
    cmd.write_text(f'@"{sys.executable}" -c "import sys; sys.exit(-1073740791)"\r\n', encoding="mbcs")
    monkeypatch.setattr(lo, "find_soffice", lambda: str(cmd))
    root = cases_dir / "崩溃"
    root.mkdir()
    shutil.copy(FIXTURES / "civil-01" / "借条.docx", root / "旧.doc")
    m = scan(make_client(), root)["旧"]
    assert m["status"] == "failed" and m["error"] == REASONS["converter_crashed"]
    assert "损坏" not in m["error"] and "加密" not in m["error"]


# ---------- 很长的案件路径（>150 字符，含中文）下 doc / xls / 无缓存值 xlsx 都成功 ----------

@needs_lo
def test_long_case_path(make_client, cases_dir, samples):
    need = 160 - len(str(cases_dir)) - 1
    root = cases_dir / ("甲某诉乙某买卖合同纠纷案" + "卷" * max(1, need - 11))
    root.mkdir()
    assert len(str(root)) > 150
    for n in ("通知.doc", "流水.xls", "银行流水.xlsx"):
        shutil.copy(samples / n, root / n)
    mats = scan(make_client(), root)
    assert {n: m["status"] for n, m in mats.items()} == {"通知": "parsed", "流水": "parsed", "银行流水": "parsed"}
    text = (root / mats["银行流水"]["text_path"]).read_text(encoding="utf-8")
    assert "| 4 | 已收利息 | 2200 |" in text   # 无缓存值的公式经 LibreOffice 重算


# ---------- xlsx：有外链关系的不交给 LibreOffice 重算 ----------

def _xlsx_with_external_image(path: pathlib.Path, url: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "表一"
    ws["A1"] = 10
    ws["A2"] = 20
    ws["A3"] = "=SUM(A1:A2)"   # openpyxl 不写缓存值
    wb.save(path)
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("xl/worksheets/_rels/sheet1.xml.rels",
                      '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                      '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                      '<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                      f'relationships/image" Target="{url}/pic.png" TargetMode="External"/></Relationships>')
    tmp.replace(path)


def test_xlsx_detector(tmp_path):
    p = tmp_path / "外链.xlsx"
    _xlsx_with_external_image(p, "http://127.0.0.1:1")
    assert links.xlsx_has_external_rels(p) is True
    assert links.xlsx_has_external_rels(FIXTURES / "civil-01" / "银行流水.xlsx") is False
    q = tmp_path / "超链接.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "网址"
    wb.active["A1"].hyperlink = "https://example.invalid/"
    wb.save(q)
    assert links.xlsx_has_external_rels(q) is False   # 超链接不算


def test_xlsx_external_not_recalculated(make_client, cases_dir, monkeypatch):
    started = []
    real_popen = lo.subprocess.Popen
    monkeypatch.setattr(lo.subprocess, "Popen", lambda *a, **k: started.append(a) or real_popen(*a, **k))
    with CountingListener() as lis:
        root = cases_dir / "外链xlsx"
        root.mkdir()
        _xlsx_with_external_image(root / "对方报表.xlsx", lis.url)
        m = scan(make_client(), root)["对方报表"]
        assert lis.count == 0
    assert started == []                                   # 没有起 soffice
    assert m["status"] == "parsed"
    text = (root / m["text_path"]).read_text(encoding="utf-8")
    assert "| 3 | =SUM(A1:A2) |" in text                   # 没有缓存值的单元格只写公式
