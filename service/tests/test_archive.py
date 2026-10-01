"""案卷归档（Spec 12.4；工单 T23 第 2–5 步）：匹配、归档方案、生成。

用例里转换程序设为 libreoffice（不动本机的真 Word；Word / WPS 的调度另见 test_convert）。
"""
from __future__ import annotations

import hashlib
import io
import json
import pathlib
import shutil
import sys

import docx as pydocx
import pytest
from pypdf import PdfReader, PdfWriter

from lawbench.archive import build as B
from lawbench.archive import documents as D
from lawbench.archive import match as am
from lawbench.config import REPO_ROOT
from lawbench.ingest import libreoffice as lo
from lawbench.office import convert as C

from t8_helpers import FIXTURES, Env, fail, ok, validator

CLOSED = FIXTURES / "closed-01"
FILES = {str(p.relative_to(CLOSED)).replace("\\", "/"): p for p in CLOSED.rglob("*") if p.is_file()}
SKILLS = REPO_ROOT / "skills"
needs_lo = pytest.mark.skipif(lo.find_soffice() is None, reason="本机没有 LibreOffice")
BUILD = validator("api/archive_build.schema.json", "#/$defs/response")


@pytest.fixture
def env(tmp_path):
    e = Env(tmp_path, FILES)
    s = ok(e.client.get("/api/settings"), "api/settings.schema.json")
    s["converter"] = "libreoffice"
    s["profile"]["lawyer_name"] = "孙律师"
    ok(e.client.put("/api/settings", json=s), "api/settings.schema.json")
    yield e
    e.close()


@pytest.fixture
def tid(env):
    return env.begin("sess-archive")["task_id"]


def originals(env) -> dict:
    return {str(p.relative_to(env.root)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in env.root.rglob("*") if p.is_file() and p.relative_to(env.root).parts[0] not in ("工作区", "成果")}


def plan(**kw) -> dict:
    p = {"catalog": "民事行政卷", "client": "赵某丁", "opponent": "钱某戊", "cause": "买卖合同纠纷", "lawyer": None,
         "entrust_date": "2026-01-05", "close_date": "2026-06-30", "jzl_no": "（2026）粤连越深圳民字第0001号",
         "result": "调解", "summary": "委托人与钱某戊因货款发生争议，委托人起诉要求支付货款。",
         "opinion": "委托人有权向钱某戊主张货款，经法院主持达成调解。", "fee_settled": True,
         "items": [{"code": 1, "name": "民事委托代理合同", "materials": ["委托代理合同"]},
                   {"code": 2, "name": "收费发票", "materials": ["律师费发票"]},
                   {"code": 3, "name": "授权委托书", "materials": ["授权委托书"]},
                   {"code": 4, "name": "民事起诉状", "materials": ["民事起诉状"]},
                   {"code": 7, "name": "证据材料", "materials": ["聊天记录", "送货单", "对账单"]},
                   {"code": 14, "name": "民事调解书", "materials": ["民事调解书"]}]}
    p.update(kw)
    return p


def save_plan(env, tid, p) -> dict:
    return env.tool_ok(tid, "case_save_archive_plan", p)


def build(env, tid, p, plan_rel=None):
    return env.client.post("/api/archive/build", json={
        "case_id": env.case_id, "task_id": tid, "plan": plan_rel or f"工作区/任务/{tid}/归档方案.json", "confirmed": p})


def pdf_pages(path: pathlib.Path) -> int:
    return len(PdfReader(str(path)).pages)


# ---------------------------------------------------------------- 匹配

def test_match_closed01(env, tid):
    v = env.tool_ok(tid, "case_archive_match", {"catalog": "民事行政卷"})
    items = {it["code"]: it for it in v["items"]}
    assert 15 not in items                                             # 结案报告（generated）不参与匹配
    got = {c: [(m["name"], m["folder"], m["reason"]) for m in it["matched"]] for c, it in items.items() if it["matched"]}
    assert got == {
        1: [("委托代理合同", "01委托手续", "文件夹和关键词")],
        2: [("律师费发票", "01委托手续", "文件夹和关键词")],
        3: [("授权委托书", "01委托手续", "文件夹和关键词")],
        4: [("民事起诉状", "03一审", "文件夹和关键词")],
        7: [("对账单", "03一审", "文件夹"), ("聊天记录", "03一审", "文件夹"), ("送货单", "03一审", "文件夹")],
        14: [("民事调解书", "03一审", "文件夹和关键词")],
    }
    assert v["unmatched"] == []
    assert {(i["name"], i["reason"]) for i in v["ignored"]} == {
        ("微信图片_20260301", "命名混乱"), ("03一审/我方文件/~$起诉状.docx", "临时文件")}


def test_match_rules_unit(tmp_path):
    cat = {"id": "民事行政卷", "attachment": "附件3", "renumber": False, "items": [
        {"code": 1, "name": "甲", "required": True, "keywords": ["调解"], "folders": ["A"]},
        {"code": 2, "name": "乙", "required": True, "keywords": ["调解书"], "folders": ["B"]},
        {"code": 3, "name": "丙", "required": False, "keywords": ["丙"], "folders": ["C/我方", "C/对方"]},
        {"code": 4, "name": "丁", "required": False, "keywords": ["丁"], "folders": ["C/我方"]},
        {"code": 9, "name": "结案报告", "required": True, "keywords": ["结案报告"], "folders": ["A"], "generated": True}]}

    def m(rel, status="parsed"):
        return {"name": pathlib.PurePosixPath(rel).stem, "rel_path": rel, "status": status}
    idx = {"materials": [m("A/民事调解书.pdf"), m("X/民事调解书2.pdf"), m("C/我方/x.pdf"), m("C/对方/y.pdf"),
                         m("D/结案报告.docx"), m("A/IMG_0012.jpg"), m("A/坏.pdf", "failed"), m("z.pdf")]}
    v = am.match(str(tmp_path), idx, cat)
    got = {it["code"]: [(x["name"], x["reason"]) for x in it["matched"]] for it in v["items"]}
    assert got[1] == [("民事调解书", "文件夹和关键词")]                    # 两项都命中：先看文件夹
    assert got[2] == [("民事调解书2", "关键词")]                           # 都不在文件夹里：取最长的关键词
    assert got[3] == [("y", "文件夹")] and got[4] == []                   # C/我方 对应两项 → 不按文件夹归
    assert 9 not in got and set(v["unmatched"]) == {"x", "结案报告", "z"}
    assert {(i["name"], i["reason"]) for i in v["ignored"]} == {("IMG_0012", "命名混乱"), ("坏", "加密或无法读取")}


# ---------------------------------------------------------------- 归档方案

def test_save_plan(env, tid):
    p = plan(items=[it for it in plan()["items"] if it["code"] != 3], result=None, jzl_no=None)
    v = save_plan(env, tid, p)
    assert v["path"] == f"工作区/任务/{tid}/归档方案.json"
    assert v["missing_required"] == [{"code": 3, "name": "授权委托书"}]   # 结案报告（generated）不算缺失
    assert any("办案结果" in w for w in v["warnings"]) and any("金助理" in w for w in v["warnings"])
    saved = json.loads((env.root / v["path"]).read_text(encoding="utf-8"))
    assert saved == p
    assert not list(validator("tools/case_save_archive_plan.schema.json", "#/$defs/plan").iter_errors(saved))


@pytest.mark.parametrize("change,code", [
    ({"items": [{"code": 1, "name": "x", "materials": ["没有这份"]}]}, "MATERIAL_NOT_FOUND"),
    ({"items": [{"code": 15, "name": "结案报告", "materials": ["委托代理合同"]}]}, "INVALID_ARGUMENT"),
    ({"items": [{"code": 99, "name": "x", "materials": ["委托代理合同"]}]}, "INVALID_ARGUMENT"),
    ({"items": [{"code": 1, "name": "x", "materials": ["委托代理合同"]},
                {"code": 1, "name": "y", "materials": ["授权委托书"]}]}, "INVALID_ARGUMENT"),
    ({"result": "大胜"}, "INVALID_ARGUMENT"),
])
def test_save_plan_errors(env, tid, change, code):
    fail(env.tool(tid, "case_save_archive_plan", plan(**change)), code)


# ---------------------------------------------------------------- 生成

def app_rows(path: pathlib.Path) -> list[tuple[str, str]]:
    t = pydocx.Document(str(path)).tables[0]
    return [(r.cells[0].text.strip(), r.cells[1].text.strip()) for r in t.rows[1:]]


@needs_lo
def test_build_closed01(env, tid):
    before = originals(env)
    p = plan()
    save_plan(env, tid, p)
    r = build(env, tid, p)
    body = r.json()
    assert not list(BUILD.iter_errors(body)), body
    v = body["value"]
    folder = env.root / v["folder"]
    assert v["folder"] == "成果/归档/赵某丁与钱某戊案件归档" and v["converter"] == "libreoffice"
    kinds = {f["kind"]: env.root / f["path"] for f in v["files"]}
    assert set(kinds) == {"卷宗", "发票", "立卷申请书", "结案报告", "归档目录"}     # 无缺失、律师费已结清：没有情况说明

    # 卷宗：总页数 = 各材料页数之和（含结案报告），页码连续、与 page_ranges 一致
    mats = {n: CLOSED / rel for rel, n in [(r_, pathlib.PurePosixPath(r_).stem) for r_ in FILES]}
    vol = PdfReader(str(kinds["卷宗"]))
    total = len(vol.pages)
    expect_codes = [1, 2, 3, 4, 7, 14, 15]
    assert [x["code"] for x in v["page_ranges"]] == expect_codes
    assert v["page_ranges"][0]["from"] == 1 and v["page_ranges"][-1]["to"] == total
    for a, b in zip(v["page_ranges"], v["page_ranges"][1:]):
        assert b["from"] == a["to"] + 1
    sizes = {x["code"]: x["to"] - x["from"] + 1 for x in v["page_ranges"]}
    for code, names in ((1, ["委托代理合同"]), (2, ["律师费发票"]), (3, ["授权委托书"]),
                        (7, ["聊天记录", "送货单", "对账单"]), (14, ["民事调解书"])):
        assert sizes[code] == sum(pdf_pages(mats[n]) for n in names), code
    assert sizes[4] >= 1 and sizes[15] == 1                             # 起诉状转出来的页数、结案报告一页
    for i, page in enumerate(vol.pages, 1):
        assert f"第{i}页 共{total}页" in page.extract_text(), i
    # 立卷申请书的页码范围与 page_ranges 一致；民事行政卷保留原编号
    rows = app_rows(kinds["立卷申请书"])
    assert [r[0] for r in rows] == [str(c) for c in expect_codes]
    for (no, text), rng in zip(rows, v["page_ranges"]):
        want = f"p{rng['from']}" if rng["from"] == rng["to"] else f"p{rng['from']}-{rng['to']}"
        assert text.endswith(" " + want), (no, text)
    assert rows[-1][1].startswith("结案报告")
    # 发票.pdf：收费发票凭证一项单独一份
    assert pdf_pages(kinds["发票"]) == pdf_pages(mats["律师费发票"])
    # 结案报告 docx 内容
    rep = "\n".join(x.text for x in pydocx.Document(str(kinds["结案报告"])).paragraphs)
    assert "赵某丁" in rep and "钱某戊" in rep and "孙律师" in rep and "八、办案结果：调解" in rep
    # 归档目录.md、提示
    md = kinds["归档目录"].read_text(encoding="utf-8")
    assert "| 7 | 证据材料 | * | 聊天记录、送货单、对账单 |" in md and "缺失必交 0 项" in md
    assert B.SIGN in v["manual"] and any("临时模板" in x for x in v["manual"])
    # 成果/索引.json 登记（只登记 md / docx）
    idx = json.loads((env.root / "成果" / "索引.json").read_text(encoding="utf-8"))
    assert not list(validator("files/outputs_index.schema.json").iter_errors(idx))
    e = idx["outputs"][-1]
    assert e["title"] == "赵某丁与钱某戊案件归档" and e["version"] == 1
    assert {f["path"].rsplit("/", 1)[1] for f in e["files"]} == {"立卷申请书.docx", "结案报告.docx", "归档目录.md"}
    # 原件不变；工作区/临时 用完为空
    assert originals(env) == before
    assert list((env.root / "工作区" / "临时").iterdir()) == []
    assert folder.is_dir()


@needs_lo
def test_build_again_v2_and_statements(env, tid):
    p = plan(items=[it for it in plan()["items"] if it["code"] not in (3, 14)], fee_settled=False)
    save_plan(env, tid, p)
    a = ok(build(env, tid, p), "api/archive_build.schema.json")
    b = ok(build(env, tid, p), "api/archive_build.schema.json")
    assert a["folder"].endswith("案件归档") and b["folder"].endswith("案件归档-v2")
    kinds = [f["kind"] for f in b["files"]]
    assert kinds.count("特殊情况说明") == 2
    paths = {pathlib.PurePosixPath(f["path"]).name: env.root / f["path"] for f in b["files"]}
    miss = "\n".join(x.text for x in pydocx.Document(str(paths["材料缺失情况说明及承诺.docx"])).paragraphs)
    assert "1. 缺失材料：授权委托书。原因：【待填写】" in miss and "2. 缺失材料：判决书、裁定书、调解书、上诉书。" in miss
    assert "（（2026）粤连越深圳民字第0001号+赵某丁）" in miss and "已办结（办案结果：调解）" in miss
    fee = "\n".join(x.text for x in pydocx.Document(str(paths["律师费未结清情况说明及承诺.docx"])).paragraphs)
    assert "关于合同律师费未结清情况说明及承诺" in fee and "（2026）粤连越深圳民字第0001号+赵某丁" in fee
    assert (env.root / a["folder"] / "卷宗.pdf").is_file()                # 旧的不动
    idx = json.loads((env.root / "成果" / "索引.json").read_text(encoding="utf-8"))
    assert [o["version"] for o in idx["outputs"]] == [1, 2]


@needs_lo
def test_image_material_and_no_opponent(env, tid):
    p = plan(opponent=None, items=[{"code": 7, "name": "证据材料", "materials": ["微信图片_20260301", "送货单"]}])
    save_plan(env, tid, p)
    v = ok(build(env, tid, p), "api/archive_build.schema.json")
    assert v["folder"] == "成果/归档/赵某丁案件归档"
    rng = {x["code"]: x for x in v["page_ranges"]}
    assert rng[7]["to"] - rng[7]["from"] + 1 == 1 + pdf_pages(CLOSED / "03一审" / "我方证据" / "送货单.pdf")
    assert any("发票.pdf 未生成" in x for x in v["manual"])


def test_result_null_rejected(env, tid):
    p = plan(result=None)
    save_plan(env, tid, p)
    fail(build(env, tid, p), "PLAN_NOT_CONFIRMED")
    assert not (env.root / "成果" / "归档").exists()


def test_build_bad_plan_path(env, tid):
    p = plan()
    fail(build(env, tid, p), "INVALID_ARGUMENT")                            # 还没保存过方案
    save_plan(env, tid, p)
    fail(build(env, tid, p, plan_rel=f"工作区/任务/{tid}/task.json"), "INVALID_ARGUMENT")
    fail(build(env, tid, plan(items=[{"code": 15, "name": "x", "materials": ["委托代理合同"]}])), "INVALID_ARGUMENT")


def test_converter_unavailable(env, tid, monkeypatch, tmp_path):
    """三种转换程序都不可用（模拟）：CONVERTER_UNAVAILABLE；不留归档文件夹，工作区/临时 为空，索引不变。"""
    s = ok(env.client.get("/api/settings"), "api/settings.schema.json")
    s["converter"] = "auto"
    ok(env.client.put("/api/settings", json=s), "api/settings.schema.json")
    log = tmp_path / "com.log"
    log.write_text("", encoding="utf-8")
    monkeypatch.setattr(C, "WORKER_CMD", [sys.executable, str(pathlib.Path(__file__).with_name("fake_com_worker.py"))])
    monkeypatch.setenv("FAKE_COM_LOG", str(log))
    monkeypatch.setenv("FAKE_COM_WORD", "fail")
    st = env.client.app.state.lb
    monkeypatch.setattr(st.archive.converter, "soffice", str(tmp_path / "没有" / "soffice.exe"))
    p = plan()
    save_plan(env, tid, p)
    fail(build(env, tid, p), "CONVERTER_UNAVAILABLE")
    assert not (env.root / "成果" / "归档").exists() and not (env.root / "成果" / "索引.json").exists()
    assert list((env.root / "工作区" / "临时").iterdir()) == []
    assert "Word.Application" in log.read_text(encoding="utf-8")


# ---------------------------------------------------------------- 律所模板在的时候

def report_template(path: pathlib.Path) -> None:
    """仿原版结案报告模板：占位符拆在几个 run 里。"""
    d = pydocx.Document()
    d.add_paragraph("结案报告")
    for parts in (["一、当事人/委托人：", "【", "        ", "】"], ["对方当事人：【  】"], ["委托日期：【YYYY-", "MM-DD】"],
                  ["结案日期：【YYYY-MM-DD】"], ["承办律师：【 】"], ["案由：【案由全称】"],
                  ["案情简介：【委托人】因【案由】等事宜向我方委托人申请【仲裁/起诉/应诉】。"],
                  ["承办律师分析与意见：委托人有权【对/向】【对方当事人】主张【诉求】。"], ["八、办案结果：【胜诉/调解】"],
                  ["备注：【不知道是什么】"]):
        p = d.add_paragraph()
        for t in parts:
            p.add_run(t)
    d.save(str(path))


def application_template(path: pathlib.Path, catalog: dict) -> None:
    """仿原版立卷申请书模板：P5 占位符拆成 5 个 run；表格 1 行表头 + 每项一行（编码 | 材料名称），名称加粗。"""
    d = pydocx.Document()
    d.add_paragraph("立卷申请书")
    p = d.add_paragraph()
    for t in ("【", "（", "金助理系统", "案件编号+客户名称）", "】"):
        p.add_run(t)
    p.add_run("符合归档条件……")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "编码", "材料名称"
    for it in catalog["items"]:
        r = t.add_row()
        r.cells[0].text = str(it["code"])
        r.cells[1].paragraphs[0].add_run(it["name"]).bold = True
    d.save(str(path))


@needs_lo
def test_templates_used(env, tid, tmp_path):
    skills = tmp_path / "skills"
    shutil.copytree(SKILLS / "case-archiving", skills / "case-archiving")
    cat = json.loads((skills / "case-archiving" / "catalogs" / "民事行政卷.json").read_text(encoding="utf-8"))
    (skills / "case-archiving" / "templates").mkdir()
    report_template(skills / "case-archiving" / "templates" / "结案报告模板.docx")
    application_template(skills / "case-archiving" / "templates" / "立卷申请书模板.docx", cat)
    env.client.app.state.lb.archive.skills_dirs = [skills]
    p = plan()
    save_plan(env, tid, p)
    v = ok(build(env, tid, p), "api/archive_build.schema.json")
    assert not any("临时" in x for x in v["manual"])
    assert any("1 处认不出" in x for x in v["manual"])                   # "备注：【不知道是什么】"
    paths = {f["kind"]: env.root / f["path"] for f in v["files"]}
    rep = [x.text for x in pydocx.Document(str(paths["结案报告"])).paragraphs]
    assert rep[1] == "一、当事人/委托人：赵某丁" and rep[2] == "对方当事人：钱某戊"
    assert rep[3] == "委托日期：2026-01-05" and rep[5] == "承办律师：孙律师" and rep[6] == "案由：买卖合同纠纷"
    assert rep[7] == "案情简介：" + p["summary"] and rep[8] == "承办律师分析与意见：" + p["opinion"]
    assert rep[9] == "八、办案结果：调解" and rep[10] == "备注：【不知道是什么】"
    doc = pydocx.Document(str(paths["立卷申请书"]))
    assert doc.paragraphs[1].text == "（（2026）粤连越深圳民字第0001号+赵某丁）符合归档条件……"
    rows = app_rows(paths["立卷申请书"])
    assert [r[0] for r in rows] == ["1", "2", "3", "4", "7", "14", "15"]
    rng = {x["code"]: x for x in v["page_ranges"]}
    assert rows[4][1] == f"证据材料 p{rng[7]['from']}-{rng[7]['to']}"
    for row in doc.tables[0].rows[1:]:
        for run in row.cells[1].paragraphs[0].runs:
            assert run.bold is False and run.font.size.pt == 12
            assert run._element.rPr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia") == "仿宋"


def test_application_rows_renumber():
    cat = json.loads((SKILLS / "case-archiving" / "catalogs" / "常法卷.json").read_text(encoding="utf-8"))
    gen = next(it for it in cat["items"] if it.get("generated"))
    first, third = cat["items"][0], cat["items"][2]
    p = plan(catalog="常法卷", items=[{"code": first["code"], "name": "常年法律顾问合同", "materials": ["委托代理合同"]},
                                     {"code": third["code"], "name": "服务记录", "materials": ["对账单"]}])
    rows = D.application_rows(p, cat, {first["code"]: (1, 2), third["code"]: (3, 3), gen["code"]: (4, 4)})
    assert [no for _c, no, _t in rows] == ["1", "2", "3"]                 # 常法卷重新连续编号
    assert rows[0][2] == "常年法律顾问合同 p1-2" and rows[1][2] == "服务记录 p3"


# ---------------------------------------------------------------- 页码层

def test_number_pages_rotated_and_offset():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    for i in range(3):
        c.drawString(100, 20, f"原页码 {i + 9}")
        c.showPage()
    c.save()
    w = PdfWriter()
    w.append(PdfReader(io.BytesIO(buf.getvalue())))
    w.pages[1].rotate(90)
    w.pages[2].mediabox.lower_left = (50, 50)
    B.number_pages(w)
    out = io.BytesIO()
    w.write(out)
    r = PdfReader(io.BytesIO(out.getvalue()))
    for i, page in enumerate(r.pages, 1):
        assert f"第{i}页 共3页" in page.extract_text()
    assert r.pages[1].rotation == 0                                        # 旋转已转进内容，页码在纸面底部


def test_folder_name():
    assert B.folder_name("赵某丁", "钱某戊") == "赵某丁与钱某戊案件归档"
    assert B.folder_name("某/公司:分部", None) == "某_公司_分部案件归档"
    assert B.folder_name("长" * 60, "短") == "长" * 40 + "与短案件归档"


# ---------------------------------------------------------------- 第一轮复核返修（0054 令）

class _Counter(__import__("http.server").server.BaseHTTPRequestHandler):
    hits: list = []

    def do_GET(self):  # noqa: N802
        type(self).hits.append(self.path)
        self.send_response(404)
        self.end_headers()

    def log_message(self, *a):
        pass


@needs_lo
def test_import_rejected_material_not_converted(tmp_path, monkeypatch):
    """P1-1 端到端：带外链图片的旧版 .doc 改名 .docx 放进案件 → 导入 failed → 放进方案（只给提醒）→ 生成：
    这份材料跳过、写进提示；监听 0 请求；假 Word 只转了起诉状和结案报告。"""
    import http.server
    import threading
    from test_convert import linked_picture_docx
    _Counter.hits = []
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Counter)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{srv.server_address[1]}"
    try:
        src = tmp_path / "linked.docx"
        linked_picture_docx(src, url)
        lo_base = tmp_path / "lo"
        import tempfile
        with tempfile.TemporaryDirectory(prefix="lblo-") as lb, lo.Converter(tmp_path / "mk", lo_base=lb).session() as s:
            doc_bytes = s.convert(src, "doc").read_bytes()
        _Counter.hits.clear()
        evil = tmp_path / "src" / "对方证据甲.docx"
        evil.parent.mkdir()
        evil.write_bytes(doc_bytes)
        files = dict(FILES)
        files["03一审/我方证据/对方证据甲.docx"] = evil
        e = Env(tmp_path / "case", files)
        try:
            m = next(x for x in json.loads((e.root / "工作区" / "材料" / "index.json").read_text(encoding="utf-8"))
                     ["materials"] if x["name"] == "对方证据甲")
            assert m["status"] == "failed"
            log = tmp_path / "com.log"
            log.write_text("", encoding="utf-8")
            monkeypatch.setattr(C, "WORKER_CMD", [sys.executable, str(pathlib.Path(__file__).with_name("fake_com_worker.py"))])
            monkeypatch.setenv("FAKE_COM_LOG", str(log))
            monkeypatch.setenv("FAKE_COM_WORD", "ok")
            s = ok(e.client.get("/api/settings"), "api/settings.schema.json")
            s["converter"] = "auto"
            ok(e.client.put("/api/settings", json=s), "api/settings.schema.json")
            tid = e.begin("sess-evil")["task_id"]
            p = plan(items=[it if it["code"] != 7 else {"code": 7, "name": "证据材料",
                                                         "materials": ["对方证据甲", "送货单"]} for it in plan()["items"]])
            saved = e.tool_ok(tid, "case_save_archive_plan", p)
            assert any("对方证据甲" in w for w in saved["warnings"])
            v = ok(build(e, tid, p), "api/archive_build.schema.json")
            assert any("对方证据甲" in x for x in v["manual"])
            rng = {x["code"]: x for x in v["page_ranges"]}
            assert rng[7]["to"] - rng[7]["from"] + 1 == pdf_pages(CLOSED / "03一审" / "我方证据" / "送货单.pdf")
            assert log.read_text(encoding="utf-8").split() == ["Word.Application"] * 2   # 起诉状、结案报告；对方证据甲没碰
            assert _Counter.hits == []
        finally:
            e.close()
    finally:
        srv.shutdown()
        srv.server_close()


def test_md_material_rejected_in_plan(tmp_path):
    """P3-2：.md 材料转不了 PDF：保存方案时就拒，不等生成到一半才失败。"""
    md = tmp_path / "备忘.md"
    md.write_text("# 备忘\n内容", encoding="utf-8")
    files = dict(FILES)
    files["03一审/我方证据/备忘.md"] = md
    e = Env(tmp_path / "case", files)
    try:
        tid = e.begin("sess-md")["task_id"]
        p = plan(items=[{"code": 7, "name": "证据材料", "materials": ["备忘"]}])
        fail(e.tool(tid, "case_save_archive_plan", p), "INVALID_ARGUMENT")
    finally:
        e.close()


@needs_lo
def test_item_with_no_readable_material_counts_missing(env, tid, monkeypatch):
    """P3-4：某项的材料全都跳过：这一项不进立卷申请书、归档目录写缺失；必交项计入缺失并生成情况说明。"""
    real = B._readable
    monkeypatch.setattr(B, "_readable", lambda p: None if p.name == "授权委托书.pdf" else real(p))
    p = plan()
    save_plan(env, tid, p)
    v = ok(build(env, tid, p), "api/archive_build.schema.json")
    assert 3 not in [x["code"] for x in v["page_ranges"]]
    paths = {pathlib.PurePosixPath(f["path"]).name: env.root / f["path"] for f in v["files"]}
    assert "3" not in [r[0] for r in app_rows(paths["立卷申请书.docx"])]
    miss = "\n".join(x.text for x in pydocx.Document(str(paths["材料缺失情况说明及承诺.docx"])).paragraphs)
    assert "缺失材料：授权委托书" in miss
    md = paths["归档目录.md"].read_text(encoding="utf-8")
    assert "| 3 | 授权委托书 | * | — |" in md and "未能放入卷宗的材料：授权委托书" in md


def test_page_numbers_inside_cropbox():
    """P3-1：CropBox 比纸面小（扫描件常见）时，页码画在可见区域里。"""
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(595, 842))
    c.drawString(100, 400, "正文")
    c.showPage()
    c.save()
    w = PdfWriter()
    w.append(PdfReader(io.BytesIO(buf.getvalue())))
    w.pages[0].cropbox.lower_left = (40, 100)
    w.pages[0].cropbox.upper_right = (560, 800)
    B.number_pages(w)
    out = io.BytesIO()
    w.write(out)
    page = PdfReader(io.BytesIO(out.getvalue())).pages[0]
    ys = []
    page.extract_text(visitor_text=lambda text, cm, tm, fd, fs: ys.append((text, tm[5] * cm[3] + cm[5])) if "页" in text else None)
    assert ys and all(100 <= y <= 800 for _t, y in ys), ys
