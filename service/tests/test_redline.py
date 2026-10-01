"""修订版 Word（Spec 12.2；工单 T15 第 3–5 步）。

contract-01 上 14 条修改：替换、插入、删除各 3 条（范围内），另有表格内、页眉页脚、跨超链接、找不到、出现多次各 1 条
（范围外）。范围内的全部生成修订和批注；范围外的 5 条全部进"需人工修改"；含修订的副本整份拒绝。
生成的 docx：全部接受 = 改后的文字，全部拒绝 = 原文（逐段核对），格式属性跟着原 run 走；LibreOffice 能打开并转 PDF。
"""
from __future__ import annotations

import hashlib
import io
import json
import pathlib
import zipfile

import docx as pydocx
import pytest
from lxml import etree

from lawbench.export import redline as R
from lawbench.ingest import docx as dx
from lawbench.ingest import libreoffice
from lawbench.tools import edit_list as el

from conftest import short_dir
from t8_helpers import FIXTURES, Env, fail, ok, validator

W = dx.W
CONTRACT = FIXTURES / "contract-01" / "采购合同.docx"
FILES = {"合同/采购合同.docx": CONTRACT,
         "合同/采购合同-含未处理修订.docx": FIXTURES / "contract-01" / "采购合同-含未处理修订.docx",
         "起诉意见书.pdf": FIXTURES / "criminal-01" / "起诉意见书.pdf"}

IN_SCOPE = [
    {"para": 27, "action": "replace", "find": "九十日", "text": "三十日"},
    {"para": 23, "action": "replace", "find": "七日", "text": "十五日"},
    {"para": 31, "action": "replace", "find": "万分之五", "text": "万分之三"},
    {"para": 21, "action": "insert_after", "find": "通知甲方收货", "text": "（书面通知）"},
    {"para": 17, "action": "insert_after", "find": "甲方有权拒收该批货物", "text": "，并有权解除合同"},
    {"para": 42, "action": "insert_after", "find": "签字盖章", "text": "（法定代表人或授权代表签字）"},
    {"para": 24, "action": "delete", "find": "隐蔽瑕疵的"},
    {"para": 32, "action": "delete", "find": "还应"},
    {"para": 16, "action": "delete", "find": "，检测费用先由青禾建材垫付"},
]
OUT_OF_SCOPE = [
    ({"para": 12, "action": "replace", "find": "HW200", "text": "HW300"}, "表格"),
    ({"para": 47, "action": "delete", "find": "虚构"}, "页眉页脚"),
    ({"para": 39, "action": "replace", "find": "电子版可在http", "text": "电子版可在"}, "超链接"),
    ({"para": 28, "action": "replace", "find": "不存在的原文", "text": "x"}, "找不到"),
    ({"para": 30, "action": "replace", "find": "逾期", "text": "迟延"}, "多次"),
]
EXPECTED = {   # 全部接受后的段落文字
    16: "2.2 青禾建材有权委托第三方检测机构对标的物进行抽检。",
    17: "2.3 抽检不合格的，检测费用由乙方承担，甲方有权拒收该批货物，并有权解除合同。",
    21: "3.3 乙方应提前两日通知甲方收货（书面通知），甲方应安排人员清点签收。",
    23: "4.1 甲方应于到货后十五日内完成验收，逾期未提出书面异议的，视为验收合格。",
    24: "4.2 异议期为验收合格之日起六个月。",
    27: "5.2 甲方应于收到货物后三十日内付款。",
    31: "6.2 甲方逾期付款的，每逾期一日，按逾期付款金额的万分之三向乙方支付违约金。",
    32: "6.3 违约金不足以弥补损失的，违约方赔偿损失。",
    42: "10.1 本合同自双方签字盖章（法定代表人或授权代表签字）之日起生效。",
}


def all_edits() -> list[dict]:
    items = IN_SCOPE + [e for e, _ in OUT_OF_SCOPE]
    return [dict(id=i, comment=f"第{i}条的理由", **e) for i, e in enumerate(items, 1)]


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    e = Env(tmp_path_factory.mktemp("t15red"), FILES)
    yield e
    e.close()


@pytest.fixture
def tid(env, request):
    return env.begin(f"sess-{request.node.name}")["task_id"]


def save_list(env, tid, name: str, edits: list[dict]) -> str:
    return env.tool_ok(tid, "case_save_edit_list", {"name": name, "edits": edits})["path"]


def redline(env, tid, edit_list: str):
    return env.client.post("/api/redline", json={"case_id": env.case_id, "task_id": tid, "edit_list": edit_list})


# ---------------------------------------------------------------- 读生成的 docx

def view(p, mode: str) -> str:
    """段落文字：accept = 跳过 w:del、算 w:ins；reject = 跳过 w:ins、w:delText 当原文。"""
    out: list[str] = []

    def walk(node):
        for ch in node:
            tag = ch.tag
            if not isinstance(tag, str):
                continue
            if tag == f"{{{W}}}ins" and mode == "reject" or tag == f"{{{W}}}del" and mode == "accept":
                continue
            if tag in (f"{{{W}}}t",) or tag == f"{{{W}}}delText" and mode == "reject":
                out.append(ch.text or "")
            elif tag == f"{{{W}}}tab":
                out.append("\t")
            elif tag in (f"{{{W}}}br", f"{{{W}}}cr"):
                out.append(" ")
            elif tag not in (f"{{{W}}}instrText", f"{{{W}}}delInstrText"):
                walk(ch)

    walk(p)
    return "".join(out)


def body_paras(data: bytes) -> list:
    root = dx._parse_xml(zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml"))
    body = root.find("w:body", dx.NS)
    return [e for k, e in dx._body_items(body) if (dx._text(e).strip() if k == "p" else dx._table(e))]


# ---------------------------------------------------------------- 接口：contract-01 的 14 条

def test_contract01_fourteen_edits(env, tid):
    edits = all_edits()
    saved = env.tool_ok(tid, "case_save_edit_list", {"name": "采购合同", "edits": edits})
    assert saved["accepted"] == 9 and len(saved["out_of_scope"]) == 5
    v = ok(redline(env, tid, saved["path"]), "api/redline.schema.json")
    assert v["path"] == f"工作区/任务/{tid}/草稿/采购合同-修订版-v1.docx"
    assert v["applied"] == 9
    manual = {m["id"]: m["reason"] for m in v["manual"]}
    assert sorted(manual) == [10, 11, 12, 13, 14]
    for i, (_e, word) in enumerate(OUT_OF_SCOPE, 10):
        assert word in manual[i], (i, manual[i])
    assert manual == {o["id"]: o["reason"] for o in saved["out_of_scope"]}   # 与保存修改清单时说的一致

    data = (env.root / v["path"]).read_bytes()
    orig = body_paras(CONTRACT.read_bytes())
    new = body_paras(data)
    assert len(new) == len(orig) == 46
    for n, (a, b) in enumerate(zip(orig, new), 1):
        assert view(b, "reject") == view(a, "reject"), n                  # 全部拒绝 = 原文，逐段
        assert view(b, "accept") == EXPECTED.get(n, view(a, "accept")), n  # 全部接受 = 改后

    z = zipfile.ZipFile(io.BytesIO(data))
    doc = dx._parse_xml(z.read("word/document.xml"))
    ins = doc.findall(f".//{{{W}}}ins")
    dels = doc.findall(f".//{{{W}}}del")
    assert len(ins) == 6 and len(dels) == 6                               # 替换 3（各一删一插）+ 插入 3 + 删除 3
    assert {x.get(f"{{{W}}}author") for x in ins + dels} == {R.AUTHOR}
    ids = [x.get(f"{{{W}}}id") for x in ins + dels]
    assert len(set(ids)) == len(ids)
    assert all(t.tag == f"{{{W}}}delText" for d in dels for t in d.iter() if t.tag in
               (f"{{{W}}}t", f"{{{W}}}delText"))
    comments = dx._parse_xml(z.read("word/comments.xml"))
    texts = ["".join(c.itertext()) for c in comments]
    assert texts == [f"第{i}条的理由" for i in range(1, 10)]
    cids = {c.get(f"{{{W}}}id") for c in comments}
    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
        assert {x.get(f"{{{W}}}id") for x in doc.iter(f"{{{W}}}{tag}")} == cids
    rels = z.read("word/_rels/document.xml.rels").decode()
    assert R.COMMENTS_REL in rels and 'Target="comments.xml"' in rels
    assert "/word/comments.xml" in z.read("[Content_Types].xml").decode()
    # 页眉页脚、表格等未改的部件逐字节不变
    zo = zipfile.ZipFile(CONTRACT)
    for n in zo.namelist():
        if n not in ("word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"):
            assert z.read(n) == zo.read(n), n
    # 修订版登进 result.json 的草稿列表
    res = env.read_json(tid, "result.json", "files/result.schema.json")
    assert {"title": "采购合同-修订版", "path": v["path"], "version": 1} in res["drafts"]


def test_original_untouched_and_version_increments(env, tid):
    before = hashlib.sha256((env.root / "合同" / "采购合同.docx").read_bytes()).hexdigest()
    p = save_list(env, tid, "采购合同", [dict(id=1, comment="理由", **IN_SCOPE[0])])
    a = ok(redline(env, tid, p), "api/redline.schema.json")
    b = ok(redline(env, tid, p), "api/redline.schema.json")
    assert a["path"].endswith("采购合同-修订版-v1.docx") and b["path"].endswith("采购合同-修订版-v2.docx")
    assert (env.root / a["path"]).is_file() and (env.root / b["path"]).is_file()
    assert hashlib.sha256((env.root / "合同" / "采购合同.docx").read_bytes()).hexdigest() == before


def test_revised_original_rejected_whole(env, tid):
    p = save_list(env, tid, "采购合同-含未处理修订", [dict(id=1, comment="理由", **IN_SCOPE[0])])
    fail(redline(env, tid, p), "INVALID_ARGUMENT")
    assert not (env.task_dir(tid) / "草稿").exists() or \
        not list((env.task_dir(tid) / "草稿").glob("*修订版*"))
    with pytest.raises(R.Revised):
        R.generate((FIXTURES / "contract-01" / "采购合同-含未处理修订.docx").read_bytes(),
                   [dict(id=1, comment="x", **IN_SCOPE[0])])


def test_original_changed_after_scan(env, tid, tmp_path):
    p = save_list(env, tid, "采购合同", [dict(id=1, comment="理由", **IN_SCOPE[0])])
    src = env.root / "合同" / "采购合同.docx"
    keep = src.read_bytes()
    try:
        d = pydocx.Document(io.BytesIO(keep))
        d.paragraphs[0].insert_paragraph_before("新加的一段")              # 段号整体后移
        d.save(str(src))
        fail(redline(env, tid, p), "INPUT_CHANGED")
    finally:
        src.write_bytes(keep)


@pytest.mark.parametrize("rel", [
    "工作区/任务/{tid}/修改清单/../task.json", "工作区/任务/{tid}/草稿/x.json", "合同/采购合同.docx",
    "工作区/任务/T-20260101000000-0000/修改清单/采购合同.json", "工作区/任务/{tid}/修改清单/没有.json",
])
def test_bad_edit_list_path(env, tid, rel):
    r = redline(env, tid, rel.format(tid=tid))
    assert r.json()["ok"] is False and r.json()["error"]["code"] in ("INVALID_ARGUMENT", "OUT_OF_CASE")


def test_not_docx_and_other_case(env, tid):
    lst = env.task_dir(tid) / "修改清单"
    lst.mkdir(parents=True, exist_ok=True)
    (lst / "起诉意见书.json").write_text(json.dumps({"name": "起诉意见书", "edits": [
        {"id": 1, "para": 1, "action": "delete", "find": "x", "comment": "y"}]}, ensure_ascii=False), encoding="utf-8")
    fail(redline(env, tid, f"工作区/任务/{tid}/修改清单/起诉意见书.json"), "INVALID_ARGUMENT")
    (lst / "坏.json").write_text('{"name": "采购合同", "edits": []}', encoding="utf-8")
    fail(redline(env, tid, f"工作区/任务/{tid}/修改清单/坏.json"), "INVALID_ARGUMENT")
    r = env.client.post("/api/redline", json={"case_id": "00000000-0000-4000-8000-000000000000", "task_id": tid,
                                              "edit_list": f"工作区/任务/{tid}/修改清单/坏.json"})
    assert r.json()["ok"] is False


# ---------------------------------------------------------------- 库：格式、拆 run、重叠、已有批注

def make_docx(paras: list[list[tuple[str, dict]]], comments: bool = False) -> bytes:
    """每段若干 run：(文字, {"bold": True, "italic": True, "size": 14})。"""
    d = pydocx.Document()
    for runs in paras:
        p = d.add_paragraph()
        for text, fmt in runs:
            r = p.add_run(text)
            r.bold = fmt.get("bold")
            r.italic = fmt.get("italic")
            if fmt.get("size"):
                r.font.size = pydocx.shared.Pt(fmt["size"])
    if comments:
        d.add_comment(d.paragraphs[0].runs[0], text="律师原有批注", author="张律师")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def run_fmt(r) -> tuple:
    rpr = r.find(f"{{{W}}}rPr")
    return (rpr is not None and rpr.find(f"{{{W}}}b") is not None,
            rpr is not None and rpr.find(f"{{{W}}}i") is not None)


def test_split_runs_keep_formatting():
    """find 跨三个不同格式的 run、两头都在 run 中间：拆开后每一截都带原来的格式；新文字取相邻原 run 的格式。"""
    data = make_docx([[("甲方应于", {}), ("收到货物后九十日", {"bold": True}), ("内付款。", {"italic": True})]])
    out, applied, manual = R.generate(data, [
        {"id": 1, "para": 1, "action": "replace", "find": "于收到货物后九十日内", "text": "验收后三十日内",
         "comment": "c"}])
    assert applied == [1] and manual == []
    p = body_paras(out)[0]
    assert view(p, "accept") == "甲方应验收后三十日内付款。" and view(p, "reject") == "甲方应于收到货物后九十日内付款。"
    d = p.find(f"{{{W}}}del")
    deleted = [("".join(t.text for t in r.iter(f"{{{W}}}delText")), run_fmt(r)) for r in d.iter(f"{{{W}}}r")]
    assert deleted == [("于", (False, False)), ("收到货物后九十日", (True, False)), ("内", (False, True))]
    kept = [("".join(t.text for t in r.iter(f"{{{W}}}t")), run_fmt(r)) for r in p if r.tag == f"{{{W}}}r"
            and r.find(f"{{{W}}}t") is not None]
    assert kept == [("甲方应", (False, False)), ("付款。", (False, True))]
    new = p.find(f"{{{W}}}ins").find(f"{{{W}}}r")
    assert run_fmt(new) == (False, True)                                   # 取最后一个被删 run 的格式


def test_overlap_and_same_paragraph():
    data = make_docx([[("甲方应于收到货物后九十日内付款，乙方应开具发票。", {})]])
    out, applied, manual = R.generate(data, [
        {"id": 1, "para": 1, "action": "replace", "find": "九十日", "text": "三十日", "comment": "a"},
        {"id": 2, "para": 1, "action": "delete", "find": "九十日内", "comment": "b"},           # 与 1 重叠
        {"id": 3, "para": 1, "action": "insert_after", "find": "开具发票", "text": "（增值税专用发票）", "comment": "c"},
        {"id": 4, "para": 1, "action": "delete", "find": "收到货物后", "comment": "d"},          # 紧挨着 1，不重叠
    ])
    assert applied == [1, 3, 4] and manual == [{"id": 2, "reason": R.R_OVERLAP}]
    p = body_paras(out)[0]
    assert view(p, "accept") == "甲方应于三十日内付款，乙方应开具发票（增值税专用发票）。"
    assert view(p, "reject") == "甲方应于收到货物后九十日内付款，乙方应开具发票。"


def test_existing_comments_appended():
    data = make_docx([[("甲方应于收到货物后九十日内付款。", {})]], comments=True)
    out, applied, _ = R.generate(data, [
        {"id": 1, "para": 1, "action": "replace", "find": "九十日", "text": "三十日", "comment": "新批注"}])
    z = zipfile.ZipFile(io.BytesIO(out))
    names = [n for n in z.namelist() if n.startswith("word/comments")]
    assert names == ["word/comments.xml"]
    comments = dx._parse_xml(z.read("word/comments.xml"))
    assert [c.get(f"{{{W}}}author") for c in comments] == ["张律师", R.AUTHOR]
    assert len({c.get(f"{{{W}}}id") for c in comments}) == 2
    assert z.read("word/_rels/document.xml.rels").decode().count(R.COMMENTS_REL) == 1


def test_multiline_insert_and_tab():
    data = make_docx([[("第一条\t付款方式：转账。", {})]])
    out, applied, _ = R.generate(data, [
        {"id": 1, "para": 1, "action": "insert_after", "find": "\t付款方式", "text": "及期限", "comment": "c"},
        {"id": 2, "para": 1, "action": "insert_after", "find": "转账。", "text": "\n另附：账户信息", "comment": "多行"}])
    assert applied == [1, 2]
    p = body_paras(out)[0]
    assert view(p, "accept") == "第一条\t付款方式及期限：转账。 另附：账户信息"


def test_hyperlink_inside_find_applied_cross_rejected():
    """find 整个在超链接文字里：可以改（修订包在超链接里面）；跨出超链接：进需人工修改。"""
    data = CONTRACT.read_bytes()
    p39 = body_paras(data)[38]
    link_text = "".join(t.text for t in p39.find(f".//{{{W}}}hyperlink").iter(f"{{{W}}}t"))
    out, applied, manual = R.generate(data, [
        {"id": 1, "para": 39, "action": "delete", "find": link_text[-5:], "comment": "c"},
        {"id": 2, "para": 39, "action": "delete", "find": "电子版可在" + link_text[:4], "comment": "d"}])
    assert applied == [1] and [m["id"] for m in manual] == [2]
    d = body_paras(out)[38].find(f".//{{{W}}}hyperlink/{{{W}}}del")
    assert d is not None


# ---------------------------------------------------------------- LibreOffice 能打开（结构有效）

@pytest.mark.skipif(libreoffice.find_soffice() is None, reason="本机没有 LibreOffice")
def test_libreoffice_opens_and_keeps_revisions(tmp_path):
    out, applied, _ = R.generate(CONTRACT.read_bytes(), all_edits())
    src = tmp_path / "修订版.docx"
    src.write_bytes(out)
    with short_dir("lblo-") as lb, libreoffice.Converter(tmp_path / "t", lo_base=lb).session() as s:
        pdf = s.convert(src, "pdf")
        assert pdf.read_bytes().startswith(b"%PDF") and pdf.stat().st_size > 1000
        back = s.convert(src, "docx")                                       # LibreOffice 读进再存：修订和批注还在
        z = zipfile.ZipFile(back)
        doc = z.read("word/document.xml").decode()
        assert doc.count("<w:ins ") >= 6 and doc.count("<w:del ") >= 6
        assert R.AUTHOR in doc
        assert "word/comments.xml" in z.namelist() and "第9条的理由" in z.read("word/comments.xml").decode()
