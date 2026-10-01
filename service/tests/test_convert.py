"""Word 转 PDF 的调度（Spec 12.3、14.3；工单 T23 第 1 步与两条补充）。

开发机装有 Word（没有 WPS）；用例默认不动真 Word：调度、超时、杀进程用假 COM 子进程（fake_com_worker.py）测，
LibreOffice 真跑。真 Word 实测另见 docs/plan/evidence/T23/printer.txt（LAWBENCH_REAL_WORD=1 时本文件最后一条也跑）。
"""
from __future__ import annotations

import http.server
import io
import os
import pathlib
import shutil
import sys
import threading
import time
import zipfile

import docx as pydocx
import pytest
from PIL import Image
from pypdf import PdfReader

from lawbench.case import gate
from lawbench.errors import ApiError
from lawbench.ingest import libreoffice as lo
from lawbench.office import convert as C

from conftest import short_dir

FIXTURES = pathlib.Path(__file__).resolve().parents[2] / "tests" / "fixtures"
COMPLAINT = FIXTURES / "closed-01" / "03一审" / "我方文件" / "民事起诉状.docx"
HAS_LO = lo.find_soffice() is not None
needs_lo = pytest.mark.skipif(not HAS_LO, reason="本机没有 LibreOffice")


@pytest.fixture
def case(tmp_path):
    root = tmp_path / "案件"
    (root / "工作区" / "临时").mkdir(parents=True)
    gate.mkdir_work(str(root), "工作区/临时/j")
    return root


@pytest.fixture
def lo_base():
    with short_dir("lblo-") as d:
        yield pathlib.Path(d)


@pytest.fixture
def fake(tmp_path, monkeypatch):
    log, pids = tmp_path / "com.log", tmp_path / "com.pids"
    log.write_text("", encoding="utf-8")
    pids.write_text("", encoding="utf-8")
    monkeypatch.setattr(C, "WORKER_CMD", [sys.executable, str(pathlib.Path(__file__).with_name("fake_com_worker.py"))])
    monkeypatch.setenv("FAKE_COM_LOG", str(log))
    monkeypatch.setenv("FAKE_COM_PIDS", str(pids))
    monkeypatch.setattr(C, "QUIT_WAIT", 0.5)

    class F:
        def calls(self):
            return log.read_text(encoding="utf-8").split()

        def pids(self):
            return [int(x) for x in pids.read_text(encoding="utf-8").split()]

        def set(self, **modes):
            for k, v in modes.items():
                monkeypatch.setenv(f"FAKE_COM_{k.upper()}", v)
    return F()


def pages(p: pathlib.Path) -> int:
    return len(PdfReader(str(p)).pages)


# ---------------------------------------------------------------- 调度

@needs_lo
def test_libreoffice_real(case, lo_base):
    out, used = C.OfficeConverter(lo_base).to_pdf(str(case), COMPLAINT, "工作区/临时/j", "libreoffice")
    assert used == "libreoffice" and pages(out) >= 1
    assert out.parent.parent == case / "工作区" / "临时" / "j"           # 只在本次的临时目录里
    assert sorted(p.name for p in out.parent.iterdir()) == ["in.docx", "out.pdf"]


def test_word_ok_is_first(case, lo_base, fake):
    fake.set(word="ok")
    out, used = C.OfficeConverter(lo_base).to_pdf(str(case), COMPLAINT, "工作区/临时/j")
    assert used == "word" and fake.calls() == ["Word.Application"] and pages(out) == 1


def test_word_missing_then_wps(case, lo_base, fake):
    fake.set(word="none", kwps="none", wps="ok")
    out, used = C.OfficeConverter(lo_base).to_pdf(str(case), COMPLAINT, "工作区/临时/j")
    assert used == "wps" and fake.calls() == ["Word.Application", "KWPS.Application", "WPS.Application"]


@needs_lo
def test_hang_times_out_kills_and_falls_through(case, lo_base, fake, monkeypatch):
    """Word 卡住不返回（如"等待打印机连接"）：超时后子进程和它起的进程都已结束，换 WPS，再换 LibreOffice。"""
    monkeypatch.setattr(C, "TIMEOUT", 3)
    fake.set(word="hang", kwps="none", wps="none")
    t0 = time.monotonic()
    out, used = C.OfficeConverter(lo_base).to_pdf(str(case), COMPLAINT, "工作区/临时/j")
    assert used == "libreoffice" and pages(out) >= 1
    assert time.monotonic() - t0 < 60
    worker, child = fake.pids()
    alive = C.processes()
    assert worker not in alive and child not in alive


def test_all_three_fail(case, lo_base, fake):
    fake.set(word="fail", kwps="none", wps="none")
    conv = C.OfficeConverter(lo_base, soffice=str(case / "没有" / "soffice.exe"))
    with pytest.raises(ApiError) as e:
        conv.to_pdf(str(case), COMPLAINT, "工作区/临时/j")
    assert e.value.code == "CONVERTER_UNAVAILABLE"


def test_explicit_choice_only_that_one(case, lo_base, fake):
    """设置里选了 libreoffice：不碰 Word / WPS（律所为避免留痕改的设置不能被悄悄绕过）；选了 word 而 Word 失败：不换。"""
    fake.set(word="ok")
    conv = C.OfficeConverter(lo_base, soffice=str(case / "没有" / "soffice.exe"))
    with pytest.raises(ApiError):
        conv.to_pdf(str(case), COMPLAINT, "工作区/临时/j", "libreoffice")
    assert fake.calls() == []
    fake.set(word="fail", kwps="ok")
    with pytest.raises(ApiError):
        conv.to_pdf(str(case), COMPLAINT, "工作区/临时/j", "word")
    assert fake.calls() == ["Word.Application"]


def test_spreadsheet_only_libreoffice(case, lo_base, fake, tmp_path):
    fake.set(word="ok")
    x = tmp_path / "表.xlsx"
    shutil.copy(FIXTURES / "civil-01" / "银行流水.xlsx", x)
    conv = C.OfficeConverter(lo_base, soffice=str(case / "没有" / "soffice.exe"))
    with pytest.raises(ApiError):
        conv.to_pdf(str(case), x, "工作区/临时/j")
    assert fake.calls() == []                                          # Word 不打开表格


def test_doc_with_external_picture_not_given_to_libreoffice(case, lo_base, fake, monkeypatch, tmp_path):
    monkeypatch.setattr(C.links, "has_external_picture", lambda p: True)
    started = []
    monkeypatch.setattr(C.lo.Converter, "session", lambda self: started.append(1))
    d = tmp_path / "旧.doc"
    d.write_bytes(b"\xd0\xcf\x11\xe0" + b"\0" * 100)
    with pytest.raises(ApiError):
        C.OfficeConverter(lo_base).to_pdf(str(case), d, "工作区/临时/j", "libreoffice")
    assert started == []


def test_com_started_only_ours():
    before = {1: (0, "system"), 10: (0, "svchost.exe"), 20: (0, "explorer.exe"), 30: (20, "winword.exe")}
    after = dict(before)
    after.update({40: (20, "winword.exe"),       # 律师刚双击打开的 Word（父进程 explorer）
                  50: (10, "winword.exe"),       # 本次 COM 启动的
                  60: (10, "excel.exe")})
    assert C.com_started(before, after, ("winword.exe",)) == [50]


def test_original_untouched(case, lo_base, fake):
    fake.set(word="ok")
    import hashlib
    h = hashlib.sha256(COMPLAINT.read_bytes()).hexdigest()
    C.OfficeConverter(lo_base).to_pdf(str(case), COMPLAINT, "工作区/临时/j")
    assert hashlib.sha256(COMPLAINT.read_bytes()).hexdigest() == h


# ---------------------------------------------------------------- 14.3：外链图片不取

class _Counter(http.server.BaseHTTPRequestHandler):
    hits: list = []

    def do_GET(self):  # noqa: N802
        type(self).hits.append(self.path)
        self.send_response(404)
        self.end_headers()

    def log_message(self, *a):
        pass


@pytest.fixture
def listener():
    _Counter.hits = []
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _Counter)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_address[1]}", _Counter.hits
    srv.shutdown()
    srv.server_close()


def linked_picture_docx(path: pathlib.Path, url: str) -> None:
    """一张"链接到文件"的图片：r:embed 改成 r:link，关系指向外部地址。"""
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(buf, "PNG")
    d = pydocx.Document()
    d.add_paragraph("正文")
    d.add_picture(io.BytesIO(buf.getvalue()))
    tmp = io.BytesIO()
    d.save(tmp)
    zin = zipfile.ZipFile(io.BytesIO(tmp.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            data = zin.read(n)
            if n == "word/document.xml":
                data = data.replace(b"r:embed=", b"r:link=")
            elif n == "word/_rels/document.xml.rels":
                s = data.decode()
                i = s.index('Target="media/')
                j = s.index('"', i + 8)
                s = s[:i] + f'Target="{url}/linked.png" TargetMode="External"' + s[j + 1:]
                data = s.encode()
            zout.writestr(n, data)


@needs_lo
def test_libreoffice_no_fetch(case, lo_base, listener, tmp_path):
    url, hits = listener
    p = tmp_path / "外链图片.docx"
    linked_picture_docx(p, url)
    out, used = C.OfficeConverter(lo_base).to_pdf(str(case), p, "工作区/临时/j", "libreoffice")
    assert used == "libreoffice" and pages(out) >= 1 and hits == []


def field_picture_docx(path: pathlib.Path, url: str) -> None:
    """正文里一个 INCLUDEPICTURE 域，地址指向外部。"""
    d = pydocx.Document()
    p = d.add_paragraph()
    for kind, text in (("begin", None), (None, f' INCLUDEPICTURE "{url}/f.png" \\d '), ("separate", None),
                       (None, None), ("end", None)):
        r = p.add_run()
        if kind:
            fc = r._r.makeelement(pydocx.oxml.ns.qn("w:fldChar"), {pydocx.oxml.ns.qn("w:fldCharType"): kind})
            r._r.append(fc)
        elif text:
            it = r._r.makeelement(pydocx.oxml.ns.qn("w:instrText"), {})
            it.text = text
            r._r.append(it)
    d.save(str(path))


def test_word_has_external(tmp_path):
    a, b, c = tmp_path / "链接图.docx", tmp_path / "域.docx", tmp_path / "普通.docx"
    linked_picture_docx(a, "http://127.0.0.1:9")
    field_picture_docx(b, "http://127.0.0.1:9")
    pydocx.Document().save(str(c))
    assert C.word_has_external(a) and C.word_has_external(b) and not C.word_has_external(c)
    assert not C.word_has_external(COMPLAINT)
    bad = tmp_path / "坏.docx"
    bad.write_bytes(b"PK\x03\x04 not a zip")
    assert C.word_has_external(bad)                                    # 查不了按有外链：不交 Word


@needs_lo
def test_external_link_skips_word_and_wps(case, lo_base, fake, tmp_path, listener):
    """auto：带外链图片的 docx 不交给 Word / WPS（它们会去取图），直接用 LibreOffice；选了 word 的转不了。"""
    url, hits = listener
    fake.set(word="ok", kwps="ok")
    p = tmp_path / "外链图片.docx"
    linked_picture_docx(p, url)
    out, used = C.OfficeConverter(lo_base).to_pdf(str(case), p, "工作区/临时/j")
    assert used == "libreoffice" and fake.calls() == [] and hits == []
    with pytest.raises(ApiError) as e:
        C.OfficeConverter(lo_base).to_pdf(str(case), p, "工作区/临时/j", "word")
    assert e.value.code == "CONVERTER_UNAVAILABLE" and fake.calls() == []


@pytest.mark.skipif(os.environ.get("LAWBENCH_REAL_WORD") != "1", reason="真 Word 实测：设 LAWBENCH_REAL_WORD=1")
def test_real_word(case, lo_base, listener, tmp_path):
    """真 Word：普通 docx 用 Word 转、转完不留 Word 进程；带外链图片的 docx 不交给 Word（监听 0 请求）。"""
    url, hits = listener
    before = C.processes()
    out, used = C.OfficeConverter(lo_base).to_pdf(str(case), COMPLAINT, "工作区/临时/j", "word")
    assert used == "word" and pages(out) >= 1
    assert C.com_started(before, C.processes(), C.SERVER_EXE["word"]) == []
    p = tmp_path / "外链图片.docx"
    linked_picture_docx(p, url)
    out, used = C.OfficeConverter(lo_base).to_pdf(str(case), p, "工作区/临时/j")
    assert used == "libreoffice" and hits == []


# ---------------------------------------------------------------- 第一轮复核返修（0054 令）：按文件头分流

@pytest.fixture
def linked_doc_bytes(tmp_path, lo_base):
    """LibreOffice 把"链接到文件"图片的 docx 存成旧版 .doc（OLE），带外链图片；地址指向一个不存在的端口。"""
    if not HAS_LO:
        pytest.skip("本机没有 LibreOffice")
    src = tmp_path / "linked.docx"
    linked_picture_docx(src, "http://127.0.0.1:9")
    with lo.Converter(tmp_path / "mk", lo_base=lo_base).session() as s:
        return s.convert(src, "doc").read_bytes()


def test_word_has_external_by_header(tmp_path):
    """P1-2：按文件头判断，扩展名怎么改都一样。"""
    linked = tmp_path / "a.docx"
    linked_picture_docx(linked, "http://127.0.0.1:9")
    for ext in (".doc", ".wps", ".docx", ".rtf"):
        x = tmp_path / f"改名{ext}"
        x.write_bytes(linked.read_bytes())
        assert C.word_has_external(x), ext                                   # docx 内容改名：照查 OOXML
    for content in (b"{\\rtf1\\ansi hello}", b"<html><img src='http://127.0.0.1:9/a.png'></html>", b"plain"):
        x = tmp_path / "冒充.doc"
        x.write_bytes(content)
        assert C.word_has_external(x)                                       # RTF、HTML、文本冒充：不交 Word
    plain = tmp_path / "普通.doc"
    plain.write_bytes(COMPLAINT.read_bytes())                               # 正常 docx 内容配 .doc 扩展名：可以交 Word
    assert not C.word_has_external(plain)


def altchunk_docx(path: pathlib.Path, url: str) -> None:
    d = pydocx.Document()
    d.add_paragraph("正文")
    tmp = io.BytesIO()
    d.save(tmp)
    zin = zipfile.ZipFile(io.BytesIO(tmp.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            zout.writestr(n, zin.read(n))
        zout.writestr("word/afchunk.html", f"<html><body><img src=\"{url}/c.png\"></body></html>")


def test_altchunk_html_external(tmp_path):
    p = tmp_path / "块.docx"
    altchunk_docx(p, "http://127.0.0.1:9")
    assert C.word_has_external(p) and C._altchunk_external(p)


def test_docx_renamed_doc_goes_to_libreoffice_not_word(case, lo_base, fake, tmp_path, listener):
    """P1-2 端到端：外链 docx 改名 .doc / .wps，auto 下假 Word / WPS 没被调用，LibreOffice 转、监听 0 请求。"""
    if not HAS_LO:
        pytest.skip("本机没有 LibreOffice")
    url, hits = listener
    fake.set(word="ok", kwps="ok", wps="ok")
    linked = tmp_path / "a.docx"
    linked_picture_docx(linked, url)
    for ext in (".doc", ".wps"):
        x = tmp_path / f"对方证据{ext}"
        x.write_bytes(linked.read_bytes())
        out, used = C.OfficeConverter(lo_base).to_pdf(str(case), x, "工作区/临时/j")
        assert used == "libreoffice" and pages(out) >= 1
    assert fake.calls() == [] and hits == []


def test_ole_doc_renamed_docx_refused_by_libreoffice(case, lo_base, fake, linked_doc_bytes, tmp_path, monkeypatch):
    """P1-1：旧版 .doc（有外链图片）改名 .docx：Word / WPS 不交（OLE 里查到外链），LibreOffice 也不交 → 转不了。"""
    started = []
    real = C.lo.Converter.session
    monkeypatch.setattr(C.lo.Converter, "session", lambda self: (started.append(1), real(self))[1])
    fake.set(word="ok", kwps="ok", wps="ok")
    x = tmp_path / "对方证据甲.docx"
    x.write_bytes(linked_doc_bytes)
    with pytest.raises(ApiError) as e:
        C.OfficeConverter(lo_base).to_pdf(str(case), x, "工作区/临时/j")
    assert e.value.code == "CONVERTER_UNAVAILABLE" and fake.calls() == [] and started == []


def test_xls_shell_checked_by_header(case, lo_base, tmp_path, monkeypatch):
    """NOTE-3：带外部图片关系的 xlsx 改名 .xls：按文件头走 14.3 ②b，不交 LibreOffice。"""
    started = []
    monkeypatch.setattr(C.lo.Converter, "session", lambda self: started.append(1))
    zin = zipfile.ZipFile(FIXTURES / "civil-01" / "银行流水.xlsx")
    x = tmp_path / "表.xls"
    with zipfile.ZipFile(x, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            zout.writestr(n, zin.read(n))
        zout.writestr("xl/drawings/_rels/drawing9.xml.rels",
                      '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/'
                      'package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
                      'officeDocument/2006/relationships/image" Target="http://127.0.0.1:9/x.png" TargetMode="External"/>'
                      "</Relationships>")
    with pytest.raises(ApiError):
        C.OfficeConverter(lo_base).to_pdf(str(case), x, "工作区/临时/j", "libreoffice")
    assert started == []


@pytest.mark.parametrize("content", [b"<html><body><img src='http://127.0.0.1:9/a.png'></body></html>",
                                     b"{\\rtf1 x}", b"PK\x03\x04 not text"])
def test_text_material_that_is_not_text_refused(case, lo_base, tmp_path, monkeypatch, content):
    started = []
    monkeypatch.setattr(C.lo.Converter, "session", lambda self: started.append(1))
    x = tmp_path / "说明.txt"
    x.write_bytes(content)
    with pytest.raises(ApiError):
        C.OfficeConverter(lo_base).to_pdf(str(case), x, "工作区/临时/j", "libreoffice")
    assert started == []


def test_com_worker_calls(monkeypatch):
    """P3-3：com_worker 对 COM 的调用（假 win32com）：新起实例、不可见、不弹窗、禁宏、只读打开且不进最近文档、
    只用导出接口不打印、UpdateLinksAtOpen 先关后恢复、关闭不保存、Quit。"""
    import types
    calls = []

    class Doc:
        def ExportAsFixedFormat(self, out, fmt):
            calls.append(("Export", out, fmt))

        def Close(self, save):
            calls.append(("Close", save))

        def PrintOut(self, *a):
            calls.append(("PrintOut",))

    class Opt:
        def __init__(self):
            object.__setattr__(self, "UpdateLinksAtOpen", True)

        def __setattr__(self, k, v):
            calls.append(("Options." + k, v))
            object.__setattr__(self, k, v)

    class Docs:
        def Open(self, *a, **kw):
            calls.append(("Open", a, kw))
            return Doc()

    class App:
        def __init__(self):
            object.__setattr__(self, "Options", Opt())
            object.__setattr__(self, "Documents", Docs())

        def __setattr__(self, k, v):
            calls.append(("App." + k, v))
            object.__setattr__(self, k, v)

        def Quit(self, *a):
            calls.append(("Quit", a))

    w32, cl = types.ModuleType("win32com"), types.ModuleType("win32com.client")
    cl.DispatchEx = lambda progid: (calls.append(("DispatchEx", progid)), App())[1]
    cl.Dispatch = lambda progid: (calls.append(("Dispatch", progid)), App())[1]
    w32.client = cl
    pc, pt = types.ModuleType("pythoncom"), types.ModuleType("pywintypes")
    pc.CoInitialize = lambda: None
    pc.CoUninitialize = lambda: None
    pt.com_error = type("com_error", (Exception,), {})
    for k, v in (("win32com", w32), ("win32com.client", cl), ("pythoncom", pc), ("pywintypes", pt)):
        monkeypatch.setitem(sys.modules, k, v)
    from lawbench.office import com_worker
    assert com_worker.main("Word.Application", "C:/x/in.docx", "C:/x/out.pdf") == 0
    assert calls == [
        ("DispatchEx", "Word.Application"),
        ("App.Visible", False), ("App.DisplayAlerts", 0), ("App.AutomationSecurity", 3),
        ("Options.UpdateLinksAtOpen", False),
        ("Open", ("C:/x/in.docx", False, True, False), {}),           # FileName, ConfirmConversions, ReadOnly, AddToRecentFiles
        ("Export", "C:/x/out.pdf", 17), ("Close", 0),
        ("Options.UpdateLinksAtOpen", True), ("Quit", (0,)),
    ]
    assert not any(c[0] in ("PrintOut", "Dispatch") for c in calls)


def test_com_conversions_one_at_a_time(lo_base, monkeypatch, tmp_path):
    """NOTE-1：两路同时转换时 Word / WPS 那一步串行（收尾按"本次期间新出现"认进程，并发会互相误杀）。"""
    import threading
    spans = []

    def slow(self, name, local):
        t0 = time.monotonic()
        time.sleep(0.5)
        spans.append((t0, time.monotonic()))
        return local
    monkeypatch.setattr(C.OfficeConverter, "_com_locked", slow)
    conv = C.OfficeConverter(lo_base)
    ts = [threading.Thread(target=conv._com, args=("word", tmp_path / "x")) for _ in range(2)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    (a0, a1), (b0, b1) = sorted(spans)
    assert b0 >= a1
