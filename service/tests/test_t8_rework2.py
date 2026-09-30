"""T8 第二轮返修（执行令 致B-ORCH-执行令-T8第二轮及T5第二轮返修-20260930-1318 第一节）逐条测试。"""
from __future__ import annotations

import pathlib
import sys
import threading
import time

import openpyxl
import pytest

from lawbench import contracts
from lawbench.case.task import TaskStore

from t8_helpers import CASE_FILES, Env, fail, ok


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    e = Env(tmp_path_factory.mktemp("t8r2"), {k: v for k, v in CASE_FILES.items() if k.endswith(".txt")})
    yield e
    e.close()


def store(e) -> TaskStore:
    return e.client.app.state.lb.tasks


def create(e, session, skill="contract-review"):
    req = {"case_id": e.case_id, "session_id": session, "entry": skill, "skill": skill, "inputs": [],
           "params": {"thinking": "中", "window": "64K", "max_tokens": 4096}}
    return ok(e.client.post("/api/task", json=req), "api/task_create.schema.json")["task_id"]


def end(e, tid, reason="completed"):
    return e.client.post("/core/task/end", json={"task_id": tid, "reason": reason, "model_calls": 1,
                                                 "tool_calls": 1, "elapsed_s": 1})


# ---------- F1：begin 与重开案件的竞态 ----------

def test_f1_reopen_during_begin_does_not_mark_new_task(env, monkeypatch):
    """begin 写完 running 的 task.json 后停住，这时调 mark_abnormal：新任务不能被标异常。"""
    s = store(env)
    real_write = TaskStore._write
    paused = threading.Event()

    def slow_write(self, root, rel, data, schema):
        real_write(self, root, rel, data, schema)
        if rel.endswith("task.json") and data.get("state") == "running" and \
                threading.current_thread().name == "begin":
            paused.set()
            time.sleep(0.5)

    monkeypatch.setattr(TaskStore, "_write", slow_write)
    out = {}
    t = threading.Thread(target=lambda: out.update(v=s.begin("sess-f1", str(env.root))), name="begin")
    t.start()
    assert paused.wait(5)
    marked = s.mark_abnormal(env.case_id)
    t.join(10)
    tid = out["v"]["task_id"]
    assert marked == 0
    assert env.read_json(tid, "task.json", "files/task.schema.json")["state"] == "running"
    monkeypatch.undo()
    env.tool_ok(tid, "case_save_draft", {"title": "F1", "content": "x"})


# ---------- F2：end 的任务锁；end 之后迟到的 save_draft ----------

def test_f2_save_draft_during_end_is_rejected(env, monkeypatch):
    """end 读完 result.json、正在算覆盖清单时 save_draft 插进来：有锁时 save_draft 等 end 做完、按已结束拒绝，
    不会给已结束的任务追加草稿，也不会被 end 写回的旧 result.json 冲掉。"""
    tid = env.begin("sess-f2")["task_id"]
    real_cov = TaskStore.coverage
    in_end = threading.Event()

    def slow_cov(self, root, case_id, task_id):
        if sys._getframe(1).f_code.co_name == "_end":
            in_end.set()
            time.sleep(0.5)
        return real_cov(self, root, case_id, task_id)

    monkeypatch.setattr(TaskStore, "coverage", slow_cov)
    t = threading.Thread(target=lambda: end(env, tid, "aborted"))
    t.start()
    assert in_end.wait(5)
    r = env.tool(tid, "case_save_draft", {"title": "迟到", "content": "x"})
    t.join(10)
    fail(r, "TASK_NOT_FOUND")
    res = env.read_json(tid, "result.json", "files/result.schema.json")
    assert res["status"] == "cancelled" and res["drafts"] == []
    assert not (env.task_dir(tid) / "草稿" / "迟到-v1.md").exists()


def test_f2_save_draft_after_end_rejected(env):
    tid = env.begin("sess-f2b")["task_id"]
    ok(end(env, tid), "core/task_end.schema.json")
    s = store(env)
    from lawbench.tools import ToolContext, run
    root = env.client.app.state.lb.cases.root_of(env.case_id)
    task = env.read_json(tid, "task.json", "files/task.schema.json")
    task["state"] = "running"   # 模拟 /core/tool 在锁外查状态时它还在执行
    ctx = ToolContext(root=root, case_id=env.case_id, task=task, tasks=s,
                      materials=env.client.app.state.lb.materials)
    with pytest.raises(Exception) as ei:
        run(ctx, "case_save_draft", {"title": "晚了", "content": "x"})
    assert getattr(ei.value, "code", None) == "TASK_NOT_FOUND"


# ---------- F3：读 json 遇 PermissionError 有上限地重试 ----------

def test_f3_read_json_retries_permission_error(tmp_path, monkeypatch):
    p = tmp_path / "a.json"
    p.write_text('{"x": 1}', encoding="utf-8")
    real = pathlib.Path.read_text
    calls = {"n": 0}

    def flaky(self, *a, **k):
        if self == p and calls["n"] < 2:
            calls["n"] += 1
            raise PermissionError(13, "正在被替换")
        return real(self, *a, **k)

    monkeypatch.setattr(pathlib.Path, "read_text", flaky)
    assert contracts.read_json(p) == {"x": 1} and calls["n"] == 2


def test_f3_read_json_gives_up(tmp_path, monkeypatch):
    p = tmp_path / "b.json"
    p.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(contracts, "REPLACE_WAIT", 0)
    monkeypatch.setattr(pathlib.Path, "read_text", lambda self, *a, **k: (_ for _ in ()).throw(PermissionError(13, "x")))
    with pytest.raises(PermissionError):
        contracts.read_json(p)


# ---------- F4：表头很宽时分段读也不超过 max_chars ----------

@pytest.fixture
def wide(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "宽表"
    for c in range(1, 181):
        ws.cell(row=1, column=c, value=f"列{c}")
    ws.cell(row=2, column=1, value="长" * 3000)
    wb.save(src / "宽表.xlsx")
    e = Env(tmp_path / "case", {"宽表.xlsx": src / "宽表.xlsx"})
    yield e
    e.close()


def test_f4_wide_header_part_within_limit(wide):
    tid = wide.begin("sess-f4")["task_id"]
    args = {"name": "宽表", "start": 2, "max_chars": 500}
    v = wide.tool_ok(tid, "case_read_material", args)
    got = ""
    n = 0
    while True:
        assert len(v["text"]) <= 500, len(v["text"])
        got += v["text"]
        n += 1
        if v["next_offset"] is None:
            break
        v = wide.tool_ok(tid, "case_read_material", dict(args, offset=v["next_offset"]))
    assert got.count("长") == 3000 and n > 1


# ---------- F5：同一秒平局以后写的为准 ----------

def test_f5_same_second_tie_prefers_latest(env, monkeypatch):
    from lawbench.case import task as task_mod
    for i in range(6):
        sess = f"sess-f5-{i}"
        monkeypatch.setattr(task_mod, "now_iso", lambda: "2026-09-30T10:00:00+08:00")
        a = create(env, sess, "contract-review")
        (env.task_dir(a) / "草稿").mkdir()
        (env.task_dir(a) / "草稿" / "人放的.md").write_text("x", encoding="utf-8")   # 旧单拒删
        create(env, sess, "case-summary")
        monkeypatch.undo()
        assert env.task_dir(a).exists()
        assert env.begin(sess)["skill"] == "case-summary", i
