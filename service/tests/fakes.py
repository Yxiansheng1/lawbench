"""假 6000D、假 395（Starlette 应用）。只返回固定内容，记下收到的请求头，用来核对转发和探测。"""
from __future__ import annotations

import threading

import anyio
from starlette.applications import Starlette
from starlette.requests import Request
from starlette.responses import JSONResponse, RedirectResponse, Response, StreamingResponse
from starlette.routing import Route

CHUNK1 = b'data: {"choices":[{"delta":{"content":"LBFAKE-CHUNK-1"}}]}\n\n'
CHUNK2 = b'data: {"choices":[{"delta":{"content":"LBFAKE-CHUNK-2"},"finish_reason":"stop"}]}\n\ndata: [DONE]\n\n'


class Fake6000D:
    def __init__(self, key_status: int = 200):
        self.key_status = key_status
        self.seen: list[dict] = []
        self.release = threading.Event()  # 流式第二段在测试读到第一段后才发出
        self.hits = 0

    def app(self) -> Starlette:
        async def models(request: Request):
            self.hits += 1
            self.seen.append({"path": request.url.path, "headers": dict(request.headers)})
            return JSONResponse({"object": "list", "data": [{"id": "qwen38-27b"}]})

        async def chat(request: Request):
            self.hits += 1
            body = await request.json()
            self.seen.append({"path": request.url.path, "headers": dict(request.headers), "body": body,
                              "query": str(request.url.query)})
            if not request.headers.get("authorization", "").startswith("Bearer "):
                return JSONResponse({"error": "no key"}, status_code=401)
            if not body.get("stream"):
                if self.key_status != 200:
                    return JSONResponse({"error": "key"}, status_code=self.key_status)
                return JSONResponse({"choices": [{"message": {"content": "x"}, "finish_reason": "length"}]})

            async def gen():
                yield CHUNK1
                await anyio.to_thread.run_sync(self.release.wait, 10)
                yield CHUNK2

            return StreamingResponse(gen(), media_type="text/event-stream", headers={"X-Queue-Wait-Ms": "12"})

        async def redirect(request: Request):
            self.hits += 1
            return RedirectResponse("http://127.0.0.1:1/elsewhere", status_code=302)

        async def other(request: Request):
            self.hits += 1
            return Response("upstream-other", status_code=200)

        return Starlette(routes=[
            Route("/v1/models", models, methods=["GET"]),
            Route("/v1/chat/completions", chat, methods=["POST"]),
            Route("/v1/redirect", redirect, methods=["GET"]),
            Route("/v1/embeddings", other, methods=["POST", "GET"]),
            Route("/admin", other, methods=["GET"]),
        ])


class Fake395:
    def __init__(self):
        self.hits = 0

    def app(self) -> Starlette:
        async def health(request: Request):
            self.hits += 1
            return JSONResponse({"status": "ok", "contract_version": "1.1"})

        return Starlette(routes=[Route("/health", health, methods=["GET"])])


# ---------- 外链图片样本与本机监听（Spec 14.3：转换不得按文档链接联网取图） ----------

class CountingListener:
    """127.0.0.1 随机端口上的 HTTP 监听，只数收到了几次请求。"""

    def __init__(self):
        import http.server
        import socketserver
        outer = self
        self.count = 0

        class H(http.server.BaseHTTPRequestHandler):
            def do_GET(self):  # noqa: N802
                outer.count += 1
                self.send_response(404)
                self.end_headers()

            do_HEAD = do_POST = do_GET  # noqa: N815

            def log_message(self, *a):
                pass

        class S(socketserver.ThreadingMixIn, http.server.HTTPServer):
            daemon_threads = True

        self.server = S(("127.0.0.1", 0), H)
        self.port = self.server.server_address[1]
        self.url = f"http://127.0.0.1:{self.port}"
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)

    def __enter__(self) -> "CountingListener":
        self.thread.start()
        return self

    def __exit__(self, *exc) -> None:
        self.server.shutdown()
        self.server.server_close()


def linked_image_docx(path, url: str, text: str = "正文里有一张外链图片。") -> None:
    """一份 docx：图片以外部链接引用（r:link + TargetMode=External），不内嵌。"""
    import zipfile
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    doc = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W}" xmlns:r="{R}"
 xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">
<w:body>
<w:p><w:r><w:t>{text}</w:t></w:r></w:p>
<w:p><w:r><w:drawing><wp:inline><wp:extent cx="914400" cy="914400"/><wp:docPr id="1" name="p"/>
<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">
<pic:pic><pic:nvPicPr><pic:cNvPr id="1" name="p"/><pic:cNvPicPr/></pic:nvPicPr>
<pic:blipFill><a:blip r:link="rIdImg"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>
<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm><a:prstGeom prst="rect"/></pic:spPr>
</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>
<w:p><w:r><w:t>另见网址 https://example.invalid/普通文字链接</w:t></w:r></w:p>
</w:body></w:document>'''
    rels = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rIdImg" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
 Target="{url}/pic.png" TargetMode="External"/>
</Relationships>'''
    ct = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml"
 ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>'''
    root_rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
 Target="word/document.xml"/>
</Relationships>'''
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("word/document.xml", doc)
        z.writestr("word/_rels/document.xml.rels", rels)


def linked_image_xlsx(path, url: str) -> None:
    """一份 xlsx：工作表里一张图片以外部链接引用（绘图部件的关系 TargetMode=External），不内嵌。
    LibreOffice Calc 打开时会去取这张图（BlockUntrustedRefererLinks 拦不住，Spec 14.3 ②b）。
    另有一个没有缓存值的公式，产品会想交给 LibreOffice 重算。"""
    import io
    import zipfile

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "流水"
    ws["A1"], ws["B1"], ws["B2"] = "金额", 100, "=B1*2"
    buf = io.BytesIO()
    wb.save(buf)
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    drawing = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="{R}">
<xdr:oneCellAnchor><xdr:from><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
<xdr:ext cx="914400" cy="914400"/>
<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="p"/><xdr:cNvPicPr/></xdr:nvPicPr>
<xdr:blipFill><a:blip r:link="rIdImg"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>
<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm><a:prstGeom prst="rect"/></xdr:spPr>
</xdr:pic><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>'''
    drawing_rels = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rIdImg" Type="{R}/image" Target="{url}/sheet.png" TargetMode="External"/>
</Relationships>'''
    sheet_rels = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rIdDr" Type="{R}/drawing" Target="../drawings/drawing1.xml"/>
</Relationships>'''
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.decode("utf-8")
                if "xmlns:r=" not in data:
                    data = data.replace("<worksheet ", f'<worksheet xmlns:r="{R}" ', 1)
                data = data.replace("</worksheet>", '<drawing r:id="rIdDr"/></worksheet>').encode("utf-8")
            elif info.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/xl/drawings/drawing1.xml" ContentType='
                                    b'"application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
            z.writestr(info.filename, data)
        z.writestr("xl/worksheets/_rels/sheet1.xml.rels", sheet_rels)
        z.writestr("xl/drawings/drawing1.xml", drawing)
        z.writestr("xl/drawings/_rels/drawing1.xml.rels", drawing_rels)


def minimal_ole(path, streams: dict[str, bytes]) -> None:
    """最小的 OLE 复合文档（CFB v3，512 字节扇区）：每个流补齐到 4096 字节以上，不用迷你流。
    给"扩展名 .doc、内容是 OLE"的分流测试用；内容不是真的 Word 文档，转换程序打开会失败。"""
    import struct
    END, FREE, FATSECT, NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    datas = [(n, d.ljust(max(4096, len(d)), b"\0")) for n, d in streams.items()]
    fat = [FATSECT, END]
    body = b""
    starts = []
    for _, d in datas:
        d = d.ljust(-(-len(d) // 512) * 512, b"\0")
        n = len(d) // 512
        first = len(fat)
        starts.append(first)
        fat += [first + i + 1 for i in range(n - 1)] + [END]
        body += d
    assert len(fat) <= 128
    fat_sector = struct.pack("<128I", *(fat + [FREE] * (128 - len(fat))))

    def entry(name, typ, child, right, start, size):
        raw = name.encode("utf-16-le") + b"\0\0"
        return (raw.ljust(64, b"\0") + struct.pack("<HBB3I", len(raw), typ, 1, NOSTREAM, right, child)
                + b"\0" * 16 + b"\0" * 4 + b"\0" * 16 + struct.pack("<IIi", start, size, 0))

    entries = [entry("Root Entry", 5, 1 if datas else NOSTREAM, NOSTREAM, END, 0)]
    for i, ((name, d), st) in enumerate(zip(datas, starts)):
        entries.append(entry(name, 2, NOSTREAM, i + 2 if i + 1 < len(datas) else NOSTREAM, st, len(d)))
    while len(entries) % 4:
        entries.append(b"\0" * 64 + struct.pack("<HBB3I", 0, 0, 0, NOSTREAM, NOSTREAM, NOSTREAM) + b"\0" * 48)
    assert len(entries) == 4
    header = (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 16 + struct.pack("<HHHHH", 0x3E, 3, 0xFFFE, 9, 6)
              + b"\0" * 6 + struct.pack("<IIIIIIIII", 0, 1, 1, 0, 4096, END, 0, END, 0)
              + struct.pack("<109I", 0, *([FREE] * 108)))
    assert len(header) == 512
    with open(path, "wb") as f:
        f.write(header + fat_sector + b"".join(entries) + body)
