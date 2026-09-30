"""T5 第二轮返修（执行令 致B-ORCH-执行令-T8第二轮及T5第二轮返修-20260930-1318 第二节）与 N44 ①。编号与返修令一致。"""
from __future__ import annotations

import json
import os
import pathlib
import sqlite3
import subprocess
import sys
import zipfile

import openpyxl
import pytest

from lawbench.case import gate, materials as mat_mod
from lawbench.case.materials import Materials
from lawbench.config import REPO_ROOT
from lawbench.ingest import REASONS, ParseError, docx, xlsx

from conftest import IS_WIN

FIXTURES = REPO_ROOT / "tests" / "fixtures"
SERVICE_DIR = pathlib.Path(__file__).resolve().parents[1]


def open_case(client, root: pathlib.Path) -> str:
    root.mkdir(parents=True, exist_ok=True)
    r = client.post("/api/case/open", json={"path": str(root)}).json()
    assert r["ok"], r
    return r["value"]["case_id"]


def scan(client, cid: str) -> dict:
    r = client.post("/api/materials/scan", json={"case_id": cid}).json()
    assert r["ok"], r
    return r["value"]


def do_import(client, cid: str, paths: list, target=None) -> dict:
    r = client.post("/api/materials/import", json={"case_id": cid, "paths": [str(p) for p in paths],
                                                   "target": target, "unzip": False}).json()
    assert r["ok"], r
    return r["value"]


def by_rel(root: pathlib.Path) -> dict:
    idx = json.loads((root / "工作区" / "材料" / "index.json").read_text(encoding="utf-8"))
    return {m["rel_path"]: m for m in idx["materials"]}


def originals(root: pathlib.Path) -> list[str]:
    return sorted(str(x.relative_to(root)).replace("\\", "/") for x in root.rglob("*")
                  if x.is_file() and x.relative_to(root).parts[0] not in ("工作区", "成果"))


def _long(p: pathlib.Path) -> str:
    return "\\\\?\\" + str(p)


def _unc(p: pathlib.Path) -> str | None:
    """C:\\x → \\\\localhost\\C$\\x；本机访问不了管理共享时返回 None。"""
    drive, rest = os.path.splitdrive(str(p))
    if not drive.endswith(":"):
        return None
    unc = f"\\\\localhost\\{drive[0]}$" + rest
    return unc if os.path.exists(unc) else None


# ---------- B-P2-1：导入源的别名写法 ----------

@pytest.fixture
def nested(make_client, cases_dir):
    client = make_client()
    root = cases_dir / "上级" / "案件"
    cid = open_case(client, root)
    (root / "证据.txt").write_text("证据", encoding="utf-8")
    scan(client, cid)
    (cases_dir / "上级" / "别的.txt").write_text("别的", encoding="utf-8")
    return client, root, cid, cases_dir / "上级"


@pytest.mark.skipif(not IS_WIN, reason="Windows 路径写法")
@pytest.mark.parametrize("form", ["long", "unc"])
def test_bp2_1_alias_of_root_only_scans(nested, form):
    client, root, cid, _ = nested
    src = _long(root) if form == "long" else _unc(root)
    if src is None:
        pytest.skip("本机访问不了 \\\\localhost\\C$ 管理共享")
    before = originals(root)
    v = do_import(client, cid, [src])
    assert v["copied"] == [] and originals(root) == before


@pytest.mark.skipif(not IS_WIN, reason="Windows 路径写法")
@pytest.mark.parametrize("form", ["long", "unc"])
def test_bp2_1_alias_of_parent_skips_case_branch(nested, form):
    client, root, cid, parent = nested
    src = _long(parent) if form == "long" else _unc(parent)
    if src is None:
        pytest.skip("本机访问不了 \\\\localhost\\C$ 管理共享")
    v = do_import(client, cid, [src])
    assert [c["to"] for c in v["copied"]] == ["上级/别的.txt"]
    assert originals(root) == ["上级/别的.txt", "证据.txt"]          # 没有套娃


@pytest.mark.skipif(not IS_WIN, reason="Windows 路径写法")
def test_bp2_1_alias_of_temp_file_deleted_after_copy(nested):
    client, root, cid, _ = nested
    t = root / "工作区" / "临时"
    t.mkdir(parents=True, exist_ok=True)
    shot = t / "截图.png"
    shot.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 20)
    v = do_import(client, cid, [_long(shot)])
    assert [c["to"] for c in v["copied"]] == ["截图.png"] and not shot.exists()


def test_bp2_1_parts_in_case():
    root = str(REPO_ROOT)
    assert mat_mod._parts_in_case(root, root) == []
    assert mat_mod._parts_in_case(str(REPO_ROOT / "service" / "lawbench"), root) == ["service", "lawbench"]
    assert mat_mod._parts_in_case(str(REPO_ROOT.parent), root) is None


# ---------- A-P2-1 = B-P2-2：1.2 之前建的 case.db ----------

def test_ap2_1_old_db_upgraded_on_open_and_ids_kept(make_client, cases_dir):
    client = make_client()
    root = cases_dir / "旧库"
    cid = open_case(client, root)
    db = root / "工作区" / "case.db"
    con = sqlite3.connect(db)
    con.execute("DROP TABLE material_ids")                          # 当成 1.2 之前建的库
    con.commit()
    con.close()
    open_case(client, root)                                           # 再打开：补建
    con = sqlite3.connect(db)
    assert con.execute("SELECT 1 FROM sqlite_master WHERE name = 'material_ids'").fetchone()
    con.close()
    (root / "a.txt").write_text("a", encoding="utf-8")
    (root / "b.txt").write_text("b", encoding="utf-8")
    scan(client, cid)
    ids = {r: m["material_id"] for r, m in by_rel(root).items()}
    (root / "工作区" / "材料" / "index.json").unlink()
    scan(client, cid)
    assert {r: m["material_id"] for r, m in by_rel(root).items()} == ids   # 拿回原编号


def test_ap2_1_missing_table_recreated_not_silent(make_client, cases_dir, appdata):
    client = make_client()
    root = cases_dir / "表没了"
    cid = open_case(client, root)
    con = sqlite3.connect(root / "工作区" / "case.db")
    con.execute("DROP TABLE material_ids")                          # 打开之后被人动过
    con.commit()
    con.close()
    (root / "a.txt").write_text("a", encoding="utf-8")
    scan(client, cid)
    con = sqlite3.connect(root / "工作区" / "case.db")
    rows = con.execute("SELECT material_id, rel_path FROM material_ids").fetchall()
    con.close()
    assert rows == [("M0001", "a.txt")]
    from lawbench import logs
    logs.close()
    text = "".join(p.read_text(encoding="utf-8") for p in (appdata / "logs").glob("*"))
    assert "TABLE_MISSING" in text


# ---------- B-P2-3：常规格式按 15 位有效数字 ----------

def _raw_values(path: pathlib.Path, values: dict[str, str]) -> None:
    """openpyxl 自造的样本只写短小数；这里直接改工作表 XML 的缓存值，写成 Excel 常见的 17 位。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, cell in enumerate(values, 1):
        ws[cell] = float(i) + 0.5
    wb.save(path.with_suffix(".tmp"))
    with zipfile.ZipFile(path.with_suffix(".tmp")) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                for i, (cell, raw) in enumerate(values.items(), 1):
                    text = text.replace(f"<v>{float(i) + 0.5}</v>", f"<v>{raw}</v>", 1)
                    assert raw in text
                data = text.encode("utf-8")
            zout.writestr(info, data)
    path.with_suffix(".tmp").unlink()


def test_bp2_3_general_uses_15_significant_digits(tmp_path):
    p = tmp_path / "尾数.xlsx"
    _raw_values(p, {"A1": "1234.6599999999999", "B1": "0.30000000000000004", "C1": "12345.670000000002",
                    "D1": "1.4999999999999999E-7"})
    text = xlsx.parse(p).blocks[0].text
    assert "| 1 | 1234.66 | 0.3 | 12345.67 | 0.00000015 |" in text, text


# ---------- B-P2-4：密集表格、段落极多的 docx，加载前按过大 ----------

PEAK = r'''
import ctypes, ctypes.wintypes as w, json, pathlib, sys, time
sys.path.insert(0, sys.argv[1])
from lawbench.ingest import ParseError, docx, xlsx
class PMC(ctypes.Structure):
    _fields_ = [("cb", w.DWORD), ("PageFaultCount", w.DWORD), ("PeakWorkingSetSize", ctypes.c_size_t),
                ("WorkingSetSize", ctypes.c_size_t), ("a", ctypes.c_size_t), ("b", ctypes.c_size_t),
                ("c", ctypes.c_size_t), ("d", ctypes.c_size_t), ("PagefileUsage", ctypes.c_size_t),
                ("PeakPagefileUsage", ctypes.c_size_t)]
def peak():
    m = PMC(); m.cb = ctypes.sizeof(PMC)
    ctypes.windll.psapi.GetProcessMemoryInfo(ctypes.windll.kernel32.GetCurrentProcess(), ctypes.byref(m), m.cb)
    return m.PeakWorkingSetSize
base = peak()
t0 = time.monotonic()
try:
    (xlsx if sys.argv[3] == "xlsx" else docx).parse(pathlib.Path(sys.argv[2]))
    reason = None
except ParseError as e:
    reason = e.reason
print(json.dumps({"reason": reason, "seconds": time.monotonic() - t0, "peak_mb": (peak() - base) / 2**20}))
'''


def _measure(tmp_path, path: pathlib.Path, kind: str) -> dict:
    script = tmp_path / "peak.py"
    script.write_text(PEAK, encoding="utf-8")
    out = subprocess.run([sys.executable, str(script), str(SERVICE_DIR), str(path), kind], capture_output=True,
                         text=True, timeout=300)
    assert out.returncode == 0, out.stderr
    return json.loads(out.stdout.strip().splitlines()[-1])


def _dense_xlsx(path: pathlib.Path, rows: int, cols: int) -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    wb.save(path.with_suffix(".tmp"))
    letters = [openpyxl.utils.get_column_letter(c) for c in range(1, cols + 1)]
    with zipfile.ZipFile(path.with_suffix(".tmp")) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename != "xl/worksheets/sheet1.xml":
                zout.writestr(info, zin.read(info))
                continue
            head, _, tail = zin.read(info).decode("utf-8").partition("<sheetData>")
            _, _, tail = tail.partition("</sheetData>")
            with zout.open(info.filename, "w") as f:
                f.write((head + "<sheetData>").encode("utf-8"))
                for r in range(1, rows + 1):
                    f.write((f'<row r="{r}">' + "".join(f'<c r="{c}{r}"><v>{r}</v></c>' for c in letters)
                             + "</row>").encode("utf-8"))
                f.write(("</sheetData>" + tail).encode("utf-8"))
    path.with_suffix(".tmp").unlink()


@pytest.mark.skipif(not IS_WIN, reason="峰值内存用 Windows API 量")
def test_bp2_4_dense_sheet_rejected_before_load(tmp_path):
    p = tmp_path / "密集.xlsx"
    _dense_xlsx(p, 100_000, 30)                                        # 300 万格
    r = _measure(tmp_path, p, "xlsx")
    assert r["reason"] == "too_large"
    assert r["seconds"] < 60 and r["peak_mb"] < 200, r


def test_bp2_4_dense_sheet_under_limit_still_parsed(tmp_path):
    p = tmp_path / "不太密.xlsx"
    _dense_xlsx(p, 2000, 30)
    out = xlsx.parse(p)
    assert "| 2000 | 2000 |" in out.blocks[0].text


@pytest.mark.skipif(not IS_WIN, reason="峰值内存用 Windows API 量")
def test_bp2_4_docx_too_many_paragraphs(tmp_path):
    p = tmp_path / "段落.docx"
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with zipfile.ZipFile(FIXTURES / "civil-01" / "借条.docx") as zin, \
            zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename != "word/document.xml":
                zout.writestr(info, zin.read(info))
        with zout.open("word/document.xml", "w") as f:
            f.write(f'<?xml version="1.0"?><w:document xmlns:w="{W}"><w:body>'.encode("utf-8"))
            chunk = "<w:p><w:r><w:t>段</w:t></w:r></w:p>".encode("utf-8") * 1000
            for _ in range(600):                                          # 60 万段
                f.write(chunk)
            f.write(b"</w:body></w:document>")
    r = _measure(tmp_path, p, "docx")
    assert r["reason"] == "too_large"
    assert r["seconds"] < 60 and r["peak_mb"] < 200, r


# ---------- X6：只认本服务自己的形状 ----------

def test_x6_only_own_shapes_cleaned(make_client, cases_dir):
    client = make_client()
    root = cases_dir / "x6形状"
    cid = open_case(client, root)
    t = root / "工作区" / "临时"
    t.mkdir(parents=True, exist_ok=True)
    (t / "lo-1a2b3c4d").mkdir()
    (t / "unzip-5e6f7a8b").mkdir()
    (t / "imp-0000aaaa").write_bytes(b"x")
    keep = ["lo-截图.png", "imp-合同扫描.pdf", "unzip-说明.txt", "lo-9a9a9a9a", "imp-0000aaaa.pdf", "lo-ABCDEF12"]
    for n in keep:
        (t / n).write_bytes(b"lawyer")
    (t / "imp-1111bbbb").mkdir()                                      # 形状对、但 imp- 应是文件：不动
    scan(client, cid)
    assert sorted(p.name for p in t.iterdir()) == sorted(keep + ["imp-1111bbbb"])


# ---------- X9：第一次读不了、重试成功 ----------

def test_x9_first_unreadable_then_retry_counts_and_sha(make_client, cases_dir, monkeypatch):
    client = make_client()
    root = cases_dir / "x9计数"
    cid = open_case(client, root)
    (root / "占用.txt").write_text("被占用", encoding="utf-8")
    real = mat_mod.sha256_file
    monkeypatch.setattr(mat_mod, "sha256_file",
                        lambda p: (_ for _ in ()).throw(PermissionError(32, "x")) if str(p).endswith("占用.txt")
                        else real(p))
    assert scan(client, cid)["added"] == 1
    monkeypatch.setattr(mat_mod, "sha256_file", real)
    v = scan(client, cid)
    assert v["changed"] == 0 and v["added"] == 0 and v["review_needed"] is True
    m = by_rel(root)["占用.txt"]
    assert m["status"] == "parsed" and m["sha256"] == real(root / "占用.txt")
    con = sqlite3.connect(root / "工作区" / "case.db")
    assert con.execute("SELECT sha256 FROM material_ids WHERE material_id = ?", (m["material_id"],)).fetchone() \
        == (m["sha256"],)
    con.close()


# ---------- N44 ①：内容不是 Office 格式的 .xls / .doc ----------

@pytest.mark.parametrize("name", ["银行流水.xls", "说明.doc"])
def test_n44_not_office_reason_and_no_retry(make_client, cases_dir, monkeypatch, name):
    client = make_client()
    root = cases_dir / "n44"
    cid = open_case(client, root)
    (root / name).write_text("<html><body><table><tr><td>80000</td></tr></table></body></html>", encoding="utf-8")
    scan(client, cid)
    m = by_rel(root)[name]
    assert m["status"] == "failed" and m["error"] == REASONS["not_office"]
    assert "银行导出" in m["error"]
    called = []
    monkeypatch.setattr(Materials, "_parse", lambda self, *a, **k: called.append(a))
    scan(client, cid)
    assert called == []                                                 # 文件本身的失败：原件没变不重试


# ---------- A-P3-3：PDF、图片在格式版本变了时也重扫 ----------

@pytest.mark.parametrize("src", ["criminal-01/起诉意见书.pdf", "image"])
def test_reformat_pdf_and_image(make_client, cases_dir, src):
    client = make_client()
    root = cases_dir / "重扫"
    cid = open_case(client, root)
    f = (next(FIXTURES.rglob("*.png"), None) or next(FIXTURES.rglob("*.jpg"))) if src == "image" else FIXTURES / src
    (root / ("材料" + f.suffix)).write_bytes(f.read_bytes())
    scan(client, cid)
    m = by_rel(root)["材料" + f.suffix]
    text = root / m["text_path"]
    status = root / "工作区" / "材料" / "_处理状态.md"
    text.write_text("旧格式", encoding="utf-8")
    status.write_text(status.read_text(encoding="utf-8").replace("材料文本格式版本：", "旧："), encoding="utf-8")
    scan(client, cid)
    assert text.read_text(encoding="utf-8") != "旧格式"
