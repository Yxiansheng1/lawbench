"""T5 第三轮小修（执行令 致B-ORCH-执行令-T5第三轮小修-20260930-1559）。编号与执行令一致。"""
from __future__ import annotations

import json
import pathlib
import sqlite3
import zipfile

import openpyxl
import pytest

from lawbench.ingest import ParseError, xlsx
from lawbench.ingest import MAX_SHEET_CELLS


def _raw_value(path: pathlib.Path, raw: str) -> None:
    """用原始 XML 写单元格缓存值（openpyxl 自造的样本写不出 Excel 那种科学计数的缓存值）。"""
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1.5
    tmp = path.with_suffix(".tmp")
    wb.save(tmp)
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                assert "<v>1.5</v>" in text
                data = text.replace("<v>1.5</v>", f"<v>{raw}</v>").encode("utf-8")
            zout.writestr(info, data)
    tmp.unlink()


# ---------- P2-1：15 位以上的整数 ----------

@pytest.mark.parametrize("raw,shown", [
    ("1.10101199003071E+17", "110101199003071000"),   # 证件号被存成数字：Excel 显示 15 位有效数字
    ("6.2284801234567898E+18", "6228480123456790000"),
    ("123456789012345", "123456789012345"),           # 15 位以内照旧
    ("80000", "80000"),
])
def test_p2_1_big_integers_15_digits(tmp_path, raw, shown):
    p = tmp_path / "证件号.xlsx"
    _raw_value(p, raw)
    assert f"| 1 | {shown} |" in xlsx.parse(p).blocks[0].text


# ---------- P3-1：补表后按 index.json 回填 ----------

def test_p3_1_backfill_old_db_keeps_ids(make_client, cases_dir):
    client = make_client()
    root = cases_dir / "旧库三份"
    root.mkdir()
    cid = client.post("/api/case/open", json={"path": str(root)}).json()["value"]["case_id"]
    for n in ("甲.txt", "乙.txt", "丙.txt"):   # 分三次加进来：编号顺序与重扫时按文件名排的顺序不同
        (root / n).write_text(n, encoding="utf-8")
        assert client.post("/api/materials/scan", json={"case_id": cid}).json()["ok"]
    idx_path = root / "工作区" / "材料" / "index.json"
    ids = {m["rel_path"]: m["material_id"] for m in json.loads(idx_path.read_text(encoding="utf-8"))["materials"]}
    assert ids == {"甲.txt": "M0001", "乙.txt": "M0002", "丙.txt": "M0003"}
    assert sorted(ids) == ["丙.txt", "乙.txt", "甲.txt"]                    # 不回填的话重扫会按这个顺序重新编号
    db = root / "工作区" / "case.db"
    con = sqlite3.connect(db)
    con.execute("DROP TABLE material_ids")                                 # 当成 1.2 之前建的、已有三份材料的库
    con.execute("DELETE FROM meta WHERE key = 'next_material_seq'")
    con.commit()
    con.close()
    client.post("/api/case/open", json={"path": str(root)})                 # 升级：补表（空的）
    assert client.post("/api/materials/scan", json={"case_id": cid}).json()["ok"]   # 扫描时回填
    idx_path.unlink()
    assert client.post("/api/materials/scan", json={"case_id": cid}).json()["ok"]
    again = {m["rel_path"]: m["material_id"] for m in json.loads(idx_path.read_text(encoding="utf-8"))["materials"]}
    assert again == ids


# ---------- P3-2：工作表部件不在 xl/worksheets/ 下 ----------

def _moved_sheet(path: pathlib.Path, rows: int, cols: int) -> None:
    """工作表部件放在 xl/sheets/big.xml，由 workbook.xml.rels 指过去。"""
    wb = openpyxl.Workbook()
    wb.active["A1"] = 1
    tmp = path.with_suffix(".tmp")
    wb.save(tmp)
    letters = [openpyxl.utils.get_column_letter(c) for c in range(1, cols + 1)]
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == "xl/worksheets/sheet1.xml":
                head, _, tail = data.decode("utf-8").partition("<sheetData>")
                _, _, tail = tail.partition("</sheetData>")
                with zout.open("xl/sheets/big.xml", "w") as f:
                    f.write((head + "<sheetData>").encode("utf-8"))
                    for r in range(1, rows + 1):
                        f.write((f'<row r="{r}">' + "".join(f'<c r="{c}{r}"><v>{r}</v></c>' for c in letters)
                                 + "</row>").encode("utf-8"))
                    f.write(("</sheetData>" + tail).encode("utf-8"))
                continue
            if info.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"worksheets/sheet1.xml", b"sheets/big.xml").replace(
                    b"/xl/worksheets/sheet1.xml", b"/xl/sheets/big.xml")
            if info.filename == "[Content_Types].xml":
                data = data.replace(b"/xl/worksheets/sheet1.xml", b"/xl/sheets/big.xml")
            zout.writestr(info, data)
    tmp.unlink()


def test_p3_2_sheet_part_located_by_workbook_rels(tmp_path):
    p = tmp_path / "挪了位置.xlsx"
    _moved_sheet(p, MAX_SHEET_CELLS // 30 + 10, 30)                         # 略超上限
    with zipfile.ZipFile(p) as z:
        assert "xl/sheets/big.xml" in z.namelist() and not any(n.startswith("xl/worksheets/") for n in z.namelist())
    with pytest.raises(ParseError) as ei:
        xlsx._check_cells(p)
    assert ei.value.reason == "too_large"


def test_p3_2_moved_small_sheet_still_parsed(tmp_path):
    p = tmp_path / "挪了位置小.xlsx"
    _moved_sheet(p, 5, 3)
    assert "| 5 | 5 | 5 | 5 |" in xlsx.parse(p).blocks[0].text
