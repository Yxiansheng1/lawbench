"""归档用的 Word 文书（Spec 12.4 第 3、5、7、8 步；原版规则 lawyer-archiving 第四、五、七、八步）。

- 结案报告：skills/case-archiving/templates/结案报告模板.docx 在就填它，不在用临时版式（仿宋 12 号，一页）。
- 立卷申请书：templates/立卷申请书模板_<卷类>.docx（民事行政卷另认原版的 立卷申请书模板.docx）在就按原版第七步
  填写，不在用临时版式生成同样内容的表格。
- 情况说明：references/特殊情况说明模板.md 的模板一、模板二。
- 归档目录.md。
全部在内存里生成，返回字节；不碰文件系统（模板只读）。
"""
from __future__ import annotations

import io
import pathlib
import re

import docx as pydocx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FIRM = "广东连越（深圳）律师事务所"
TODO = "【待填写】"
_PH = re.compile(r"【[^】]*】")
_PAREN = re.compile(r"（[^）]*）$")


def plain_name(name: str) -> str:
    """目录项名称去掉末尾括号里的提示（"民事委托代理合同（合同编号……必填）" → "民事委托代理合同"）。"""
    return _PAREN.sub("", name).strip()


def page_label(rng: tuple[int, int] | None) -> str:
    if not rng:
        return ""
    a, b = rng
    return f"p{a}" if a == b else f"p{a}-{b}"


def _font(run, name: str = "仿宋", size: float = 12, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _para(doc, text: str, *, font="仿宋", size=12, bold=False, center=False, before=0, after=2, indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.15
    if indent:
        pf.first_line_indent = Pt(size * 2)
    _font(p.add_run(text), font, size, bold)
    return p


def _merge_runs(p) -> None:
    """占位符可能被拆成多个 run：把整段文字放进第一个 run（保留它的格式），其余 run 清空。"""
    runs = p.runs
    if len(runs) > 1:
        runs[0].text = "".join(r.text for r in runs)
        for r in runs[1:]:
            r.text = ""


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ================================================================ 结案报告

def report_fields(plan: dict, lawyer: str | None) -> dict[str, str]:
    changfa = plan["catalog"] == "常法卷"
    opp = plan["opponent"] or ("无（常年法律顾问）" if changfa else "无")
    return {"client": plan["client"], "opponent": opp, "entrust": plan["entrust_date"] or TODO,
            "close": plan["close_date"] or TODO, "lawyer": lawyer or TODO, "cause": plan["cause"] or TODO,
            "summary": plan["summary"] or TODO, "opinion": plan["opinion"] or TODO, "result": plan["result"] or TODO}


# 段落标签（"："之前的文字）→ 字段；先查长的、具体的（"承办律师分析与意见"含"承办律师"）
_REPORT_LABELS = (("分析与意见", "opinion"), ("服务结果", "opinion"), ("案情简介", "summary"), ("服务概况", "summary"),
                  ("办案结果", "result"), ("对方当事人", "opponent"), ("委托日期", "entrust"), ("结案日期", "close"),
                  ("承办律师", "lawyer"), ("案由", "cause"), ("委托人", "client"))


def _label_field(label: str) -> str | None:
    return next((f for key, f in _REPORT_LABELS if key in label), None)


def fill_report_template(template: pathlib.Path, f: dict[str, str]) -> tuple[bytes, list[str]]:
    """按段落标签填【…】：段落是"案情简介：…""分析与意见：…"的，冒号后整段换成律师确认过的内容（AI 写的是一整段）；
    其余段落里的每个【…】按该段标签填。认不出标签的占位符原样留着，返回提示。"""
    doc = pydocx.Document(str(template))
    unknown = 0
    for p in doc.paragraphs:
        if not _PH.search(p.text):
            continue
        _merge_runs(p)
        run = p.runs[0]
        text = run.text
        label = re.split(r"[：:]", text, 1)[0] if re.search(r"[：:]", text) else ""
        field = _label_field(label)
        if field in ("summary", "opinion") and label:
            run.text = text[: len(label) + 1] + f[field]
        elif field:
            run.text = _PH.sub(lambda _m: f[field], text)
        else:
            unknown += len(_PH.findall(text))
    notes = [f"结案报告模板里有 {unknown} 处认不出的【】占位符，已原样保留，请在 Word 中手工填写"] if unknown else []
    return _save(doc), notes


def make_report(f: dict[str, str], changfa: bool) -> bytes:
    """临时版式：仿宋 12 号，一页（原版第四步的结构）。"""
    doc = pydocx.Document()
    for s in doc.sections:
        s.top_margin, s.bottom_margin, s.left_margin, s.right_margin = Cm(2.0), Cm(1.8), Cm(2.4), Cm(2.0)
    _para(doc, "结案报告", font="黑体", size=16, bold=True, center=True, after=10)
    rows = [("一、当事人/委托人：", f["client"]), ("对方当事人：", f["opponent"]), ("二、委托日期：", f["entrust"]),
            ("三、结案日期：", f["close"]), ("四、承办律师：", f["lawyer"]), ("五、案由：", f["cause"]),
            ("六、服务概况：" if changfa else "六、案情简介：", f["summary"]),
            ("七、服务结果：" if changfa else "七、承办律师分析与意见：", f["opinion"]),
            ("八、办案结果：", f["result"])]
    for label, value in rows:
        _para(doc, label + value, before=3)
    return _save(doc)


# ================================================================ 立卷申请书

def application_rows(plan: dict, catalog: dict, ranges: dict[int, tuple[int, int]]) -> list[tuple[int, str, str]]:
    """[(目录原编号, 写进申请书的编号, "材料名称 p12-15")]：只列实际归档的项（方案里的 + 程序生成的结案报告），按目录编号；
    常法卷等 renumber 的连续编号，其余保留原编号（原版 7.2）。名称用方案里按本案改写的，结案报告用目录名去掉括号。"""
    names = {it["code"]: it["name"] for it in plan["items"]}
    rows = []
    for it in catalog["items"]:
        if it["code"] in names:
            name = names[it["code"]]
        elif it.get("generated") and it["code"] in ranges:
            name = plain_name(it["name"])
        else:
            continue
        rows.append((it["code"], f"{plain_name(name)} {page_label(ranges.get(it['code']))}".strip()))
    if catalog["renumber"]:
        return [(c, str(i), t) for i, (c, t) in enumerate(rows, 1)]
    return [(c, str(c), t) for c, t in rows]


def _case_no(plan: dict) -> str:
    return f"（{plan['jzl_no'] or '【待填写：金助理系统案件编号】'}+{plan['client']}）"


def fill_application_template(template: pathlib.Path, plan: dict,
                              rows: list[tuple[int, str, str]]) -> tuple[bytes, list[str]]:
    """原版第七步：①"金助理系统"那段的【…】换成（编号+客户名称）；②表格 0 只留归档的行（按第一列的原编号），
    ③常法卷等重新连续编号，④每行写本案材料名和页码，⑤内容行仿宋 12 号不加粗、表头不动。"""
    doc = pydocx.Document(str(template))
    notes: list[str] = []
    for p in doc.paragraphs:
        if "金助理系统" in p.text and _PH.search(p.text):
            _merge_runs(p)
            p.runs[0].text = _PH.sub(lambda _m: _case_no(plan), p.runs[0].text, count=1)
            break
    else:
        notes.append("立卷申请书模板里没找到\"金助理系统案件编号\"占位，请手工填写案件编号")
    if not doc.tables:
        return _save(doc), notes + ["立卷申请书模板里没有表格，材料清单请手工填写"]
    table = doc.tables[0]
    keep = {code: (no, text) for code, no, text in rows}
    for row in list(table.rows)[1:][::-1]:                # 从下往上删，免得下标错位
        first = row.cells[0].text.strip()
        if not (first.isdigit() and int(first) in keep):
            table._tbl.remove(row._tr)
    found = set()
    for row in list(table.rows)[1:]:
        code = int(row.cells[0].text.strip())
        found.add(code)
        for cell, value in zip(row.cells[:2], keep[code]):
            paras = cell.paragraphs
            for pp in paras[1:]:
                pp._p.getparent().remove(pp._p)
            for r in paras[0].runs[1:]:
                r._r.getparent().remove(r._r)
            run = paras[0].runs[0] if paras[0].runs else paras[0].add_run()
            run.text = value
            _font(run, "仿宋", 12, False)
    lost = [str(c) for c in keep if c not in found]
    if lost:
        notes.append(f"立卷申请书模板的表格里没有第 {'、'.join(lost)} 项的行，请手工补上")
    return _save(doc), notes


def make_application(plan: dict, rows: list[tuple[int, str, str]], lawyer: str | None) -> bytes:
    """临时版式：标题、律所、案件编号与申请、表格（编码 | 材料名称）、签字行。"""
    doc = pydocx.Document()
    _para(doc, "立卷申请书", font="黑体", size=16, bold=True, center=True, after=10)
    _para(doc, f"{FIRM}：", before=4)
    _para(doc, f"本人承办的{_case_no(plan)}案件已办结，符合归档条件，现申请立卷归档。归档材料如下：", indent=True)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    for cell, h in zip(t.rows[0].cells, ("编码", "材料名称")):
        _font(cell.paragraphs[0].add_run(h), "黑体", 12, True)
    for _code, no, text in rows:
        r = t.add_row()
        _font(r.cells[0].paragraphs[0].add_run(no), "仿宋", 12, False)
        _font(r.cells[1].paragraphs[0].add_run(text), "仿宋", 12, False)
    _para(doc, f"经办律师（签字）：______________　　（{lawyer or TODO}）", before=18)
    _para(doc, "年　　月　　日", before=6)
    return _save(doc)


# ================================================================ 情况说明

def _templates(md_path: pathlib.Path) -> dict[str, list[str]]:
    """特殊情况说明模板.md 里"## 模板一 / 模板二"下两条 --- 之间的正文，按行。"""
    text = md_path.read_text(encoding="utf-8")
    out: dict[str, list[str]] = {}
    for key in ("模板一", "模板二"):
        sec = text.split(f"## {key}", 1)[1]
        body = sec.split("\n---\n", 2)[1]
        out[key] = body.strip("\n").split("\n")
    return out


def _statement_doc(lines: list[str]) -> bytes:
    doc = pydocx.Document()
    first = True
    for line in lines:
        if not line.strip():
            continue
        if first:
            _para(doc, line.strip(), font="黑体", size=16, bold=True, center=True, after=10)
            first = False
        elif line.startswith(("案件负责人", "年 ")):
            _para(doc, line.strip(), before=12)
        else:
            _para(doc, line.strip(), indent=not line.endswith("：") and not line.startswith(("一、", "二、")))
    return _save(doc)


def make_missing_statement(md_path: pathlib.Path, plan: dict, missing: list[dict]) -> bytes:
    """模板一：案件编号、案件进展（办案结果、结案日期）、缺失材料逐项列出，原因处留【待填写】。"""
    src = _templates(md_path)["模板一"]
    lines: list[str] = []
    skip = False
    for line in src:
        if skip:
            if line.startswith("现本人承诺"):
                skip = False
            else:
                continue
        if line.startswith("【简述"):
            lines.append(f"已办结（办案结果：{plan['result'] or TODO}），结案日期：{plan['close_date'] or TODO}。")
            continue
        if line.startswith("【逐项列出"):
            lines += [f"{i}. 缺失材料：{plain_name(m['name'])}。原因：{TODO}" for i, m in enumerate(missing, 1)]
            lines.append("")
            skip = True
            continue
        lines.append(line.replace("【金助理系统案件编号+客户名称】", _case_no(plan)))
    return _statement_doc(lines)


def make_fee_statement(md_path: pathlib.Path, plan: dict) -> bytes:
    """模板二：只填案件编号；催收方式、未收回原因、佐证材料由律师填（模板里的【…】提示原样保留）。"""
    lines = [line.replace("【金助理系统案件编号+客户名称】", _case_no(plan)) for line in _templates(md_path)["模板二"]]
    return _statement_doc(lines)


# ================================================================ 归档目录.md

def catalog_md(plan: dict, catalog: dict, ranges: dict[int, tuple[int, int]], missing: list[dict],
               skipped: list[str]) -> str:
    planned = {it["code"]: it for it in plan["items"]}
    out = [f"# 归档目录（{catalog['id']}，{catalog['attachment']}）", "",
           f"委托人：{plan['client']}　对方当事人：{plan['opponent'] or '无'}　案由：{plan['cause']}", "",
           "| 编码 | 材料名称 | 必交 | 入卷材料 | 卷宗页码 | 说明 |", "|---|---|---|---|---|---|"]
    got = miss_req = miss_opt = 0
    for it in catalog["items"]:
        need = "*" if it["required"] else "如有"
        if it["code"] in planned:
            got += 1
            out.append(f"| {it['code']} | {plain_name(planned[it['code']]['name'])} | {need} | "
                       f"{'、'.join(planned[it['code']]['materials'])} | {page_label(ranges.get(it['code']))} | |")
        elif it.get("generated"):
            got += 1
            out.append(f"| {it['code']} | {plain_name(it['name'])} | {need} | （程序生成） | "
                       f"{page_label(ranges.get(it['code']))} | |")
        else:
            if it["required"]:
                miss_req += 1
            else:
                miss_opt += 1
            out.append(f"| {it['code']} | {plain_name(it['name'])} | {need} | — | | "
                       f"{'需补充或出具情况说明' if it['required'] else '非必交，可跳过'} |")
    out += ["", f"共 {len(catalog['items'])} 项：已归档 {got} 项，缺失必交 {miss_req} 项，缺失非必交 {miss_opt} 项。"]
    if missing:
        out += ["", "缺失的必交项：" + "；".join(f"{m['code']} {plain_name(m['name'])}" for m in missing)]
    if skipped:
        out += ["", "未能放入卷宗的材料：" + "；".join(skipped)]
    return "\n".join(out) + "\n"

