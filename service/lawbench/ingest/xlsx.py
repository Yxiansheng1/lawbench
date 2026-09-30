"""xlsx：openpyxl 读两遍——data_only=True 取显示值，data_only=False 取公式（Spec 5.2）。

显示值为空而公式不为空（文件没存计算结果）时，先用 LibreOffice 重算后再读显示值。
每个工作表转成带行号、列字母表头的 md 表格；公式单独存一份。
只遍历文件里实际存在的单元格；表格列宽到最右一个有内容的列为止（X4）。
单元格按 Excel 显示的样子写（数字格式见 numfmt.py，N27；X13）。
"""
from __future__ import annotations

import datetime as dt
import pathlib
import zipfile
from typing import Callable

import openpyxl
from openpyxl.utils import get_column_letter

from .. import logs
from . import MAX_SHEET_CELLS, Block, Parsed, ParseError, numfmt
from .detect import is_ole, ole_encrypted
from .links import open_zip


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if v.time() == dt.time() else v.isoformat(sep=" ")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return str(v).replace("\r\n", "\n").replace("|", "\\|").replace("\n", "<br>")


def _display(cell) -> str:
    """单元格按 Excel 显示的样子写（N27，T5 返修 X13）：按数字格式套；常规格式和套不出来的冷门格式写原始值。"""
    shown = numfmt.format(cell.value, cell.number_format)
    if shown is None:
        return _fmt(cell.value)
    return shown.replace("|", "\\|")


def _load(path: pathlib.Path, data_only: bool):
    """用文件对象打开：openpyxl 按文件名打开时会看扩展名，内容是 xlsx、扩展名是 .xls 的会被拒（X1 按文件头分流后会来这里）。"""
    try:
        with open(path, "rb") as f:
            return openpyxl.load_workbook(f, data_only=data_only, read_only=False)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError, TypeError):
        raise ParseError("corrupt")
    except Exception:  # noqa: BLE001 openpyxl 对坏文件抛的异常种类很多
        raise ParseError("corrupt")


def _cells(ws) -> dict:
    """文件里实际存在的单元格 {(行, 列): 单元格}（X4）：不按 max_row × max_column 逐格取——非只读模式下每取一格
    都会新建单元格对象，远处一个带格式的空单元格就能让解析跑几十秒、吃掉几百 MB。"""
    return ws._cells  # noqa: SLF001 openpyxl 没有公开的"只列已有单元格"接口


def _missing_cache(values, formulas) -> bool:
    for ws in formulas.worksheets:
        wv = _cells(values[ws.title])
        for key, c in _cells(ws).items():
            if c.data_type == "f" and (key not in wv or wv[key].value is None):
                return True
    return False


def parse(path: pathlib.Path, recalc: Callable[[pathlib.Path], pathlib.Path] | None = None,
          note: str | None = None, blocked_note: str | None = None) -> Parsed:
    """recalc(path) 用 LibreOffice 重算并返回重算后的 xlsx 路径（位于 工作区/临时/，调用方负责删除）。
    重算不成（深路径、转换程序异常等）时退回"没有缓存值的单元格写公式本身"，不让整份失败（Y7）。"""
    if is_ole(path):
        raise ParseError("encrypted" if ole_encrypted(path) else "corrupt")
    open_zip(path).close()  # 单个部件解压后超过 300 MB：按过大，不交给 openpyxl（X4）
    values = _load(path, True)
    formulas = _load(path, False)
    missing = _missing_cache(values, formulas)
    if recalc is None and missing and blocked_note:
        note = blocked_note  # 有外链、没交给 LibreOffice 重算：没有缓存值的格子写的是公式（契约 1.2 N21）
    if recalc is not None and missing:
        try:
            recalculated = recalc(path)
        except ParseError as e:
            logs.event("materials", "recalc", status="fail", error=e.reason)
        else:
            values.close()
            values = _load(recalculated, True)
    out = Parsed(unit="cell", unit_count=len(formulas.worksheets), count_word="个工作表", note=note)
    flist: list[str] = []
    for ws in formulas.worksheets:
        wv = _cells(values[ws.title])
        fcells = _cells(ws)
        texts: dict[int, dict[int, str]] = {}
        for key in sorted(set(fcells) | set(wv)):
            fc = fcells.get(key)
            # 没有缓存值又没有重算（没有 LibreOffice，或文档有外链、不交给 LibreOffice）：只写公式
            t = _display(wv[key]) if key in wv else ""
            if not t and fc is not None and fc.data_type == "f":
                t = _fmt(fc.value)
            if t:
                texts.setdefault(key[0], {})[key[1]] = t
            if fc is not None and fc.data_type == "f":
                flist.append(f"{ws.title}!{fc.coordinate}\t{fc.value}")
        if texts:
            max_col = max(c for row in texts.values() for c in row)
            if len(texts) * max_col > MAX_SHEET_CELLS:
                raise ParseError("too_large")
            header = "| 行 | " + " | ".join(get_column_letter(c) for c in range(1, max_col + 1)) + " |"
            sep = "|---|" + "---|" * max_col
            body = [f"| {r} | " + " | ".join(texts[r].get(c, "") for c in range(1, max_col + 1)) + " |"
                    for r in sorted(texts)]
            text = "\n".join([header, sep, *body])
        else:
            text = "（空工作表）"
        out.blocks.append(Block(f"表:{ws.title}", text))
    values.close()
    formulas.close()
    out.formulas = "\n".join(flist) if flist else None
    return out
