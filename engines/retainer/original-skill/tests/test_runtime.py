import contextlib
import io
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import generate as g
import runtime_utils as u
from clean_templates import merge_para_runs
from docx import Document
from docx.shared import RGBColor
from lxml import etree


class Regression(unittest.TestCase):
    def test_amounts(self):
        for value, expected in [('2.5万元', '贰万伍仟元整'), ('100000', '壹拾万元整'),
                                ('100001', '壹拾万零壹元整'), ('1,234.56', '壹仟贰佰叁拾肆元伍角陆分'),
                                ('100000001', '壹亿零壹元整'), ('0.05', '零元伍分')]:
            with self.subTest(value=value):
                self.assertEqual(u.parse_fee(value, 'fixed')['upfront'], expected)

    def test_percentages(self):
        for value in ['前期2万+12.5%风险', '前期2万+百分之十二点五', '前期2万+12.5％']:
            self.assertEqual(u.parse_fee(value, 'semi_risk')['risk_rate'], '百分之十二点五（12.5%）')

    def test_ambiguous_fee_rejected(self):
        for value, fee in [('5%', 'semi_risk'), ('前期3万+5万结果', 'fixed'),
                           ('-500元', 'fixed'), ('约定另议', 'fixed'), ('50000+5%', 'fixed'),
                           ('50000+101%', 'semi_risk'), ('50000+风险', 'semi_risk'), ('0.001元', 'fixed'), ('1,23元','fixed')]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                u.parse_fee(value, fee)

    def test_detection(self):
        self.assertEqual(g.detect_case_type('集资诈骗 一审'), 'criminal')
        self.assertEqual(g.detect_case_type('侵犯著作权纠纷 一审'), 'civil')
        self.assertEqual(g.detect_party_type('张行'), 'individual')
        self.assertEqual(g.detect_party_type('某某商行'), 'company')

    def params(self, base, party='individual', fee='fixed', criminal=False):
        return dict(case_type='criminal' if criminal else 'civil', client='张三&A', plaintiff='张三&A',
                    defendant='李四<乙>', crime='集资诈骗', cause='买卖合同纠纷', stage='二审',
                    party_type=party, fee_type=fee, output_base=str(base),
                    legal_rep='王五', legal_rep_position='经理',
                    fee_info=u.parse_fee('25000元+12.5%' if fee == 'semi_risk' else '25000元', fee))

    def test_all_five_template_routes(self):
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            for party, fee, criminal in [('individual','fixed',False), ('individual','semi_risk',False),
                                         ('company','fixed',False), ('company','semi_risk',False),
                                         ('individual','fixed',True)]:
                p = self.params(tmp, party, fee, criminal)
                result = g.generate_documents(p, create_case=True)
                self.assertTrue(result['success'])
                self.assertEqual(len(result['files']), 5 if party == 'company' else 4)
                self.assertFalse(g._verify_output(result['files'], p))
                for name, f in result['files'].items():
                    self.assertNotIn('{{', u.docx_text(f))
                    self.assertEqual(Path(f).read_bytes(), (Path(result['case_dir'])/'01委托手续'/name).read_bytes())
                    with ZipFile(f) as z:
                        for n in z.namelist():
                            if n.endswith(('.xml', '.rels')):
                                etree.fromstring(z.read(n))

    def test_rewrite_does_not_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp, contextlib.redirect_stdout(io.StringIO()):
            p = self.params(tmp)
            first = g.generate_documents(p, True)
            original = {f: Path(f).read_bytes() for f in first['files'].values()}
            p['plaintiff'] = '改名委托人'
            second = g.generate_documents(p, True)
            self.assertNotEqual(first['output_dir'], second['output_dir'])
            self.assertNotEqual(first['case_dir'], second['case_dir'])
            for f, content in original.items():
                self.assertEqual(Path(f).read_bytes(), content)

    def test_failed_validation_is_failure(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(g, '_verify_output', return_value=['broken']), patch.object(g, 'create_case_folder') as archive:
            with self.assertRaisesRegex(ValueError, 'broken'):
                g.generate_documents(self.params(tmp), True)
            archive.assert_not_called()

    def test_cross_run_and_header(self):
        with tempfile.TemporaryDirectory() as tmp:
            src, dst = Path(tmp)/'in.docx', Path(tmp)/'out.docx'
            d = Document()
            p = d.add_paragraph()
            p.add_run('前{{PLA').bold = True
            p.add_run('INTIFF}}后').italic = True
            d.sections[0].header.paragraphs[0].text = '{{UNKNOWN}}'
            d.save(src)
            u.fill_template(src, dst, {'{{PLAINTIFF}}': 'A&B<甲>'})
            new = Document(dst)
            self.assertEqual(new.paragraphs[0].text, '前A&B<甲>后')
            self.assertTrue(new.paragraphs[0].runs[0].bold)
            self.assertTrue(new.paragraphs[0].runs[1].italic)
            self.assertTrue(g._verify_output({'合同.docx': str(dst)}, self.params(tmp)))

    def test_source_suffix_preserved(self):
        d = Document()
        p = d.add_paragraph()
        p.add_run('前张')
        p.add_run('三后')
        g._merge_and_replace_in_para(p, '张三', '{{PLAINTIFF}}')
        self.assertEqual(p.text, '前{{PLAINTIFF}}后')

    def test_cleaner_preserves_format_and_break(self):
        p = Document().add_paragraph()
        p.add_run('a').bold = True
        p.add_run('b').bold = True
        p.add_run('c').italic = True
        p.add_run('d').font.color.rgb = RGBColor(255,0,0)
        p.add_run().add_break()
        merge_para_runs(p)
        self.assertEqual(p.text, 'abcd\n')
        self.assertEqual(len(p.runs), 4)
        self.assertTrue(p.runs[1].italic)
        self.assertEqual(p.runs[2].font.color.rgb, RGBColor(255,0,0))

    def test_no_ready_marker_write(self):
        with patch.object(g, '_templates_complete', return_value=True), patch.object(g, '_mark_templates_ready') as mark:
            g.ensure_templates()
            mark.assert_not_called()

    def test_path_containment(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = u.unique_directory(tmp, '../../x:<bad>')
            self.assertEqual(Path(result).parent, Path(tmp))

    def test_cli_invalid_and_help(self):
        for args, expected in [(['--help'], 0), (['only-one'], 2), ([], 2)]:
            result = subprocess.run([sys.executable, '-B', str(ROOT/'generate.py'), *args], input='', capture_output=True, text=True, encoding='utf-8', timeout=20)
            self.assertEqual(result.returncode, expected)

    def test_version_consistency(self):
        version = json.loads((ROOT/'_meta.json').read_text(encoding='utf-8'))['version']
        self.assertIn('version: '+version, (ROOT/'SKILL.md').read_text(encoding='utf-8'))
        self.assertIn('## ['+version+']', (ROOT/'CHANGELOG.md').read_text(encoding='utf-8'))

    def test_template_manifest_and_years(self):
        manifest = json.loads((ROOT/'template-checksums.json').read_text(encoding='utf-8'))
        actual = {p.relative_to(ROOT).as_posix() for p in ROOT.glob('templates/**/*.docx')}
        self.assertEqual(set(manifest), actual)
        for path, digest in manifest.items():
            self.assertEqual(hashlib.sha256((ROOT/path).read_bytes()).hexdigest(), digest)
            text = u.docx_text(ROOT/path)
            self.assertNotRegex(text, r'202[56]\s*年')
            self.assertNotIn('周海沺', text)

    def test_rebuild_failure_preserves_templates(self):
        with tempfile.TemporaryDirectory() as tmp:
            template = Path(tmp)/'templates'
            template.mkdir()
            sentinel = template/'sentinel.txt'
            sentinel.write_text('original')
            with patch.object(g, 'SKILL_ROOT', tmp), patch.object(g, 'TEMPLATE_DIR', str(template)), patch.object(g, '_rebuild_templates', side_effect=RuntimeError('mock source failure')):
                with self.assertRaises(RuntimeError):
                    g.ensure_templates()
                self.assertEqual(sentinel.read_text(), 'original')
                self.assertEqual(g.TEMPLATE_DIR, str(template))
                self.assertEqual(list(Path(tmp).iterdir()), [template])


if __name__ == '__main__':
    unittest.main(verbosity=2)
