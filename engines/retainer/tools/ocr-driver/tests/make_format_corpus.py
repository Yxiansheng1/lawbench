# -*- coding: utf-8 -*-
"""
生成多格式测试语料（全部为虚构内容）

覆盖：PNG / JPG（两种质量）/ PDF（文本层、扫描件）/ DOCX（文字＋表格＋内嵌图片）
      / 扩展名与内容不符 / 老版 .doc
"""
import io
import os
import shutil
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "format-corpus")
DRIVER = os.path.dirname(HERE)
PKG_ROOT = os.path.dirname(os.path.dirname(DRIVER))
PROBE_IMG = os.path.join(PKG_ROOT, "docs", "ocr-driver-probe", "_tmp_ocr_license.png")

LINES = [
    "营业执照",
    "名　　称　深圳市示例科技有限公司",
    "统一社会信用代码　91440300MA5EXAMPLA",
    "类　　型　有限责任公司",
    "法定代表人　张三",
    "住　　所　广东省深圳市福田区示例路1号",
    "成立日期　2020年09月01日",
]

log = []


def ensure_dirs():
    os.makedirs(OUT, exist_ok=True)


def make_images():
    from PIL import Image
    src = Image.open(PROBE_IMG).convert("RGB")
    png = os.path.join(OUT, "license.png")
    src.save(png)
    jpg = os.path.join(OUT, "license-q90.jpg")
    src.save(jpg, quality=90)
    jpg70 = os.path.join(OUT, "license-q70.jpeg")
    src.save(jpg70, quality=70)
    return png, src


def make_pdf_text(src_img):
    """文本层 PDF：文字直接写入，不需要 OCR"""
    import fitz
    path = os.path.join(OUT, "license-text.pdf")
    font = r"C:\Windows\Fonts\simhei.ttf"
    doc = fitz.open()
    page = doc.new_page(width=595, height=420)
    ok = True
    try:
        page.insert_text((60, 70), LINES[0], fontsize=20, fontfile=font, fontname="hei")
        y = 110
        for t in LINES[1:]:
            page.insert_text((60, y), t.replace("　", " "), fontsize=12, fontfile=font, fontname="hei")
            y += 30
    except Exception as e:
        ok = False
        log.append("text pdf insert failed: " + str(e)[:150])
    doc.save(path)
    doc.close()
    return path, ok


def make_pdf_scan(src_img):
    """扫描件 PDF：整页就是一张图，无文本层，必须 OCR"""
    import fitz
    path = os.path.join(OUT, "license-scan.pdf")
    buf = io.BytesIO()
    src_img.save(buf, format="PNG")
    doc = fitz.open()
    page = doc.new_page(width=595, height=420)
    page.insert_image(fitz.Rect(20, 20, 575, 400), stream=buf.getvalue())
    doc.save(path)
    doc.close()
    return path


def make_docx(src_img):
    """DOCX：段落文字 + 表格 + 一张内嵌图片"""
    from docx import Document
    from docx.shared import Inches
    path = os.path.join(OUT, "license.docx")
    doc = Document()
    doc.add_heading("委托主体资料", level=1)
    doc.add_paragraph("名称：深圳市示例科技有限公司")
    doc.add_paragraph("统一社会信用代码：91440300MA5EXAMPLA")
    doc.add_paragraph("法定代表人：张三")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "住所"
    table.cell(0, 1).text = "广东省深圳市福田区示例路1号"
    table.cell(1, 0).text = "成立日期"
    table.cell(1, 1).text = "2020年09月01日"
    img_path = os.path.join(OUT, "_inline.png")
    src_img.save(img_path)
    doc.add_picture(img_path, width=Inches(4.5))
    doc.save(path)
    os.remove(img_path)
    return path


def make_noext():
    """扩展名缺失但内容是 PNG，用于验证文件头嗅探"""
    path = os.path.join(OUT, "license-noext")
    shutil.copyfile(os.path.join(OUT, "license.png"), path)
    return path


def make_fake_doc():
    """构造最小 OLE 头 + 垃圾数据，用于验证 .doc 的报错是否友好"""
    path = os.path.join(OUT, "license-old.doc")
    with open(path, "wb") as f:
        f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
        f.write(b"\x00" * 512)
    return path


def main():
    ensure_dirs()
    png, src = make_images()
    log.append("png " + png)
    log.append("jpg q90 " + os.path.join(OUT, "license-q90.jpg"))
    log.append("jpeg q70 " + os.path.join(OUT, "license-q70.jpeg"))
    p, ok = make_pdf_text(src)
    log.append("pdf text %s insert_ok=%s" % (p, ok))
    log.append("pdf scan " + make_pdf_scan(src))
    log.append("docx " + make_docx(src))
    log.append("noext " + make_noext())
    log.append("fake doc " + make_fake_doc())

    with open(os.path.join(HERE, "_format_corpus.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(log))
        f.write("\n\n--- docx 内的 media ---\n")
        with zipfile.ZipFile(os.path.join(OUT, "license.docx")) as z:
            for n in z.namelist():
                if n.startswith("word/media/"):
                    f.write("%s  %d B\n" % (n, z.getinfo(n).file_size))


if __name__ == "__main__":
    main()
