# -*- coding: utf-8 -*-
"""
把委托代理合同第六条规整为「单锚点段」，供生成时按收费模式注入。

做法：
  1. 定位首段（含「收费模式」或「本合同律师费约定如下」）
  2. 替换为统一锚点句，保留该段的段落格式与首个 run 的字体
  3. 删除紧随其后的 1./2./（1）（2） 子段

用法（在项目根目录）： python tools/fix-clause6.py
"""
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

TARGETS = [
    "original-skill/templates/个人委托/1 民事委托代理合同.docx",
    "original-skill/templates/公司委托/1 民事委托代理合同.docx",
]
ANCHOR = "六、甲乙双方协商确定，本次法律服务采用以下半风险代理收费模式。"
SUB_RE = re.compile(r"^\s*(?:[1-9][.．、]|（[1-9]）)")


def set_text_keep_format(p, new_text):
    """替换段落文本，保留段落属性与首个 run 的字体属性"""
    if not p.runs:
        p.add_run(new_text)
        return
    p.runs[0].text = new_text
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)


def fix(path):
    from docx import Document
    doc = Document(path)
    paras = list(doc.paragraphs)
    anchor_i = None
    for i, p in enumerate(paras):
        t = p.text or ""
        if "收费模式" in t or "本合同律师费约定如下" in t:
            anchor_i = i
            break
    if anchor_i is None:
        print("  [跳过] 未找到第六条锚点")
        return False

    before = len(paras)
    old = paras[anchor_i].text
    set_text_keep_format(paras[anchor_i], ANCHOR)

    removed = 0
    for p in paras[anchor_i + 1:]:
        t = p.text or ""
        if SUB_RE.match(t):
            p._element.getparent().remove(p._element)
            removed += 1
        elif t.strip():
            break

    doc.save(path)
    after = len(Document(path).paragraphs)
    print("  [完成] 段落 %d → %d，删除子段 %d" % (before, after, removed))
    print("         原第六条首段：%s" % (old.strip()[:60] or "(空)"))
    print("         新锚点：%s" % ANCHOR)
    return True


def main():
    ok = 0
    for rel in TARGETS:
        path = os.path.join(ROOT, rel)
        print("处理：" + rel)
        if not os.path.exists(path):
            print("  [缺失] 文件不存在")
            continue
        if fix(path):
            ok += 1
    print("\n完成 %d / %d 份" % (ok, len(TARGETS)))


if __name__ == "__main__":
    main()
