from docx import Document
d = Document(r"D:\workbuddy测试\retainer-offline\original-skill\templates\公司委托\1 民事委托代理合同.docx")
out = ["段落总数：%d" % len(d.paragraphs), ""]
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t.startswith(("九、", "十、", "十一、", "十二、", "（以下无正文）", "账号", "帐号")):
        out.append("[%02d] %s" % (i, t))
with open(r"D:\workbuddy测试\_company.txt", "w", encoding="utf-8") as f:
    f.write("\n".join(out))
print("ok")