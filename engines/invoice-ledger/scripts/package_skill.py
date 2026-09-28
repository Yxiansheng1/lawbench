"""Publish only declared skill assets. Default excludes personal ledger; --include-ledger for migration;
--blank-ledger ships an empty (header-only) ledger for public distribution."""
import argparse
import io
import os
import subprocess
import sys
import tempfile
import uuid
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent


def _assert_blank(wb):
    """任何工作表都不允许残留数据行（只有表头行可保留）。"""
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(v is not None and str(v).strip() for v in row):
                raise RuntimeError(f"空白台账仍含数据：工作表「{ws.title}」")


def make_blank_ledger(src, dst):
    """由现有台账生成仅含表头的空白台账（对外发布用；绝不改动源文件）。"""
    import openpyxl
    wb = openpyxl.load_workbook(src)
    for ws in wb.worksheets:
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
    wb.save(dst)
    _assert_blank(openpyxl.load_workbook(dst, data_only=True))


def verify_blank_ledger_bytes(blob):
    """对已写入发布包的台账做独立复核。"""
    import openpyxl
    _assert_blank(openpyxl.load_workbook(io.BytesIO(blob), data_only=True))


def assets(include_ledger=False):
    files = [ROOT / n for n in ("SKILL.md", "_meta.json", "CHANGELOG.md",
             "vendor/env.lock", "vendor/env-win_amd64.zip", "LICENSE.md", "THIRD_PARTY.json")]
    for folder, suffixes in (("scripts", {".py", ".md"}), ("references", {".md"})):
        files.extend(p for p in (ROOT / folder).rglob("*") if p.is_file()
                     and p.suffix in suffixes and "__pycache__" not in p.parts)
    if include_ledger:
        files.append(ROOT / "发票主台账.xlsx")
        for folder in ("_原票", "_提取记录", "_审计", "_人工核验", "_报销批次", "_打印包"):
            if (ROOT / folder).exists():
                files.extend(p for p in (ROOT / folder).rglob("*") if p.is_file())
    for p in files:
        if not p.is_file() or p.is_symlink() or ROOT not in p.resolve().parents:
            raise RuntimeError(f"发布资产缺失或越界：{p}")
    return sorted(set(files))


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", required=True)
    ap.add_argument("--include-ledger", action="store_true")
    ap.add_argument("--blank-ledger", action="store_true",
                    help="随包提供空白台账（仅表头、无任何发票数据），供对外发布")
    a = ap.parse_args()
    if a.include_ledger and a.blank_ledger:
        ap.error("--include-ledger 与 --blank-ledger 互斥")
    out = Path(a.out).resolve()
    if out == ROOT or ROOT in out.parents:
        ap.error("发布包必须写到技能目录外")
    if out.exists():
        ap.error("输出已存在，请使用新文件名")
    r = subprocess.run([sys.executable, "-B", str(ROOT / "scripts/verify_consistency.py")])
    if r.returncode:
        return r.returncode
    from runtime_cache import read_lock, verify_archive
    verify_archive(read_lock())
    files = assets(a.include_ledger)
    extra = []  # (临时文件, 包内路径)
    if a.blank_ledger:
        tmp_ledger = ROOT / (".blank-ledger-" + uuid.uuid4().hex + ".xlsx")
        make_blank_ledger(ROOT / "发票主台账.xlsx", tmp_ledger)
        extra.append((tmp_ledger, "发票主台账.xlsx"))
    out.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=".package-", suffix=".zip", dir=out.parent)
    os.close(fd)
    temp = Path(name)
    try:
        expected = {p.relative_to(ROOT).as_posix() for p in files} | {arc for _, arc in extra}
        if any("private" in n.lower() or n.endswith((".dpapi", ".pem")) or n == "license.json" for n in expected):
            raise RuntimeError("发布包禁止包含私钥或设备许可证")
        with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as z:
            for p in files:
                z.write(p, p.relative_to(ROOT).as_posix())
            for p, arc in extra:
                z.write(p, arc)
        with zipfile.ZipFile(temp) as z:
            if set(z.namelist()) != expected or z.testzip() is not None:
                raise RuntimeError("发布包清单或 CRC 校验失败")
            forbidden = ("vendor/env/", "vendor/env.tmp/", "_日志/", "_备份/", "_候选台账/")
            if any(n.startswith(forbidden) or "__pycache__" in n for n in z.namelist()):
                raise RuntimeError("发布包混入运行产物")
            if a.blank_ledger:
                verify_blank_ledger_bytes(z.read("发票主台账.xlsx"))
        os.replace(temp, out)
        print(f"[OK] {out} ({out.stat().st_size:,} bytes; {len(files) + len(extra)} files)")
        return 0
    finally:
        if temp.exists():
            temp.unlink()
        for p, _ in extra:
            if p.exists():
                p.unlink()


if __name__ == "__main__":
    sys.exit(main())
