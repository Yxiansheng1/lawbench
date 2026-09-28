"""Real offline Chinese image OCR and isolated ledger regression. Run through invoke.py."""
import builtins
import hashlib
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
from PIL import Image, ImageDraw, ImageFont
import extract_fields as fields

SK = Path(__file__).resolve().parents[2]
OCR_BUYER = "示例科技有限公司"   # 合成购买方；验收脚本自带配置，不依赖本机 buyer.json


def main():
    os.environ["INVOICE_BUYER"] = OCR_BUYER
    work = Path(tempfile.mkdtemp(prefix="invoice-ocr-test-"))
    ledger_hash = hashlib.sha256((SK / "发票主台账.xlsx").read_bytes()).hexdigest()
    checks = []

    def check(name, value, detail=""):
        print(("PASS " if value else "FAIL ") + name + (": " + str(detail)[-400:] if not value else ""), flush=True)
        checks.append(bool(value))

    status = fields.ocr_capabilities()
    print("OCR languages:", status)
    if not status["chinese_available"]:
        print("NOT RUN: target Windows lacks Chinese OCR; this is not a passing image test")
        return 1
    font_path = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/msyh.ttc"
    image = Image.new("RGB", (1500, 820), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path), 46)
    for i, line in enumerate(["电子发票 测试样本", "发票号码：12345678901234567890",
                              "开票日期：2026年09月18日", "价税合计：￥123.45", "项目：办公用品复印纸", "购买方名称：" + OCR_BUYER]):
        draw.text((60, 50+i*120), line, font=font, fill="black")
    sample = work / "中文图片.png"
    image.save(sample)
    result = fields.extract_from_image(str(sample))
    check("real PNG number date amount", result.get("发票号码全号") == "12345678901234567890"
          and result.get("开票日期") == "2026-09-18" and result.get("价税合计（元）") == "123.45", result)
    jpeg = work / "sample.jpg"; image.save(jpeg, quality=95)
    r = fields.extract_from_image(str(jpeg))
    check("real JPEG OCR", r.get("发票号码全号") == "12345678901234567890" and r.get("价税合计（元）") == "123.45", r)
    blank = work / "blank.png"; Image.new("RGB", (800, 500), "white").save(blank)
    check("blank image is content-empty", fields.extract_from_image(str(blank)).get("_ocr_empty"))
    broken = work / "broken.png"; broken.write_bytes(b"not an image")
    check("corrupt image is execution error", bool(fields.extract_from_image(str(broken)).get("_ocr_error")))
    with patch("winsdk.windows.media.ocr.OcrEngine") as engine:
        engine.try_create_from_language.return_value = None
        r = fields.extract_from_image(str(sample))
    check("missing OS language is distinct from blank image", r.get("_ocr_language_missing") and not r.get("_ocr_empty"))
    real_import = builtins.__import__

    def without_winsdk(name, *args, **kwargs):
        if name.startswith("winsdk"):
            raise ImportError("simulated missing bundled winsdk")
        return real_import(name, *args, **kwargs)

    missing = False
    with patch("builtins.__import__", side_effect=without_winsdk):
        try:
            fields.extract_from_image(str(sample))
        except fields.DependencyMissing as e:
            missing = e.name == "winsdk"
    check("missing bundled binding is fatal dependency error", missing)
    text = "开 票 日 期：2026 年 09 月 1 8 日\n价 税 合 计：¥ 123 · 45"
    parsed = fields.parse_text_fields(fields.normalize_ocr_text(text), "test.png")
    check("OCR formatting normalization", parsed["开票日期"] == "2026-09-18" and parsed["价税合计（元）"] == "123.45")
    check("normalization does not guess letters as digits", fields.normalize_ocr_text("金额 ¥ I23.O5") == "金额 ¥ I23.O5")
    dbdir = work / "ledger"; dbdir.mkdir()
    ledger = dbdir / "发票主台账.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "发票主台账"
    ws.append(["类别序号", "发票类别", "开票日期", "价税合计（元）", "发票号码后五位",
               "文件名", "来源批次", "发票号码全号", "状态"])
    wb.save(ledger); wb.close()

    def cli(*args):
        return subprocess.run([sys.executable, "-B", str(SK / "scripts/invoice_db.py"), *args],
                              capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=45)

    first = cli("import", "--src", str(sample), "--img", "--batch", "OCR测试", "--ledger", str(dbdir))
    wb = openpyxl.load_workbook(ledger, read_only=True); rows = list(wb["发票主台账"].values); wb.close()
    check("image CLI import retains manual-review state", first.returncode in (0, 2) and len(rows) == 2
          and rows[1][7] == "12345678901234567890" and rows[1][8] == "⚠OCR待人工", first.stdout+first.stderr)
    duplicate = cli("import", "--src", str(sample), "--img", "--batch", "OCR测试", "--ledger", str(dbdir))
    wb = openpyxl.load_workbook(ledger, read_only=True); row_count = wb["发票主台账"].max_row; wb.close()
    check("OCR duplicate does not insert second row", duplicate.returncode == 2 and row_count == 2, duplicate.stdout+duplicate.stderr)
    mark = cli("mark-reimbursed", "--batch", "OCR测试", "--apply", "--ledger", str(dbdir))
    wb = openpyxl.load_workbook(ledger, read_only=True); state = wb["发票主台账"].cell(2,9).value; wb.close()
    check("OCR rows cannot be automatically marked reimbursed", mark.returncode == 0 and state == "⚠OCR待人工", mark.stdout+mark.stderr)
    check("real ledger SHA256 unchanged", ledger_hash == hashlib.sha256((SK / "发票主台账.xlsx").read_bytes()).hexdigest())
    print(json.dumps({"passed": sum(checks), "total": len(checks), "artifacts": str(work)}, ensure_ascii=False))
    return 0 if all(checks) else 1


if __name__ == "__main__":
    sys.exit(main())
