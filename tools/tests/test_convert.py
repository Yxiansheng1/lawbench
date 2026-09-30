"""格式互转（Spec 13.2）：每种转换一个样本成功、原文件 sha256 不变；找不到程序时给中文提示。"""
from __future__ import annotations

import os
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest

from conftest import FIXTURES, sha256
from convert import core, finder

SOFFICE = finder.find_soffice()
PANDOC = finder.find_pandoc()
need_lo = pytest.mark.skipif(SOFFICE is None, reason="本机没有 LibreOffice")
need_pandoc = pytest.mark.skipif(PANDOC is None, reason="本机没有 pandoc")


@pytest.fixture(scope="module")
def samples(tmp_path_factory):
    """把 tests\\fixtures 的样本复制出来；doc、wps、xls 由 LibreOffice 从 docx / xlsx 转出。"""
    d = tmp_path_factory.mktemp("convert-src")
    m = {
        "docx": FIXTURES / "civil-01" / "借条.docx",
        "md": FIXTURES / "civil-01" / "案情摘要.md",
        "pdf": FIXTURES / "criminal-01" / "起诉意见书.pdf",
        "scan": FIXTURES / "criminal-01" / "讯问笔录.pdf",
        "jpg": FIXTURES / "criminal-01" / "转账截图.jpg",
        "xlsx": FIXTURES / "civil-01" / "银行流水.xlsx",
    }
    out = {}
    for k, p in m.items():
        out[k] = d / p.name
        shutil.copyfile(p, out[k])
    if SOFFICE:
        for src, fmt in ((out["docx"], "doc"), (out["xlsx"], "xls")):
            work = tmp_path_factory.mktemp("legacy")     # 用产品自己的调用（配置目录在短路径下）
            shutil.copyfile(core._libreoffice(src, fmt, work), d / f"{src.stem}.{fmt}")
        out["doc"] = d / "借条.doc"
        out["xls"] = d / "银行流水.xls"
        # WPS 文字的 .wps 与 .doc 同为 OLE 复合文档，LibreOffice 按内容识别；这里用 .doc 改扩展名代替
        out["wps"] = d / "借条-wps.wps"
        shutil.copyfile(out["doc"], out["wps"])
    return out


def check_unchanged_and_clean(src, before):
    import tempfile
    assert sha256(src) == before
    # 源文件所在目录没有锁文件、临时文件
    extras = [p.name for p in src.parent.iterdir() if p.name.startswith((".~lock", "~$"))]
    assert extras == []
    # 系统临时目录里没有本工具留下的工作目录（复核 P2-1：原先查错了目录）
    left = [p.name for p in Path(tempfile.gettempdir()).glob(core.TEMP_PREFIX + "*")]
    assert left == [], left


@pytest.fixture(autouse=True)
def isolated_temp(monkeypatch, tmp_path_factory):
    """每个测试用自己的系统临时目录，才能断言"没有留下 lawbench-convert-*"。
    建在真实 %TEMP% 下的短路径里（与正式环境一致），不放在可能很深的 tmp_path 里。"""
    import tempfile
    tmp_path_factory.getbasetemp()      # 先让 pytest 定下自己的临时根目录，免得它建到下面这个会被删掉的目录里
    real = tempfile.gettempdir()
    t = Path(tempfile.mkdtemp(prefix="lbt-", dir=real))
    monkeypatch.setattr(tempfile, "tempdir", str(t))
    yield t
    monkeypatch.setattr(tempfile, "tempdir", real)
    shutil.rmtree(core.LONG_PREFIX + str(t) if os.name == "nt" else t, ignore_errors=True)


CASES = [
    pytest.param("doc2docx", "doc", marks=need_lo, id="doc→docx"),
    pytest.param("doc2docx", "wps", marks=need_lo, id="wps→docx"),
    pytest.param("xls2xlsx", "xls", marks=need_lo, id="xls→xlsx"),
    pytest.param("word2pdf", "docx", marks=need_lo, id="docx→pdf"),
    pytest.param("word2pdf", "doc", marks=need_lo, id="doc→pdf"),
    pytest.param("pdf2docx", "pdf", marks=need_pandoc, id="pdf→docx"),
    pytest.param("md2docx", "md", marks=need_pandoc, id="md→docx"),
    pytest.param("docx2md", "docx", marks=need_pandoc, id="docx→md"),
    pytest.param("img2pdf", "jpg", id="jpg→pdf"),
]


@pytest.mark.parametrize("kind,sample", CASES)
def test_each_conversion(samples, kind, sample, monkeypatch):
    src = samples[sample]
    if sample in ("doc", "wps"):                         # 第一版不转换旧版 Word / WPS（N24 ②）
        before = sha256(src)
        with pytest.raises(core.ConvertError, match="暂不支持旧版 Word / WPS 文件"):
            core.convert_file(kind, src)
        check_unchanged_and_clean(src, before)
        return
    before = sha256(src)
    out = core.convert_file(kind, src)
    assert out.is_file() and out.parent == src.parent / core.OUT_DIR_NAME
    assert out.suffix == core.BY_KEY[kind].output
    assert_valid_output(out)
    check_unchanged_and_clean(src, before)


def assert_valid_output(out):
    if out.suffix == ".docx":
        from docx import Document
        text = "".join(p.text for p in Document(out).paragraphs)
        assert len(text) > 10
    elif out.suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(out)
        assert "流水" in wb.sheetnames
    elif out.suffix == ".pdf":
        import pypdfium2 as pdfium
        assert len(pdfium.PdfDocument(str(out))) >= 1
    elif out.suffix == ".md":
        assert "借" in out.read_text(encoding="utf-8")
    with open(out, "rb") as f:
        head = f.read(4)
    if out.suffix in (".docx", ".xlsx"):
        assert zipfile.is_zipfile(out) and head.startswith(b"PK")


@need_pandoc
def test_pdf_to_word_keeps_text_and_paragraphs(samples):
    from docx import Document
    out = core.convert_file("pdf2docx", samples["pdf"])
    paras = [p.text for p in Document(out).paragraphs if p.text.strip()]
    joined = "".join(paras)
    assert "虚公刑诉字〔2026〕417号" in joined
    assert "骗取他人财物共计人民币126,500元" in joined.replace("\n", "")
    assert len(paras) >= 10                         # 按段落组织，不是整页一段


@need_pandoc
def test_scanned_pdf_refused(samples):
    before = sha256(samples["scan"])
    with pytest.raises(core.ConvertError, match="扫描件"):
        core.convert_file("pdf2docx", samples["scan"])
    assert sha256(samples["scan"]) == before


def test_same_name_gets_suffix(samples, tmp_path):
    src = tmp_path / samples["jpg"].name
    shutil.copyfile(samples["jpg"], src)
    a = core.convert_file("img2pdf", src)
    b = core.convert_file("img2pdf", src)
    assert a.name == f"{src.stem}.pdf" and b.name == f"{src.stem}(2).pdf"


def test_wrong_extension_and_batch(samples, tmp_path):
    with pytest.raises(core.ConvertError, match="不是这种转换能处理的格式"):
        core.convert_file("xls2xlsx", samples["md"])
    res = core.convert_many("img2pdf", [samples["jpg"], tmp_path / "不存在.png"])
    assert res[0][1].suffix == ".pdf" and "不存在" in res[1][1]


def test_missing_programs_chinese_message(samples, monkeypatch):
    monkeypatch.setattr(finder, "find_soffice", lambda: None)
    monkeypatch.setattr(finder, "find_pandoc", lambda: None)
    with pytest.raises(core.ConvertError, match="没有找到 LibreOffice"):
        core.convert_file("word2pdf", samples["docx"])
    with pytest.raises(core.ConvertError, match="没有找到 pandoc"):
        core.convert_file("md2docx", samples["md"])
    assert not (samples["md"].parent / core.OUT_DIR_NAME / "案情摘要(9).docx").exists()


def _alive(pid: int) -> bool:
    r = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/NH"], capture_output=True, text=True,
                       encoding="mbcs", errors="replace")
    return str(pid) in r.stdout


def test_hung_converter_killed_with_children(samples, tmp_path, monkeypatch):
    """转换程序卡住（例如弹出"等待打印机连接"）：到超时结束它和它启动的子进程，给中文原因，原文件不动。"""
    import sys
    import time
    pidfile = tmp_path / "pids.txt"
    script = tmp_path / "hang.py"
    script.write_text(
        "import os, subprocess, sys, time\n"
        "child = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(600)'])\n"
        f"open(r'{pidfile}', 'w').write(f'{{os.getpid()}} {{child.pid}}')\n"
        "time.sleep(600)\n", encoding="utf-8")
    fake = tmp_path / "soffice.cmd"
    fake.write_text(f'@"{sys.executable}" "{script}" %*\n', encoding="mbcs")
    monkeypatch.setattr(core, "_soffice", lambda: fake)
    monkeypatch.setattr(core, "TIMEOUT_S", 3)
    before = sha256(samples["docx"])
    t0 = time.time()
    with pytest.raises(core.ConvertError, match="转换超时"):
        core.convert_file("word2pdf", samples["docx"])
    assert time.time() - t0 < 30
    pids = [int(x) for x in pidfile.read_text().split()]
    time.sleep(0.5)
    assert not any(_alive(p) for p in pids), "卡住的转换程序或其子进程还在"
    check_unchanged_and_clean(samples["docx"], before)


def test_no_print_calls_in_converter():
    src = (core.__file__)
    text = open(src, encoding="utf-8").read()
    for bad in ("PrintOut", "print_to", "-p ", "--print-to-file", "--pt ", "SetDefaultPrinter"):
        assert bad not in text


def test_finder_order(isolated_temp, monkeypatch):
    tmp_path = isolated_temp
    fake_client = tmp_path / "local"
    exe = fake_client / "Programs" / "lawbench" / "resources" / "libreoffice" / "program" / "soffice.exe"
    exe.parent.mkdir(parents=True)
    exe.write_bytes(b"")
    pan = fake_client / "Programs" / "lawbench" / "resources" / "pandoc" / "pandoc.exe"
    pan.parent.mkdir(parents=True)
    pan.write_bytes(b"")
    monkeypatch.setenv("LOCALAPPDATA", str(fake_client))
    monkeypatch.delenv("LAWBENCH_SOFFICE", raising=False)
    monkeypatch.delenv("LAWBENCH_PANDOC", raising=False)
    assert finder.find_soffice() == exe                 # 客户端内置优先于系统安装
    assert finder.find_pandoc() == pan
    override = tmp_path / "my-soffice.exe"
    override.write_bytes(b"")
    monkeypatch.setenv("LAWBENCH_SOFFICE", str(override))
    assert finder.find_soffice() == override            # 环境变量最先


def test_gui_builds():
    import tkinter as tk
    from convert.app import ConvertApp
    try:
        root = tk.Tk()
    except tk.TclError:
        pytest.skip("没有图形界面环境")
    root.withdraw()
    app = ConvertApp(root)
    app.kind.set(core.BY_KEY["pdf2docx"].label)
    app._update_note()
    assert "只保留文字和段落" in app.note["text"]
    root.destroy()


# ---------------------------------------------------------------- 不联网（复核 P1-1）

@pytest.fixture
def listener():
    """本机 127.0.0.1 随机端口的 HTTP 监听，记录收到的请求数。"""
    import http.server
    import threading
    hits = []

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            hits.append(self.path)
            self.send_response(200)
            self.send_header("Content-Type", "image/png")
            self.end_headers()
            self.wfile.write(_png())

        def log_message(self, *a):
            pass

    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    t = threading.Thread(target=srv.serve_forever, daemon=True)
    t.start()
    yield srv.server_address[1], hits
    srv.shutdown()


def _png() -> bytes:
    import io
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 0, 0)).save(buf, "PNG")
    return buf.getvalue()


def _docx_with_external_image(path, url):
    """docx 里一张以外部链接引用（r:link、TargetMode=External）的图片。"""
    import io
    import re
    from docx import Document
    doc = Document()
    doc.add_paragraph("外部链接图片测试（虚构）")
    doc.add_picture(io.BytesIO(_png()))
    tmp = path.with_suffix(".tmp.docx")
    doc.save(tmp)
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(b'r:embed="', b'r:link="')
            elif item.filename == "word/_rels/document.xml.rels":
                data = re.sub(rb'Target="media/image1\.png"', f'Target="{url}" TargetMode="External"'.encode(), data)
            zout.writestr(item, data)
    tmp.unlink()


@need_lo
@need_pandoc
def test_conversions_do_not_fetch_remote_resources(tmp_path, listener):
    port, hits = listener
    md = tmp_path / "远程图片.md"
    md.write_text(f"# 测试\n\n![图](http://127.0.0.1:{port}/md.png)\n\n正文（虚构）\n", encoding="utf-8")
    dx = tmp_path / "外链图片.docx"
    _docx_with_external_image(dx, f"http://127.0.0.1:{port}/docx.png")
    core.convert_file("md2docx", md)
    core.convert_file("word2pdf", dx)
    core.convert_file("docx2md", dx)
    assert hits == [], f"转换时访问了网络：{hits}"


@need_lo
def test_doc_with_external_image_refused_without_fetch(tmp_path, listener, monkeypatch):
    """带外链图片的 .doc / .wps：小工具不交给转换程序，监听 0 次请求，原文件不变。留档的 extlinks 仍能认出外链。"""
    from convert import extlinks
    port, hits = listener
    dx = tmp_path / "外链图片.docx"
    _docx_with_external_image(dx, f"http://127.0.0.1:{port}/docx.png")
    doc = _as_doc(dx, tmp_path)
    wps = tmp_path / "外链图片.wps"
    shutil.copyfile(doc, wps)
    before = sha256(doc)
    called = []
    monkeypatch.setattr(core, "_libreoffice", lambda *a, **k: called.append(a))
    for kind, src in (("doc2docx", doc), ("word2pdf", doc), ("doc2docx", wps)):
        with pytest.raises(core.ConvertError, match="暂不支持旧版 Word / WPS 文件"):
            core.convert_file(kind, src)
    assert called == [] and hits == [] and sha256(doc) == before
    assert extlinks.has_external_links(doc)              # 留档模块（工作台服务 T5 同一思路）


@need_lo
def test_extlinks_ignores_hyperlink_and_plain_url(tmp_path):
    """留档的 extlinks：普通超链接和正文网址文字不算外链。"""
    from convert import extlinks
    from docx import Document
    dx = tmp_path / "超链接.docx"
    shutil.copyfile(FIXTURES / "contract-01" / "采购合同.docx", dx)       # 第 39 段有超链接
    d = Document(dx)
    d.add_paragraph("参考网址：http://127.0.0.1:9/page（纯文字）")
    d.save(dx)
    assert not extlinks.has_external_links(_as_doc(dx, tmp_path))


def test_disguised_rtf_or_html_named_doc_refused(tmp_path):
    """扩展名是 .doc、内容其实是 RTF / 网页：不是 docx，一律不交给转换程序。"""
    rtf = tmp_path / "伪装.doc"
    rtf.write_bytes(b'{\\rtf1 {\\field{\\*\\fldinst INCLUDEPICTURE "http://127.0.0.1:9/x.png"}}}')
    html = tmp_path / "网页.doc"
    html.write_text('<html><body><img src="https://example.invalid/a.png"></body></html>', encoding="utf-8")
    for f in (rtf, html):
        with pytest.raises(core.ConvertError, match="暂不支持旧版 Word / WPS 文件"):
            core.convert_file("doc2docx", f)


@need_lo
def test_xls_external_references_do_not_fetch(tmp_path, listener):
    """.xls 里的外部工作簿引用、WEBSERVICE、HYPERLINK：转换时不发请求（Calc 链接不更新）。"""
    from openpyxl import Workbook
    port, hits = listener
    wb = Workbook()
    ws = wb.active
    ws["A1"] = f"='http://127.0.0.1:{port}/[ext.xlsx]Sheet1'!A1"
    ws["A2"] = f'=WEBSERVICE("http://127.0.0.1:{port}/ws")'
    ws["A3"] = f'=HYPERLINK("http://127.0.0.1:{port}/hl","链接")'
    ws["A4"] = "虚构数据"
    x = tmp_path / "外部引用.xlsx"
    wb.save(x)
    work = tmp_path / "mkxls"
    work.mkdir()
    xls = tmp_path / "外部引用.xls"
    shutil.copyfile(core._libreoffice(x, "xls", work), xls)
    out = core.convert_file("xls2xlsx", xls)
    assert out.is_file()
    assert hits == [], f"转换时访问了网络：{hits}"


def _as_doc(dx, tmp_path):
    """用 LibreOffice 把带外链图片的 docx 另存为 doc（配置同样禁止取外链）。"""
    work = tmp_path / "mkdoc"
    work.mkdir()
    out = core._libreoffice(dx, "doc", work)
    dst = tmp_path / "外链图片.doc"
    shutil.copyfile(out, dst)
    return dst


# ---------------------------------------------------------------- 复核 P1-2、P1-3、P3-5、(e)

def test_batch_survives_unexpected_errors(samples, tmp_path, monkeypatch):
    """批量中间夹一个文件夹和一个超大图：各自给中文原因，前后文件照常转换。"""
    from PIL import Image
    folder = tmp_path / "只读文件夹.png"
    folder.mkdir()
    big = tmp_path / "超大图.png"
    Image.new("RGB", (3000, 3000)).save(big)
    monkeypatch.setattr(Image, "MAX_IMAGE_PIXELS", 2_000_000)   # 900 万像素 > 2 倍上限，触发"图片过大"
    res = core.convert_many("img2pdf", [samples["jpg"], folder, big, samples["jpg"]])
    assert isinstance(res[0][1], Path) and isinstance(res[3][1], Path)
    for _, r, _ in res[1:3]:
        assert isinstance(r, str) and not any(w in r for w in ("Error", "Traceback", "Exception"))


def test_unexpected_exception_becomes_chinese_reason(samples, monkeypatch):
    def boom(kind, src):
        raise RuntimeError("内部细节 xyz")
    monkeypatch.setattr(core, "convert_file_ex", boom)
    res = core.convert_many("img2pdf", [samples["jpg"]])
    assert res[0][1] == core.INTERNAL and "xyz" not in res[0][1]


def test_failure_and_timeout_leave_no_workdir(samples, tmp_path, monkeypatch):
    bad = tmp_path / "坏.pdf"
    bad.write_bytes(b"%PDF-1.4 broken")
    with pytest.raises(core.ConvertError):
        core.convert_file("pdf2docx", bad)
    check_unchanged_and_clean(bad, sha256(bad))
    import sys
    fake = tmp_path / "soffice.cmd"
    fake.write_text(f'@"{sys.executable}" -c "import time; time.sleep(60)" %*\n', encoding="mbcs")
    monkeypatch.setattr(core, "_soffice", lambda: fake)
    monkeypatch.setattr(core, "TIMEOUT_S", 2)
    with pytest.raises(core.ConvertError, match="转换超时"):
        core.convert_file("word2pdf", samples["docx"])
    check_unchanged_and_clean(samples["docx"], sha256(samples["docx"]))


def test_cleanup_stale_only_old_own_dirs(isolated_temp, tmp_path, monkeypatch):
    import os
    import time
    monkeypatch.setattr(core, "profile_root", lambda: tmp_path / "lo")
    old = isolated_temp / (core.TEMP_PREFIX + "old")
    old.mkdir()
    (old / "src.docx").write_bytes(b"x")
    fresh = isolated_temp / (core.TEMP_PREFIX + "fresh")
    fresh.mkdir()
    other = isolated_temp / "other-old"
    other.mkdir()
    afile = isolated_temp / (core.TEMP_PREFIX + "file")
    afile.write_bytes(b"x")
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "keep.txt").write_bytes(b"x")
    link = isolated_temp / (core.TEMP_PREFIX + "junction")
    subprocess.run(["cmd", "/c", "mklink", "/J", str(link), str(outside)], check=True, capture_output=True)
    past = time.time() - 3600
    for p in (old, other, afile):
        os.utime(p, (past, past))
    assert core.cleanup_stale() == 1
    assert not old.exists() and fresh.exists() and other.exists() and afile.exists()
    assert link.exists() and (outside / "keep.txt").exists()
    os.rmdir(link)


@need_pandoc
def test_mixed_pdf_reports_scan_pages(samples, tmp_path):
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for f in (samples["pdf"], samples["scan"]):
        for pg in PdfReader(f).pages[:2]:
            w.add_page(pg)
    mixed = tmp_path / "混合.pdf"
    with open(mixed, "wb") as fh:
        w.write(fh)
    res = core.convert_many("pdf2docx", [mixed])
    _, out, notes = res[0]
    assert isinstance(out, Path)
    assert notes == ["第 3 页是扫描页，没有转出文字", "第 4 页是扫描页，没有转出文字"]


def test_gui_notes_and_close_guard(monkeypatch):
    import tkinter as tk
    from convert.app import ConvertApp
    try:
        root = tk.Tk()
    except tk.TclError:
        pytest.skip("没有图形界面环境")
    root.withdraw()
    app = ConvertApp(root)
    for key in ("md2docx", "docx2md"):
        app.kind.set(core.BY_KEY[key].label)
        app._update_note()
        assert "不保留图片" in app.note["text"]
    shown = []
    monkeypatch.setattr("convert.app.messagebox.showinfo", lambda *a: shown.append(a))
    app.busy = True
    app.on_close()
    assert shown and root.winfo_exists()        # 转换中不关窗
    app.busy = False
    app.on_close()


# ---------------------------------------------------------------- 第二轮返修：域指令范围、权限与损坏、LibreOffice 短路径


def _rewrite_stream(path, name, old: bytes, new: bytes):
    """按原长度改写 OLE 流里的字节（照复核员的做法）。"""
    import olefile
    assert len(old) == len(new)
    with olefile.OleFileIO(str(path), write_mode=True) as ole:
        data = ole.openstream(name).read()
        assert old in data, (name, old)
        ole.write_stream(name, data.replace(old, new))


def _linked_doc_variants(tmp_path, port):
    """底子：LibreOffice 生成的带 INCLUDEPICTURE 外链的 .doc。返回 (开关在前且 Data 无地址, 域代码原样且 Data 无地址)。"""
    url = f"http://127.0.0.1:{port}/docx.png"
    dx = tmp_path / "外链图片.docx"
    _docx_with_external_image(dx, url)
    base = _as_doc(dx, tmp_path)
    garble = (url.encode("ascii"), ("x" + url[1:]).encode("ascii"))      # Data 流里的单字节地址改掉
    v1 = tmp_path / "开关在前.doc"
    shutil.copyfile(base, v1)
    old = f' INCLUDEPICTURE  "{url}" \\d'.encode("utf-16-le")
    new = f' INCLUDEPICTURE \\d "{url}" '.encode("utf-16-le")
    _rewrite_stream(v1, "WordDocument", old, new)
    _rewrite_stream(v1, "Data", *garble)
    v2 = tmp_path / "只靠域代码.doc"
    shutil.copyfile(base, v2)
    _rewrite_stream(v2, "Data", *garble)
    return v1, v2


@need_lo
def test_extlinks_field_check_switch_first_and_no_data_address(tmp_path):
    """留档的 extlinks：①开关在前、Data 流无地址；②域代码原样、Data 流无地址——都认得出外链。"""
    from convert import extlinks
    v1, v2 = _linked_doc_variants(tmp_path, 9)
    assert extlinks.has_external_links(v1) and extlinks.has_external_links(v2)


@need_lo
def test_extlinks_plain_english_link_text(tmp_path):
    """留档的 extlinks：英文正文 "see the link http://…" 不在域指令里，不算外链。"""
    from convert import extlinks
    from docx import Document
    dx = tmp_path / "英文正文.docx"
    d = Document()
    d.add_paragraph("Please see the link http://127.0.0.1:9/page for details.")
    d.add_paragraph("To import https://127.0.0.1:9/data use the tool.")
    d.save(dx)
    assert not extlinks.has_external_links(_as_doc(dx, tmp_path))


def test_readonly_source_dir_reason(samples, tmp_path):
    ro = tmp_path / "只读目录"
    ro.mkdir()
    src = ro / "转账截图.jpg"
    shutil.copyfile(samples["jpg"], src)
    subprocess.run(["icacls", str(ro), "/deny", "*S-1-1-0:(AD)"], check=True, capture_output=True)  # 不许建子文件夹
    try:
        res = core.convert_many("img2pdf", [src])
    finally:
        subprocess.run(["icacls", str(ro), "/remove:d", "*S-1-1-0"], capture_output=True)
    assert "没有写入权限" in res[0][1]


@need_lo
def test_corrupt_doc_refused_as_legacy(tmp_path):
    """结构损坏的 .doc（OLE 文件头还在）：同样按旧版文件不转换，给同一句提示，不报"内部错误"。"""
    from docx import Document
    dx = tmp_path / "a.docx"
    Document().save(dx)
    doc = _as_doc(dx, tmp_path)
    bad = tmp_path / "结构损坏.doc"
    bad.write_bytes(doc.read_bytes()[:1536] + b"\xff" * 512)
    res = core.convert_many("doc2docx", [bad])
    assert "暂不支持旧版 Word / WPS 文件" in res[0][1]


def test_profile_path_too_long_refused(samples, tmp_path, monkeypatch):
    deep = tmp_path / ("很长的目录名" * 20)
    monkeypatch.setattr(core, "profile_root", lambda: deep)
    called = []
    monkeypatch.setattr(core, "_run", lambda *a, **k: called.append(a))
    with pytest.raises(core.ConvertError, match="路径太长"):
        core.convert_file("word2pdf", samples["docx"])
    assert called == []


def test_soffice_crash_code_message(samples, tmp_path, monkeypatch):
    fake = tmp_path / "soffice.cmd"
    fake.write_text("@exit /b -1073740791\n", encoding="mbcs")                 # 0xC0000409
    monkeypatch.setattr(core, "_soffice", lambda: fake)
    with pytest.raises(core.ConvertError, match="转换程序异常退出"):
        core.convert_file("word2pdf", samples["docx"])


@need_lo
def test_profile_dir_short_and_removed(samples, isolated_temp, monkeypatch, short_profile_root):
    """配置目录在 <用户数据目录>/lawbench/lo/ 下（路径短），用完删除。用户数据目录指到测试目录，不写真实的 %LOCALAPPDATA%。"""
    monkeypatch.setattr(core, "profile_root", short_profile_root)      # 用产品里真正的 profile_root()
    monkeypatch.setenv("LOCALAPPDATA", str(isolated_temp / "appdata"))
    root = core.profile_root()
    assert root == isolated_temp / "appdata" / "lawbench" / "lo"
    core.convert_file("word2pdf", samples["docx"])
    assert len(str(root)) + 9 <= core.MAX_PROFILE_PATH
    assert list(root.iterdir()) == []                          # 用完删除


# ---------------------------------------------------------------- 第三轮返修：网络共享路径、删不掉要报、隔离


def test_long_form_shapes():
    B = chr(92)
    assert core.long_form("C:" + B + "a" + B + "b.doc") == B * 2 + "?" + B + "C:" + B + "a" + B + "b.doc"
    unc = B * 2 + "srv" + B + "share" + B + "x.doc"
    assert core.long_form(unc) == B * 2 + "?" + B + "UNC" + B + "srv" + B + "share" + B + "x.doc"
    done = B * 2 + "?" + B + "UNC" + B + "srv" + B + "share"
    assert core.long_form(done) == done                          # 已带前缀的不重复加
    # 映射成盘符的网络驱动器被 resolve() 还原后的样子（推断，本机不做映射）
    mapped = B * 2 + "fileserver" + B + "cases" + B + "张某甲案" + B + "借条.doc"
    assert core.long_form(mapped).startswith(B * 2 + "?" + B + "UNC" + B + "fileserver" + B)


def _unc_of(p: Path) -> Path:
    """本机路径 C:\\x → \\\\127.0.0.1\\c$\\x（管理共享）。"""
    B = chr(92)
    s = str(p.resolve())
    return Path(B * 2 + "127.0.0.1" + B + s[0].lower() + "$" + s[2:])


@pytest.mark.parametrize("kind,sample", [
    pytest.param("img2pdf", "jpg", id="jpg→pdf"),
    pytest.param("md2docx", "md", marks=need_pandoc, id="md→docx"),
    pytest.param("word2pdf", "docx", marks=need_lo, id="docx→pdf"),
])
def test_convert_from_network_share(samples, isolated_temp, kind, sample):
    """源文件在网络共享路径（\\\\127.0.0.1\\c$\\…）上：转换成功，结果写在原文件旁，原文件不变。"""
    d = isolated_temp / "share"
    d.mkdir()
    local = d / samples[sample].name
    shutil.copyfile(samples[sample], local)
    unc = _unc_of(local)
    if not os.path.exists(str(unc)):
        pytest.skip("本机管理共享不可访问")
    before = sha256(local)
    res = core.convert_many(kind, [unc])
    src, out, _ = res[0]
    assert isinstance(out, Path), out
    assert (d / core.OUT_DIR_NAME / out.name).is_file()
    assert sha256(local) == before


def test_profile_left_behind_is_reported(samples, isolated_temp, monkeypatch):
    """LibreOffice 配置目录重试后仍删不掉：本次转换照常成功，结果里带一条提示（Spec 5.2：不能静默）。"""
    if finder.find_soffice() is None:
        pytest.skip("本机没有 LibreOffice")
    monkeypatch.setenv("LOCALAPPDATA", str(isolated_temp / "appdata"))
    real = core._remove_workdir
    monkeypatch.setattr(core, "_remove_workdir", lambda w: False if w.parent == core.profile_root() else real(w))
    res = core.convert_many("word2pdf", [samples["docx"]])
    _, out, notes = res[0]
    assert isinstance(out, Path) and core.PROFILE_LEFT in notes


# ---------------------------------------------------------------- 第四轮返修 F1：提示在各条路径上都不丢


def _left_behind(monkeypatch):
    real = core._remove_workdir
    monkeypatch.setattr(core, "_remove_workdir", lambda w: False if w.parent == core.profile_root() else real(w))


@need_lo
def test_left_behind_notice_survives_convert_error(samples, tmp_path, monkeypatch):
    """转换失败（ConvertError）+ 配置目录删不掉：失败原因里带上提示。"""
    _left_behind(monkeypatch)
    bad = tmp_path / "坏.docx"
    bad.write_bytes(b"PK\x03\x04 not really a docx")
    res = core.convert_many("word2pdf", [bad])
    _, r, _ = res[0]
    assert isinstance(r, str) and core.PROFILE_LEFT in r
    with pytest.raises(core.ConvertError, match="临时目录没能删除"):
        core.convert_file("word2pdf", bad)


def test_left_behind_notice_survives_timeout(samples, tmp_path, monkeypatch):
    """超时 + 配置目录删不掉：失败原因里带上提示。"""
    import sys
    _left_behind(monkeypatch)
    fake = tmp_path / "soffice.cmd"
    fake.write_text(f'@"{sys.executable}" -c "import time; time.sleep(60)" %*\n', encoding="mbcs")
    monkeypatch.setattr(core, "_soffice", lambda: fake)
    monkeypatch.setattr(core, "TIMEOUT_S", 2)
    res = core.convert_many("word2pdf", [samples["docx"]])
    _, r, _ = res[0]
    assert "转换超时" in r and core.PROFILE_LEFT in r


def test_left_behind_notice_survives_other_exception(samples, monkeypatch):
    """LibreOffice 之后出现意外异常 + 配置目录删不掉：内部错误的原因里带上提示，不带堆栈。"""
    _left_behind(monkeypatch)
    monkeypatch.setattr(core, "_soffice", lambda: Path(__file__))   # 让 _run 之前的准备照常
    monkeypatch.setattr(core, "_run", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("内部细节 xyz")))
    res = core.convert_many("word2pdf", [samples["docx"]])
    _, r, _ = res[0]
    assert r.startswith(core.INTERNAL[:10]) and core.PROFILE_LEFT in r and "xyz" not in r


@need_lo
def test_left_behind_notice_on_success_stays_in_notes(samples, monkeypatch):
    """转换成功 + 删不掉：结果照常，提示在提示列表里。"""
    _left_behind(monkeypatch)
    res = core.convert_many("word2pdf", [samples["docx"]])
    _, out, notes = res[0]
    assert isinstance(out, Path) and core.PROFILE_LEFT in notes


def test_no_test_writes_real_profile_root(short_profile_root):
    """F3：测试期间 profile_root() 不指向真实的 %LOCALAPPDATA%\\lawbench\\lo。"""
    import os
    real = Path(os.environ.get("LOCALAPPDATA", "")) / "lawbench" / "lo"
    assert core.profile_root() != real and core.profile_root().parent.name.startswith("lbp-")


# ---------------------------------------------------------------- N24 ②：本版小工具不转换 .doc / .wps


@need_lo
@pytest.mark.parametrize("case", ["doc", "wps", "docx改名doc", "doc改名docx"])
def test_legacy_word_by_file_header(samples, tmp_path, monkeypatch, case):
    """按文件头判断：.doc、.wps、改名成 .docx 的 doc 都不转换；改名成 .doc 的 docx 照常按 docx 转换。"""
    src = {"doc": samples["doc"], "wps": samples["wps"]}.get(case)
    if case == "docx改名doc":
        src = tmp_path / "其实是docx.doc"
        shutil.copyfile(samples["docx"], src)
    if case == "doc改名docx":
        src = tmp_path / "其实是doc.docx"
        shutil.copyfile(samples["doc"], src)
    before = sha256(src)
    if case == "docx改名doc":
        out = core.convert_file("doc2docx", src)
        assert out.is_file()
    else:
        called = []
        monkeypatch.setattr(core, "_libreoffice", lambda *a, **k: called.append(a))
        kind = "word2pdf" if src.suffix == ".docx" else "doc2docx"
        with pytest.raises(core.ConvertError, match="请用 Word 或 WPS 打开后另存为 .docx"):
            core.convert_file(kind, src)
        assert called == []
    assert sha256(src) == before


@need_lo
def test_batch_with_legacy_word_not_interrupted(samples, tmp_path):
    """批量里夹着旧版文件：只跳过它并给提示，前后文件照常转换。"""
    res = core.convert_many("word2pdf", [samples["docx"], samples["doc"], samples["docx"]])
    assert isinstance(res[0][1], Path) and isinstance(res[2][1], Path)
    assert isinstance(res[1][1], str) and "暂不支持旧版 Word / WPS 文件" in res[1][1]


def test_xls_still_converts_with_ole_header(samples):
    """.xls 也是 OLE 文件头，但走 xls → xlsx，不受影响。"""
    if "xls" not in samples:
        pytest.skip("本机没有 LibreOffice")
    assert samples["xls"].read_bytes()[:8] == bytes.fromhex("d0cf11e0a1b11ae1")    # OLE 复合文档文件头
    assert core.convert_file("xls2xlsx", samples["xls"]).is_file()


def test_rtf_reason_says_maybe_not_word(tmp_path):
    """RTF 改名 .docx：提示里带"或文件不是 Word 文档"。"""
    f = tmp_path / "其实是rtf.docx"
    f.write_bytes(b"{\rtf1 hello}")
    with pytest.raises(core.ConvertError, match="或文件不是 Word 文档"):
        core.convert_file("word2pdf", f)


def test_legacy_word_hidden_in_gui_and_messages(tmp_path):
    """下拉框不出现 doc / wps → docx；Word → PDF 的选文件筛选和扩展名提示只列 .docx。"""
    assert "doc2docx" not in [k.key for k in core.VISIBLE]
    assert core.BY_KEY["word2pdf"].listed == (".docx",)
    f = tmp_path / "a.txt"
    f.write_text("x", encoding="utf-8")
    with pytest.raises(core.ConvertError) as e:
        core.convert_file("word2pdf", f)
    assert ".doc、" not in str(e.value) and ".wps" not in str(e.value)


def test_legacy_word_note_in_gui():
    import tkinter as tk
    from convert.app import ConvertApp
    try:
        root = tk.Tk()
    except tk.TclError:
        pytest.skip("没有图形界面环境")
    root.withdraw()
    app = ConvertApp(root)
    from tkinter import ttk
    boxes = [w for f in root.winfo_children() for w in f.winfo_children() if isinstance(w, ttk.Combobox)]
    assert boxes and "doc / wps → docx" not in boxes[0]["values"]
    app.kind.set(core.BY_KEY["word2pdf"].label)
    app._update_note()
    assert "另存为 .docx" in app.note["text"]
    root.destroy()
