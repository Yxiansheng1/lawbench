"""civil-01：李某乙诉王某丙民间借贷纠纷（虚构）。

借条 docx（一处插入、一处删除的修订痕迹，含表格、页眉页脚）、银行流水 xlsx（两个工作表，
一格只有公式没有缓存值）、csv（gb18030）、txt、md。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook

import common as C

CASE = "civil-01"
FEAT = C.FEATURE[CASE]

FILES = ["借条.docx", "银行流水.xlsx", "还款记录.csv", "情况说明.txt", "案情摘要.md"]

# 借条正文段落（段号按 Spec 5.2：非空段落从 1 起，表格整体算一段）
JIETIAO = [
    "借条",                                                                                   # 1
    "今借到出借人李某乙人民币捌万元整（￥80,000.00），用于本人经营资金周转。",                       # 2
    "借款日期：2025年3月10日。",                                                                # 3
    None,  # 4 插入修订段：借款年利率为百分之十二，按月付息，到期还本。
    "还款期限：2025年9月10日前一次性归还全部本金。",                                             # 5
    None,  # 6 删除修订段
    "TABLE",                                                                                  # 7 当事人信息表
    "借款人（签字）：王某丙",                                                                    # 8
    "2025年3月10日",                                                                          # 9
    f"样本编号：{FEAT}",                                                                       # 10
]

REVISIONS = {
    4: [("keep", "借款年利率为百分之十二，按月付息"), ("ins", "，到期还本"), ("keep", "。")],
    6: [("keep", "借款人逾期还款的，"), ("del", "应按日加收千分之五的违约金，并"),
        ("keep", "应承担出借人为实现债权支出的律师费。")],
}

TABLE = [
    ["身份", "姓名", "公民身份号码（虚构）", "住址（虚构）"],
    ["出借人", "李某乙", "990103197901010000", "虚构市虚构区不存在街8号"],
    ["借款人", "王某丙", "990104198203030000", "虚构市虚构区无此巷16号"],
]

FLOW_ROWS = [
    ["交易日期", "摘要", "收入", "支出", "余额", "对方户名", "备注"],
    ["2025-03-01", "工资", 12000, None, 95000, "虚构市某某科技有限公司", ""],
    ["2025-03-10", "转账", None, 80000, 15000, "王某丙", "借款80,000元（借条）"],
    ["2025-04-10", "转账", 800, None, 15800, "王某丙", "利息"],
    ["2025-05-10", "转账", 800, None, 16600, "王某丙", "利息"],
    ["2025-06-10", "转账", 20000, None, 36600, "王某丙", "还款"],
    ["2025-07-10", "转账", 600, None, 37200, "王某丙", "利息"],
    ["2025-09-12", "消费", None, 3200, 34000, "虚构超市", ""],
]

SUMMARY_ROWS = [
    ["项目", "金额", "说明"],
    ["借出本金", 80000, "2025-03-10 转出"],
    ["已收回本金", 20000, "2025-06-10"],
    ["已收利息", None, "公式单元格，文件未存计算结果"],   # B4 = 公式
    ["样本编号", FEAT, ""],
]

CSV_LINES = [
    "日期,金额,方式,备注",
    "2025-03-10,80000,手机银行,借款本金",
    "2025-04-10,800,微信转账,3月利息",
    "2025-05-10,800,微信转账,4月利息",
    "2025-06-10,20000,手机银行,归还部分本金",
    "2025-07-10,600,微信转账,5月利息",
    f"合计,102200,,样本编号{FEAT}",
]

TXT_LINES = [
    "情况说明",
    "",
    "本人李某乙，与王某丙系朋友关系，经朋友孙某介绍相识。",
    "2025.3.10，王某丙称经营周转困难，向我借款8万元，当天我通过手机银行转账给他，他出具了借条。",
    "借款后王某丙支付了三个月利息，于2025年6月10日归还本金2万元，此后再未还款。",
    "2025年9月10日还款期限届满后，我多次催要，王某丙均以资金紧张为由拖延。",
    "孙某知道借款的经过，可以作证。",
    "",
    "说明人：李某乙",
    "2026年2月1日",
    f"样本编号：{FEAT}",
]

MD_LINES = [
    "# 案情摘要",
    "",
    "- 案号：（2026）虚0103民初2046号",
    "- 案由：民间借贷纠纷",
    "- 原告：李某乙；被告：王某丙",
    "- 诉讼请求：归还借款本金60,000元及逾期利息",
    "",
    "## 争议焦点",
    "",
    "1. 借条约定的利率是否有效；",
    "2. 2025年6月10日的20,000元是归还本金还是支付利息。",
    "",
    f"样本编号：{FEAT}",
]


def _jietiao(path: Path) -> None:
    doc = Document()
    C.set_east_asian_font(doc)
    C.fix_core(doc, "借条", FEAT)
    C.header_footer(doc, f"借条（虚构测试样本 {FEAT}）", "第1页 共1页　本借条一式两份")
    for no, item in enumerate(JIETIAO, start=1):
        if item == "TABLE":
            t = doc.add_table(rows=len(TABLE), cols=len(TABLE[0]))
            t.style = "Table Grid"
            for i, row in enumerate(TABLE):
                for j, v in enumerate(row):
                    t.cell(i, j).text = v
        elif item is None:
            C.revision_paragraph(doc.add_paragraph(), REVISIONS[no])
        else:
            doc.add_paragraph(item)
    doc.save(path)


def _flow(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "流水"
    for r in FLOW_ROWS:
        ws.append(r)
    ws2 = wb.create_sheet("汇总")
    for r in SUMMARY_ROWS:
        ws2.append(r)
    ws2["B4"] = "=SUM(流水!C4:C5,流水!C7)"   # 只有公式、没有缓存值（openpyxl 不写计算结果）
    wb.properties.creator = "lawbench 虚构样本"
    wb.properties.keywords = FEAT
    wb.properties.created = C.FIXED_TIME
    wb.properties.modified = C.FIXED_TIME
    wb.save(path)


def build(root: Path) -> None:
    d = root / CASE
    d.mkdir(parents=True, exist_ok=True)
    _jietiao(d / "借条.docx")
    _flow(d / "银行流水.xlsx")
    (d / "还款记录.csv").write_bytes(("\r\n".join(CSV_LINES) + "\r\n").encode("gb18030"))
    (d / "情况说明.txt").write_text("\n".join(TXT_LINES) + "\n", encoding="utf-8")
    (d / "案情摘要.md").write_text("\n".join(MD_LINES) + "\n", encoding="utf-8")
