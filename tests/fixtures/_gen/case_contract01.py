"""contract-01：采购合同（虚构）——修订版生成与范围外情况的测试对象。

采购合同.docx：正文段落 ≥ 40，含一个表格、页眉页脚、一个超链接（位于句中，用于"find 跨越超链接"）。
采购合同-含未处理修订.docx：同一份合同，另带一处未接受的插入和一处未接受的删除。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

import common as C

CASE = "contract-01"
FEAT = C.FEATURE[CASE]

FILES = ["采购合同.docx", "采购合同-含未处理修订.docx"]

HYPER = "HYPERLINK"  # 该段在句中插入超链接
TABLE = "TABLE"

PARAS = [
    "钢材采购合同",                                                                         # 1
    "合同编号：QH-CG-2026-0088",                                                            # 2
    "甲方（采购方）：虚构市青禾建材贸易有限公司（以下简称“青禾建材”）",                      # 3
    "住所：虚构省虚构市虚构区不存在大道100号",                                                  # 4
    "法定代表人：冯某",                                                                      # 5
    "乙方（供货方）：虚构市某某钢构有限公司（以下简称“某某钢构”）",                         # 6
    "住所：虚构省虚构市虚构区无此路200号",                                                      # 7
    "法定代表人：褚某",                                                                      # 8
    "甲乙双方根据《中华人民共和国民法典》及相关法律法规，经平等协商，就甲方向乙方采购钢材事宜订立本合同。",   # 9
    "第一条　标的物",                                                                        # 10
    "1.1 标的物的名称、规格、数量、单价见下表：",                                                # 11
    TABLE,                                                                                  # 12
    "1.2 上表单价已含增值税及运至甲方指定地点的运输费用。",                                        # 13
    "第二条　质量标准",                                                                       # 14
    "2.1 标的物应符合国家现行标准，并随货提供产品质量证明书。",                                     # 15
    "2.2 青禾建材有权委托第三方检测机构对标的物进行抽检，检测费用先由青禾建材垫付。",                   # 16
    "2.3 抽检不合格的，检测费用由乙方承担，甲方有权拒收该批货物。",                                  # 17
    "第三条　交货",                                                                          # 18
    "3.1 交货时间：2026年5月20日前分两批交付。",                                                # 19
    "3.2 交货地点：虚构市虚构区不存在大道100号青禾建材一号仓库。",                                   # 20
    "3.3 乙方应提前两日通知甲方收货，甲方应安排人员清点签收。",                                      # 21
    "第四条　验收",                                                                          # 22
    "4.1 甲方应于到货后七日内完成验收，逾期未提出书面异议的，视为验收合格。",                          # 23
    "4.2 隐蔽瑕疵的异议期为验收合格之日起六个月。",                                                # 24
    "第五条　价款与支付",                                                                     # 25
    "5.1 合同总价款为人民币壹佰贰拾陆万元整（￥1,260,000.00）。",                                  # 26
    "5.2 甲方应于收到货物后九十日内付款。",                                                       # 27
    "5.3 乙方应在甲方付款前开具等额增值税专用发票。",                                               # 28
    "第六条　违约责任",                                                                       # 29
    "6.1 乙方逾期交货的，每逾期一日，按逾期交货部分价款的千分之一向甲方支付违约金。",                    # 30
    "6.2 甲方逾期付款的，每逾期一日，按逾期付款金额的万分之五向乙方支付违约金。",                        # 31
    "6.3 违约金不足以弥补损失的，违约方还应赔偿损失。",                                              # 32
    "第七条　不可抗力",                                                                       # 33
    "7.1 因不可抗力不能履行合同的，根据不可抗力的影响部分或全部免除责任。",                             # 34
    "7.2 遭遇不可抗力的一方应在三日内书面通知对方，并在十五日内提供证明。",                             # 35
    "第八条　争议解决",                                                                       # 36
    "8.1 因本合同发生的争议，双方协商解决；协商不成的，提交甲方所在地人民法院诉讼解决。",                   # 37
    "第九条　通知与送达",                                                                     # 38
    HYPER,                                                                                  # 39
    "9.2 一方变更通知地址的，应于变更后三日内书面通知对方。",                                         # 40
    "第十条　其他",                                                                          # 41
    "10.1 本合同自双方签字盖章之日起生效。",                                                     # 42
    "10.2 本合同一式四份，双方各执两份，具有同等法律效力。",                                         # 43
    "甲方（盖章）：虚构市青禾建材贸易有限公司　　乙方（盖章）：虚构市某某钢构有限公司",                    # 44
    "签订日期：2026年4月1日",                                                                  # 45
    f"样本编号：{FEAT}",                                                                     # 46
]

HYPER_PARTS = ("9.1 双方指定的电子邮件通知地址以本合同附件为准，附件的电子版可在",
               "http://127.0.0.1/lbfx-attachment", "查阅，通知自发出之日起视为送达。")

TABLE_ROWS = [
    ["序号", "名称", "规格", "数量（吨）", "单价（元/吨）", "金额（元）"],
    ["1", "热轧H型钢", "HW200×200", "150", "4,200", "630,000"],
    ["2", "热轧H型钢", "HN300×150", "150", "4,200", "630,000"],
]


def _build_doc(path: Path, with_revisions: bool) -> None:
    doc = Document()
    C.set_east_asian_font(doc, size_pt=11)
    C.fix_core(doc, "钢材采购合同", FEAT)
    C.header_footer(doc, "钢材采购合同（QH-CG-2026-0088）", f"虚构测试样本 {FEAT}　第1页")
    for no, item in enumerate(PARAS, start=1):
        if item == TABLE:
            t = doc.add_table(rows=len(TABLE_ROWS), cols=len(TABLE_ROWS[0]))
            t.style = "Table Grid"
            for i, row in enumerate(TABLE_ROWS):
                for j, v in enumerate(row):
                    t.cell(i, j).text = v
        elif item == HYPER:
            p = doc.add_paragraph(HYPER_PARTS[0])
            C.add_hyperlink(p, HYPER_PARTS[1], HYPER_PARTS[1])
            p.add_run(HYPER_PARTS[2])
        elif with_revisions and no == 27:
            C.revision_paragraph(doc.add_paragraph(), [("keep", "5.2 甲方应于收到货物后"),
                                                      ("del", "九十"), ("ins", "六十"), ("keep", "日内付款。")])
        else:
            doc.add_paragraph(item)
    doc.save(path)


def hyper_para_text() -> str:
    return "".join(HYPER_PARTS)


def build(root: Path) -> None:
    d = root / CASE
    d.mkdir(parents=True, exist_ok=True)
    _build_doc(d / "采购合同.docx", with_revisions=False)
    _build_doc(d / "采购合同-含未处理修订.docx", with_revisions=True)
