"""核对 tests\\fixtures\\ 的样本是否完整、可用。

用法：
    python tests\\fixtures\\_gen\\verify.py               # 核对 tests\\fixtures\\
    python tests\\fixtures\\_gen\\verify.py --root <目录>  # 核对 make_all.py --out 生成的目录

核对项：
1. 每个案件含 README.md 所列文件（链接类文件生成失败时记"跳过"，不算缺失）；
2. 每份材料含本案特征字符串（没有文字层的扫描页、图片用渲染前文字核对；加密文件解密后核对）；
3. search-cases.json 每条 query 在对应材料、对应位置的源文字中找得到（found_as 条目核对扩展写法）；
4. cite-cases.json 每条原文片段在对应位置找得到，出处写法与位置对象符合 contracts\\common.schema.json；
5. sentence-cases.json 每条的输入符合 contracts\\tools\\case_calc_sentence.schema.json 的 args；
6. 样本特征：扫描页无文字层、图文混排页的文字与图片面积、docx 修订、xlsx 无缓存公式、csv 编码、
   加密与损坏文件确实打不开、ZIP 内容；
7. 8 个样本目录总大小 < 30 MB。
全部通过退出码 0，否则 1。
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import common as C  # noqa: E402
import case_attack01  # noqa: E402
import case_closed01  # noqa: E402
import case_criminal01  # noqa: E402
from make_all import CASES, tree_size  # noqa: E402

REPO = C.FIXTURES.parent.parent
sys.path.insert(0, str(REPO / "contracts"))
from check_examples import validator  # noqa: E402  契约校验的统一写法
SCAN_TEXT = {("criminal-01", k): v for k, v in case_criminal01.SCAN_TEXT.items()}
SCAN_TEXT.update({("closed-01", k): v for k, v in case_closed01.SCAN_TEXT.items()})

fails: list[str] = []
passes = 0


def check(ok: bool, what: str) -> None:
    global passes
    if ok:
        passes += 1
    else:
        fails.append(what)
        print(f"  [失败] {what}")


def squash(s: str) -> str:
    return re.sub(r"\s+", "", s)


# ---------------------------------------------------------------- 读材料


def pdf_pages(p: Path | bytes) -> list[str]:
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(p if isinstance(p, bytes) else str(p))
    return [doc[i].get_textpage().get_text_range() for i in range(len(doc))]


def xlsx_cells(p: Path) -> dict[tuple[str, str], str]:
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=False)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out[(ws.title, c.coordinate)] = str(c.value)
    return out


def text_lines(p: Path) -> list[str]:
    raw = p.read_bytes()
    enc = "gb18030" if p.suffix.lower() == ".csv" else "utf-8"
    return raw.decode(enc).splitlines()


def material_units(case: str, rel: str, root: Path):
    """返回 (unit, {位置: 文字})；位置为页/段/行号，或 (工作表, 单元格)。"""
    p = root / case / rel
    ext = p.suffix.lower()
    if (case, rel) in SCAN_TEXT:
        return "page", {i: t for i, t in enumerate(SCAN_TEXT[(case, rel)], 1)}
    if ext == ".pdf":
        return "page", {i: t for i, t in enumerate(pdf_pages(p), 1)}
    if ext == ".docx":
        return "para", {i: t for i, t in enumerate(C.body_paragraphs(p), 1)}
    if ext == ".xlsx":
        return "cell", xlsx_cells(p)
    if ext in (".csv", ".txt", ".md"):
        return "line", {i: t for i, t in enumerate(text_lines(p), 1)}
    raise ValueError(f"不认识的材料类型：{rel}")


def name_map(case: str, root: Path) -> dict[str, str]:
    """材料名（文件名去扩展名；本样本内不重名）→ 相对路径。"""
    out = {}
    for p in (root / case).rglob("*"):
        if p.is_file() and not p.is_symlink():
            out[p.stem] = p.relative_to(root / case).as_posix()
    return out


def at_loc(units, loc: dict) -> str | None:
    unit, table = units
    if loc["unit"] != unit:
        return None
    if unit == "cell":
        return table.get((loc["sheet"], loc["ref"]))
    lo, hi = loc["from"], loc.get("to", loc["from"])
    parts = [table.get(i) for i in range(lo, hi + 1)]
    return None if any(x is None for x in parts) else "\n".join(parts)


# ---------------------------------------------------------------- README


def readme_files() -> dict[str, list[str]]:
    text = (C.FIXTURES / "README.md").read_text(encoding="utf-8")
    out: dict[str, list[str]] = {}
    case = None
    for line in text.splitlines():
        m = re.match(r"^###\s+(\S+)", line)
        if m:
            case = m.group(1) if m.group(1) in CASES else None
            if case:
                out[case] = []
            continue
        m = re.match(r"^\|\s*`([^`]+)`\s*\|", line)
        if case and m:
            out[case].append(m.group(1))
    return out


# ---------------------------------------------------------------- 检索扩展


def _date_key(s: str):
    m = re.fullmatch(r"(\d{4})(?:年|-|\.)(\d{1,2})(?:月|-|\.)(\d{1,2})日?", s)
    return tuple(int(x) for x in m.groups()) if m else None


def _money_key(s: str):
    m = re.fullmatch(r"(\d+(?:\.\d+)?)万", s)
    if m:
        return round(float(m.group(1)) * 10000)
    if re.fullmatch(r"\d{1,3}(,\d{3})+|\d+", s):
        return int(s.replace(",", ""))
    return None


def equivalent(query: str, found: str) -> bool:
    for key in (_date_key, _money_key):
        a, b = key(query), key(found)
        if a is not None and a == b:
            return True
    return False


# ---------------------------------------------------------------- 各项核对


def check_files(root: Path, listed: dict[str, list[str]]) -> None:
    print("1. README 所列文件")
    check(set(listed) == set(CASES), f"README 应列出 8 个案件，实际 {sorted(listed)}")
    for case, files in listed.items():
        check(bool(files), f"{case}：README 没有列出文件")
        for rel in files:
            p = root / case / rel
            if rel in case_attack01.LINKS:
                if p.exists() or p.is_symlink():
                    check(p.is_symlink() or _isjunction(p), f"{case}/{rel} 应为链接")
                else:
                    print(f"  [跳过] {case}/{rel}（链接未生成，见 make_all 输出）")
                continue
            check(p.is_file(), f"{case}/{rel} 缺失")


def _isjunction(p: Path) -> bool:
    import os
    return os.path.isjunction(p)


def feature_text(case: str, rel: str, root: Path) -> str:
    p = root / case / rel
    ext = p.suffix.lower()
    if (case, rel) in SCAN_TEXT:
        return "\n".join(SCAN_TEXT[(case, rel)])
    if case == "broken":
        if ext == ".docx":
            from msoffcrypto import OfficeFile
            of = OfficeFile(open(p, "rb"))
            of.load_key(password=C.TEST_PASSWORD)
            buf = io.BytesIO()
            of.decrypt(buf)
            return C.docx_all_text(buf)
        if rel == "加密.pdf":
            from pypdf import PdfReader
            r = PdfReader(p)
            r.decrypt(C.TEST_PASSWORD)
            return "".join(pg.extract_text() for pg in r.pages)
        return p.read_bytes().decode("latin-1")
    if ext == ".pdf":
        return "".join(pdf_pages(p))
    if ext == ".docx":
        if rel.split("/")[-1].startswith("~$"):
            return p.read_bytes().decode("latin-1")
        return C.docx_all_text(p)
    if ext == ".xlsx":
        return "\n".join(xlsx_cells(p).values())
    if ext in (".csv", ".txt", ".md"):
        return "\n".join(text_lines(p))
    if ext == ".zip":
        with zipfile.ZipFile(p) as z:
            return "\n".join("".join(pdf_pages(z.read(n))) for n in z.namelist())
    raise ValueError(rel)


def check_features(root: Path, listed: dict[str, list[str]]) -> None:
    print("2. 特征字符串")
    for case, files in listed.items():
        feat = C.FEATURE[case]
        for rel in files:
            if rel in case_attack01.LINKS:
                continue
            try:
                txt = feature_text(case, rel, root)
            except Exception as e:  # noqa: BLE001
                check(False, f"{case}/{rel} 读取失败：{e}")
                continue
            check(feat in squash(txt), f"{case}/{rel} 不含特征字符串 {feat}")


def check_search(root: Path) -> None:
    print("3. search-cases.json")
    cases = json.loads((C.FIXTURES / "search-cases.json").read_text(encoding="utf-8"))
    check(len(cases) >= 20, f"search-cases 应 ≥ 20 条，实际 {len(cases)}")
    kinds = {c["kind"] for c in cases}
    check({"两字姓名", "简称", "日期", "金额", "案号"} <= kinds, f"search-cases 类别不全：{kinds}")
    for kind, forms in (("日期", {"年", "-", "."}), ("金额", {"万", ",", "plain"})):
        seen = set()
        for c in cases:
            if c["kind"] == kind and "found_as" not in c:
                q = c["query"]
                seen |= {f for f in forms if f != "plain" and f in q}
                if kind == "金额" and re.fullmatch(r"\d+", q):
                    seen.add("plain")
        check(seen == forms, f"{kind}三种写法没有都覆盖：{seen}")
    for i, c in enumerate(cases, 1):
        names = name_map(c["case"], root)
        rel = names.get(c["expect_material"])
        if rel is None:
            check(False, f"search #{i}：找不到材料 {c['case']}/{c['expect_material']}")
            continue
        src = at_loc(material_units(c["case"], rel, root), c["expect_loc"])
        target = c.get("found_as", c["query"])
        check(src is not None and squash(target) in squash(src),
              f"search #{i}：{c['query']!r} 不在 {c['expect_material']} {c['expect_loc']}")
        if "found_as" in c:
            check(equivalent(c["query"], c["found_as"]),
                  f"search #{i}：{c['query']!r} 与 {c['found_as']!r} 不是同一日期或金额的不同写法")


def check_cite(root: Path) -> None:
    print("4. cite-cases.json")
    cases = json.loads((C.FIXTURES / "cite-cases.json").read_text(encoding="utf-8"))
    check(len(cases) >= 15, f"cite-cases 应 ≥ 15 条，实际 {len(cases)}")
    check({c["loc"]["unit"] for c in cases} >= {"page", "para", "cell", "line"}, "cite-cases 没有覆盖页、段、单元格、行")
    common = json.loads((REPO / "contracts" / "common.schema.json").read_text(encoding="utf-8"))
    cite_re = re.compile(common["$defs"]["citation_text"]["pattern"])
    loc_v = validator("common.schema.json", "#/$defs/loc")
    for i, c in enumerate(cases, 1):
        check(bool(cite_re.fullmatch(c["citation"])), f"cite #{i}：出处写法不符合 citation_text：{c['citation']}")
        errs = list(loc_v.iter_errors(c["loc"]))
        check(not errs, f"cite #{i}：loc 不符合 $defs/loc：{[e.message for e in errs]}")
        rel = name_map(c["case"], root).get(c["material"])
        if rel is None:
            check(False, f"cite #{i}：找不到材料 {c['case']}/{c['material']}")
            continue
        src = at_loc(material_units(c["case"], rel, root), c["loc"])
        check(src is not None and squash(c["text"]) in squash(src),
              f"cite #{i}：{c['text']!r} 不在 {c['material']} {c['loc']}")


def check_sentence() -> None:
    print("5. sentence-cases.json")
    cases = json.loads((C.FIXTURES / "sentence-cases.json").read_text(encoding="utf-8"))
    check(len(cases) >= 10, f"sentence-cases 应 ≥ 10 条，实际 {len(cases)}")
    v = validator("tools/case_calc_sentence.schema.json", "#/$defs/args")
    for c in cases:
        args = {k: c[k] for k in ("penalty", "years", "months", "execution_start", "custody") if k in c}
        errs = list(v.iter_errors(args))
        check(not errs, f"sentence {c['id']}：输入不符合契约：{[e.message for e in errs]}")
        check(bool(c.get("note")), f"sentence {c['id']}：缺少 note 算式")


def check_traits(root: Path) -> None:
    print("6. 样本特征")
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c

    crim = root / "criminal-01"
    pages = pdf_pages(crim / "起诉意见书.pdf")
    check(len(pages) == 5 and all(len(squash(t)) >= 30 for t in pages), "起诉意见书应为 5 页文字版")
    pages = pdf_pages(crim / "讯问笔录.pdf")
    check(len(pages) == 3 and all(len(squash(t)) < 30 for t in pages), "讯问笔录应为 3 页、无文字层（需识别）")

    doc = pdfium.PdfDocument(str(crim / "现场勘验图文.pdf"))
    pg = doc[0]
    w, h = pg.get_size()
    area = 0.0
    for obj in pg.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE]):
        l, b, r, t = obj.get_bounds()
        area += (r - l) * (t - b)
    text = pg.get_textpage().get_text_range()
    check(len(doc) == 1 and len(squash(text)) >= 30 and area / (w * h) > 0.3,
          f"现场勘验图文应为 1 页、≥30 字、图片 >30%（实际 {len(squash(text))} 字、{area / (w * h):.0%}）")

    def rev_counts(p: Path) -> tuple[int, int]:
        with zipfile.ZipFile(p) as z:
            x = z.read("word/document.xml").decode("utf-8")
        return x.count("<w:ins "), x.count("<w:del ")

    civ = root / "civil-01"
    check(rev_counts(civ / "借条.docx") == (1, 1), "借条应有一处插入、一处删除修订")
    with zipfile.ZipFile(civ / "借条.docx") as z:
        names = z.namelist()
        check(any(n.startswith("word/header") for n in names) and any(n.startswith("word/footer") for n in names),
              "借条应有页眉页脚")
        check("<w:tbl>" in z.read("word/document.xml").decode("utf-8"), "借条应含表格")
    from openpyxl import load_workbook
    wb_v = load_workbook(civ / "银行流水.xlsx", data_only=True)
    wb_f = load_workbook(civ / "银行流水.xlsx", data_only=False)
    check(len(wb_v.sheetnames) == 2, "银行流水应有两个工作表")
    formula_cells = [(ws.title, c.coordinate) for ws in wb_f.worksheets for row in ws.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith("=")]
    check(formula_cells == [("汇总", "B4")] and wb_v["汇总"]["B4"].value is None,
          f"银行流水应只有 汇总!B4 一个公式且无缓存值（实际 {formula_cells}）")
    raw = (civ / "还款记录.csv").read_bytes()
    try:
        raw.decode("utf-8")
        check(False, "还款记录.csv 应为 gb18030 编码（不是 utf-8）")
    except UnicodeDecodeError:
        check(raw.decode("gb18030").startswith("日期,"), "还款记录.csv 应能按 gb18030 解码")

    con = root / "contract-01"
    paras = C.body_paragraphs(con / "采购合同.docx")
    check(len(paras) >= 40, f"采购合同正文段落应 ≥ 40（实际 {len(paras)}）")
    with zipfile.ZipFile(con / "采购合同.docx") as z:
        x = z.read("word/document.xml").decode("utf-8")
        check("<w:hyperlink" in x and "<w:tbl>" in x, "采购合同应含超链接和表格")
        check(any(n.startswith("word/header") for n in z.namelist()), "采购合同应有页眉页脚")
    check(rev_counts(con / "采购合同.docx") == (0, 0), "采购合同原件不应有修订")
    check(rev_counts(con / "采购合同-含未处理修订.docx") != (0, 0), "采购合同副本应含未处理修订")

    brk = root / "broken"
    from pypdf import PdfReader
    check(PdfReader(brk / "加密.pdf").is_encrypted, "加密.pdf 应加密")
    import olefile
    ole = olefile.OleFileIO(str(brk / "加密.docx"))
    check(ole.exists("EncryptionInfo"), "加密.docx 应为含 EncryptionInfo 流的 OLE 容器")
    ole.close()
    try:
        pdfium.PdfDocument(str(brk / "截断.pdf"))
        check(False, "截断.pdf 应无法打开")
    except pdfium.PdfiumError:
        check(True, "")

    inv = root / "invoices-01"
    with zipfile.ZipFile(inv / "邮件附件-发票两张.zip") as z:
        members = sorted(z.namelist())
    check(members == sorted(["发票02-差旅住宿.pdf", "发票03-交通.pdf"]), f"ZIP 应装发票02、03（实际 {members}）")
    t1 = squash("".join(pdf_pages(inv / "发票01-办公用品.pdf")))
    t4 = squash("".join(pdf_pages(inv / "发票04-办公用品-重复.pdf")))
    check(t1 == t4, "发票04 应与发票01 内容相同（重复发票）")
    t5 = squash("".join(pdf_pages(inv / "发票05-购买方不符.pdf")))
    check("购买方名称：某某虚构律师事务所" not in t5, "发票05 的购买方应与律所不符")


def main() -> int:
    ap = argparse.ArgumentParser(description="核对虚构样本")
    ap.add_argument("--root", type=Path, default=C.FIXTURES)
    root = ap.parse_args().root.resolve()
    listed = readme_files()
    check_files(root, listed)
    check_features(root, listed)
    check_search(root)
    check_cite(root)
    check_sentence()
    check_traits(root)
    print("7. 大小")
    total = sum(tree_size(root / c) for c in CASES)
    check(total < 30 * 1024 * 1024, f"样本总大小 {total / 1024 / 1024:.2f} MB 超过 30 MB")
    print(f"  8 个样本目录共 {total / 1024 / 1024:.2f} MB")
    print(f"\n通过 {passes} 项，失败 {len(fails)} 项。")
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
